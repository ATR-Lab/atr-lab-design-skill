#!/usr/bin/env python3
"""new_deck.py - build a real ATR Lab deck (or any ATR template piece) from a JSON or YAML spec.

A generic template-fill engine on python-pptx. It opens one of the ATR templates in assets/templates/,
removes the teaching (showcase) slides cleanly, adds one slide per spec entry from the template's own
layouts, types your text into the layout placeholders (so fonts, sizes, colors, bullets and spacing come
from the template), inserts pictures, native charts and native tables, deletes every placeholder you did
not fill, writes speaker notes, and then checks the result.

Spec (JSON, or YAML when PyYAML is installed):

  {
    "template": "presentation",            # or quad, poster-48x36, poster-36x48, flyer, one-pager,
                                           # certificate, door-sign, badge, social-square ... or a .pptx path
    "output": "talk.pptx",                 # relative to the current folder; -o overrides it
    "title": "Deck title for File > Info", # optional document properties: title, author, subject, keywords
    "slides": [
      {"layout": "ATR - Title",
       "fields": {"eyebrow": "IEEE RO-MAN 2026  ·  SEPT. 23, 2026", "title": "...", "subtitle": "...",
                  "presenter": "..."},
       "notes": "Speaker notes."},
      {"layout": "ATR - Title + Content", "section": 2,          # eyebrow "SECTION 02  ·  <agenda item 2>"
       "fields": {"title": "...",
                  "body": ["Bullet", ["Sub-bullet"], "Bullet with **bold** words"],
                  "image": "assets/illustrations/illus-telepresence.png",
                  "image_alt": "Illustration: ..."}},
      {"layout": "ATR - Chart + Takeaway",
       "fields": {"chart": {"type": "bar", "categories": ["A", "B"],
                            "series": [{"name": "Mean seconds", "values": [84, 47]}],
                            "number_format": "0 \\"s\\"", "highlight": "B", "alt": "..."}}},
      {"layout": "ATR - Table",
       "fields": {"table": {"header": ["Milestone", "Due", "Status"],
                            "rows": [["Prototype", "2026-12-01", {"status": "on-track"}]]}}}
    ]
  }

Field names come from assets/templates/presentation-layouts.json and quad-layouts.json (name / key), or,
for templates without a JSON, from the layout's placeholder names (ph<idx> when a name is generic). Aliases:
"image"/"photo"/"figure" = the first picture slot, "chart" and "table" = the chart/table slot or else the
picture slot's frame (or the content area), "video" = the media slot, "body" = the main text slot.
"<field>_alt" is the alt text of a picture field ("" = decorative). Text values: a string ("\\n" starts a
paragraph, a leading Tab makes it level 2), or a list where a nested list is one level deeper; **word** is
bold and ^{1} is superscript (affiliation marks). An empty string deletes the placeholder. Run
--list-layouts <template> for every name and budget.

Slide keys: layout, fields (or put the fields at slide level), notes (speaker notes; "speaker_notes" is the
same, and needed on ATR - Video, whose "what to watch for" placeholder is itself called notes: put that one
in "fields"), section (1-based agenda item: fills eyebrow, divider number, title and outline), from_slide (clone showcase
slide N of the template instead of starting from the layout: posters, flyers, the quad milestone table),
remove (shape or field names to delete), keep (placeholders to keep as they are on a cloned slide), shapes
({"shape name": text} on a cloned slide), tables ({"table name": {"header", "rows"}} on a cloned slide),
images ([{"path", "alt", "box": {"x","y","w","h"} | "replace": "shape name", "fit"}]),
textboxes ([{"text", "box", "size", "color", "bold", "font"}]).

Checks: text over a placeholder's max_chars / max_lines / max_paragraphs is a warning (content titles over
max_chars get the template's fallback_pt, the one sanctioned override), and every text placeholder is also
measured with the brand TTFs in assets/fonts/ at the layout's own size, indent, line spacing and insets, so a
line that wraps or a block that overflows its box is a warning even inside the character budget; tables are
measured too, and columns are sized from their measured text. A content slide whose body (Title + Content,
Two Content) holds under a third of its line budget gets an info hint naming a layout sized for that much
content (Statement, Key Numbers, Content + Image, Three Icon Columns) or asking for a visual; title, section,
statement, closing and picture-led layouts never get it. [Bracketed] placeholders, lorem ipsum, TODO and
"Click to add" anywhere on the slides (and brackets in the speaker notes) are REFUSED unless
--allow-placeholders; the facts brand_check.py's wording check treats as errors or unverified (@atr_kent,
College of Arts and Sciences, Room 236, 330-672-9060, the old building name, RoboCup placements) are refused
on slides and warned in notes; every picture, chart and video needs alt text; the package is re-read after
saving to prove there are no orphan parts. Pictures below 100 ppi and heavy crops are warnings.

Exit status: 0 built; 1 spec or usage error; 2 refused (placeholders, wording, missing alt text);
3 built but a post-build check failed (package check, --check errors, --strict warnings, render failure).

Requires python-pptx >= 1.0, lxml and Pillow (scripts/requirements.txt). PyYAML optional. --render needs
LibreOffice and poppler (scripts/render_office.sh). Run from anywhere; paths in a spec resolve against the
current folder, then the spec's folder, then the skill root (so "assets/icons/png/humanoid-navy.png" works).
"""
from __future__ import annotations

import sys

sys.dont_write_bytecode = True  # never leave __pycache__ in the skill (brand_check is imported from here)

import argparse  # noqa: E402
import copy  # noqa: E402
import datetime as _dt  # noqa: E402
import io  # noqa: E402
import json  # noqa: E402
import math  # noqa: E402
import os  # noqa: E402
import posixpath  # noqa: E402
import re  # noqa: E402
import subprocess  # noqa: E402
import textwrap  # noqa: E402
import zipfile  # noqa: E402
from dataclasses import dataclass, field  # noqa: E402
from pathlib import Path  # noqa: E402

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
TEMPLATES_DIR = SKILL / "assets" / "templates"

try:
    from lxml import etree
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData, XyChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
    from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.opc.packuri import PackURI
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Pt
except ImportError as exc:  # pragma: no cover
    print(f"new_deck.py needs python-pptx, lxml and Pillow ({exc}). Install them in a virtual environment:\n"
          "  python3 -m venv .venv && .venv/bin/pip install -r scripts/requirements.txt\n"
          "then run: .venv/bin/python scripts/new_deck.py spec.json", file=sys.stderr)
    sys.exit(1)
try:
    from PIL import Image
except ImportError:  # pragma: no cover
    Image = None

EMU_IN = 914400
FONT, MONO = "Source Sans 3", "Source Code Pro"
NAVY, GOLD, INK, SLATE, LINE = "003976", "EFAB00", "1B2533", "4A5868", "D6DEE8"
WHITE, MIST, GRID, AXIS = "FFFFFF", "F3F6FA", "E6EBF1", "B9C2CD"
# Categorical order (references/data-visualization.md 2.1): navy, gold, sky, brick, teal, orange, plum, green.
ATR_SERIES = ["003976", "EFAB00", "2C8ECD", "B63B35", "059583", "DC7533", "7D4DAD", "47A34E"]
# Milestone status as a typed glyph + label (shape + label + color). The glyph colors are the dark status
# text colors that pass 4.5:1 on white; the lighter marker fills are for drawn shapes only.
STATUS = {
    "complete": ("▲", "003976", "Complete"),
    "on-track": ("●", "137738", "On track"),
    "at-risk": ("◆", "915109", "At risk"),
    "late": ("■", "A21921", "Late"),
    "not-started": ("△", "616F7E", "Not started"),
}
TEMPLATE_ALIASES = {
    "presentation": "ATR-Presentation-Template.pptx", "deck": "ATR-Presentation-Template.pptx",
    "slides": "ATR-Presentation-Template.pptx",
    "quad": "ATR-Quad-Chart-Template.pptx", "quad-chart": "ATR-Quad-Chart-Template.pptx",
    "poster-48x36": "ATR-Research-Poster-48x36.pptx", "poster-landscape": "ATR-Research-Poster-48x36.pptx",
    "poster-36x48": "ATR-Research-Poster-36x48.pptx", "poster-portrait": "ATR-Research-Poster-36x48.pptx",
    "flyer": "ATR-Flyer-Letter.pptx", "one-pager": "ATR-One-Pager-Letter.pptx",
    "certificate": "ATR-Certificate-Letter.pptx", "door-sign": "ATR-Door-Sign-Letter.pptx",
    "badge": "ATR-Name-Badge.pptx", "name-badge": "ATR-Name-Badge.pptx",
    "social-square": "social/ATR-Social-Square-1080.pptx", "social-portrait": "social/ATR-Social-Portrait-1080x1350.pptx",
    "social-story": "social/ATR-Social-Story-1080x1920.pptx",
    "social-landscape": "social/ATR-Social-Landscape-1200x675.pptx",
    "youtube-thumbnail": "social/ATR-YouTube-Thumbnail-1280x720.pptx",
}
RESERVED = {"layout", "fields", "notes", "speaker_notes", "section", "from_slide", "remove", "keep", "shapes", "tables", "images",
            "textboxes", "comment", "_comment", "hidden"}
FIXED_SETS = [   # (layout name contains, field prefix, count, what, advice)
    ("Agenda", "item", 5, "agenda rows", "Plan the talk in five parts, or drop the Agenda slide and use "
                                          "Section Divider (Outline) slides."),
    ("Timeline", "step", 5, "timeline stations", "Use five steps, or show the plan as a Table."),
    ("Acknowledgements", "logo", 4, "sponsor-logo panels", "Put real sponsor logos in them, or give the "
                                                           "acknowledgement a Title + Content slide instead."),
]
PIC_ALIASES = ("image", "picture", "photo", "figure", "logo", "icon")
BODY_ALIASES = ("body", "text", "content", "bullets")
GENERIC_NAME = re.compile(r"^(?:text|shape|image|picture|content|content placeholder|text placeholder|title|"
                          r"subtitle|placeholder|body|object|chart|table|media|google shape;\S*)\s+\d+$|^$", re.I)
NUM_WORDS = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6, "seven": 7, "eight": 8,
             "nine": 9, "ten": 10, "twelve": 12}
_NUM = r"(\d+|one|two|three|four|five|six|seven|eight|nine|ten|twelve)"


class SpecError(Exception):
    """Bad spec or usage: exit 1."""


def _num(tok):
    return int(tok) if tok.isdigit() else NUM_WORDS.get(tok.lower())


def rgb(hexstr):
    return RGBColor.from_string(hexstr.lstrip("#").upper())


def emu(inches):
    return Emu(int(round(float(inches) * EMU_IN)))


def _lum(hexstr):
    h = hexstr.lstrip("#")
    out = []
    for i in (0, 2, 4):
        c = int(h[i:i + 2], 16) / 255
        out.append(c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4)
    return 0.2126 * out[0] + 0.7152 * out[1] + 0.0722 * out[2]


def contrast(a, b):
    la, lb = sorted((_lum(a), _lum(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


# --------------------------------------------------------------------------- brand rules (from brand_check)

def _load_rules():
    """Placeholder patterns and wording rules. Reuses brand_check.py (same folder) so both tools agree;
    falls back to a minimal copy if it cannot be imported."""
    try:
        sys.path.insert(0, str(HERE))
        import brand_check as bc  # noqa: WPS433
        ph = [(re.compile(rx, fl), label) for rx, fl, label in bc.PLACEHOLDERS]
        wording = [(cid, re.compile(rx, fl), sev, msg, fix) for cid, rx, fl, sev, msg, fix in bc.WORDING]
        return ph, wording
    except Exception:  # pragma: no cover - fallback
        ph = [(re.compile(r"\[(?!\d+(?:\s*[,–\-]\s*\d+)*\])[^\[\]\n]{1,80}\]"), "bracketed placeholder"),
              (re.compile(r"\blorem ipsum\b", re.I), "lorem ipsum filler"),
              (re.compile(r"\bclick to (?:add|edit)\b", re.I), "PowerPoint prompt text"),
              (re.compile(r"\b(?:TODO|TBD|FIXME)\b"), "TODO/TBD marker")]
        wording = [("handle", re.compile(r"(?<![\w@])@atr_kent(?![\w])", re.I), "error",
                    '"@atr_kent" does not exist.', "Use X @atrlab_kent."),
                   ("college", re.compile(r"College\s+of\s+Arts\s+(?:and|&)\s+Sciences", re.I), "error",
                    "Computer Science is in the College of Sciences and Humanities.", ""),
                   ("contact", re.compile(r"330[\s.)\-]*672[\s.\-]*9060"), "warning",
                    "330-672-9060 is not a verified lab number.", "Use 330-672-9980 (department)."),
                   ("contact", re.compile(r"\bRoom\s+236\b", re.I), "warning", "Room 236 is not verified.", ""),
                   ("building", re.compile(r"\bMath(?:ematics)?\s+(?:and|&)\s+Computer\s+Science\s+Building\b",
                                           re.I), "warning", "Older building name.",
                    'Write "Mathematical Sciences Building".')]
        return ph, wording


PH_RULES, WORDING_RULES = _load_rules()
# Wording rules that new_deck refuses outright (as social_card.py does): brand_check errors plus unverified facts.
REFUSE_IDS = {"contact", "building", "claim"}


# --------------------------------------------------------------------------- template catalog

@dataclass
class Slot:
    name: str
    idx: int
    kind: str                     # text | pic | chart | tbl | media
    ph_type: str
    box: tuple | None             # inches x, y, w, h
    meta: dict = field(default_factory=dict)

    def budget_str(self):
        m = self.meta
        bits = []
        if m.get("mc"):
            bits.append(f"{m['mc']} ch" + ("/para" if m.get("mp") else ""))
        if m.get("ml"):
            bits.append(f"{m['ml']} ln")
        if m.get("mp"):
            bits.append(f"{m['mp']} para")
        if m.get("fallback_pt"):
            bits.append(f"fallback {m['fallback_pt']} pt")
        return ", ".join(bits) or "-"


def resolve_template(name_or_path: str) -> Path:
    cand = Path(name_or_path).expanduser()
    if cand.suffix.lower() in (".pptx", ".potx"):
        for base in (Path.cwd(), SKILL, TEMPLATES_DIR):
            p = cand if cand.is_absolute() else base / cand
            if p.exists():
                return p.resolve()
        raise SpecError(f"template file not found: {name_or_path}")
    key = name_or_path.strip().lower()
    if key in TEMPLATE_ALIASES:
        return TEMPLATES_DIR / TEMPLATE_ALIASES[key]
    stems = {}
    for p in TEMPLATES_DIR.rglob("*.pptx"):
        s = p.stem.lower()
        stems[s] = p
        stems[re.sub(r"^atr-", "", s)] = p
    if key in stems:
        return stems[key]
    raise SpecError(f"unknown template '{name_or_path}'. Known: {', '.join(sorted(TEMPLATE_ALIASES))}, "
                    "or a path to a .pptx/.potx file.")


def open_presentation(path: Path):
    """python-pptx refuses a .potx by its content type; open a patched copy in memory."""
    if path.suffix.lower() != ".potx":
        return Presentation(str(path))
    src = zipfile.ZipFile(path)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.template.main+xml", b"presentationml.presentation.main+xml")
            out.writestr(item, data)
    buf.seek(0)
    return Presentation(buf)


def find_layouts_json(template: Path, override: str | None):
    if override:
        p = Path(override).expanduser()
        if not p.exists():
            raise SpecError(f"--layouts file not found: {override}")
        return json.loads(p.read_text(encoding="utf-8")), p
    for cand in (template.with_suffix(".layouts.json"), template.parent / (template.stem + "-layouts.json")):
        if cand.exists():
            return json.loads(cand.read_text(encoding="utf-8")), cand
    for p in sorted(TEMPLATES_DIR.glob("*.json")):
        try:
            d = json.loads(p.read_text(encoding="utf-8"))
        except (ValueError, OSError):
            continue
        if isinstance(d, dict) and "layouts" in d and \
                Path(posixpath.basename(str(d.get("template", "")))).stem == template.stem:   # .pptx or .potx
            return d, p
    return None, None


def _parse_cpl(text):
    m = (re.search(r"about (\d+) char\w* (?:per line|each|a line)", text) or
         re.search(r"\((?:about )?(\d+) characters? per line", text) or
         re.search(r"lines? of about (\d+) char", text))
    return int(m.group(1)) if m else None


def _parse_capacity(text):
    """Line capacity from 'N one-line bullets or M two-line bullets' style purposes."""
    one = re.search(_NUM + r" one-line", text, re.I)
    two = re.search(_NUM + r" two-line", text, re.I)
    a = _num(one.group(1)) if one else None
    b = _num(two.group(1)) if two else None
    return a, b


def _meta_from_json(p):
    purpose = p.get("purpose", "") or ""
    font = p.get("font") or {}
    m = {
        "mc": p.get("max_chars"), "ml": p.get("max_lines"), "mp": p.get("max_paragraphs"),
        "cpl": p.get("chars_per_line") or _parse_cpl(purpose),
        "fallback_pt": p.get("fallback_pt"), "size_pt": font.get("size_pt") or font.get("size"),
        "levels": p.get("levels"), "purpose": purpose, "prompt": p.get("prompt", ""),
        "format": p.get("format", ""), "overflow": p.get("overflow_behaviour", ""),
        "role": p.get("role", ""),
    }
    one, two = _parse_capacity(purpose)
    m["one_line"], m["two_line"] = one, two
    return m


def _meta_from_prompt(prompt):
    """Budgets written into a layout prompt ('up to 85 characters each (8 lines)'), for templates with no JSON."""
    m = {"prompt": prompt, "purpose": prompt}
    t = prompt or ""
    mc = re.search(r"(?<!\()up to (\d[\d,]*) char\w*", t, re.I)
    if mc:
        n = int(mc.group(1).replace(",", ""))
        if re.search(r"up to \d[\d,]* char\w*(?: \([^)]*\))? each", t, re.I):
            m["mc"] = n
            cnt = re.match(r"\s*" + _NUM + r"\s+\w+", t, re.I)
            if cnt and _num(cnt.group(1)):
                m["mp"] = _num(cnt.group(1))
        else:
            m["mc"] = n
    ml = re.search(r"up to (\d+) lines?|\((\d+) lines?", t, re.I)
    if ml and not m.get("mp"):
        m["ml"] = int(next(g for g in ml.groups() if g))
    return m


def _slug(s):
    s = re.sub(r"^(?:picture placeholder|placeholder|ph:)\s*", "", s.strip(), flags=re.I)
    return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")


def _kind(ph_type):
    return {PP_PLACEHOLDER.PICTURE: "pic", PP_PLACEHOLDER.BITMAP: "pic", PP_PLACEHOLDER.CHART: "chart",
            PP_PLACEHOLDER.TABLE: "tbl", PP_PLACEHOLDER.MEDIA_CLIP: "media"}.get(ph_type, "text")


def _box_of(shape):
    try:
        return (shape.left / EMU_IN, shape.top / EMU_IN, shape.width / EMU_IN, shape.height / EMU_IN)
    except (TypeError, AttributeError):
        return None


class Template:
    def __init__(self, name_or_path, layouts_json=None):
        self.path = resolve_template(name_or_path)
        self.prs = open_presentation(self.path)
        self.spec, self.spec_path = find_layouts_json(self.path, layouts_json)
        self.slide_w = self.prs.slide_width / EMU_IN
        self.slide_h = self.prs.slide_height / EMU_IN
        self._json_by_name = {}
        if self.spec:
            for L in self.spec.get("layouts", []):
                self._json_by_name[L.get("name")] = L
        self.layouts = {l.name: l for l in self.prs.slide_layouts}

    # -- layouts
    def find_layout(self, name):
        if isinstance(name, int) or (isinstance(name, str) and name.isdigit()):
            i = int(name)
            lays = list(self.prs.slide_layouts)
            if 0 <= i < len(lays):
                return lays[i]
            raise SpecError(f"layout index {i} out of range (0-{len(lays) - 1})")
        if name in self.layouts:
            return self.layouts[name]
        low = {k.lower(): v for k, v in self.layouts.items()}
        for cand in (name.lower(), ("atr - " + name).lower()):
            if cand in low:
                return low[cand]
        raise SpecError(f"no layout named '{name}' in {self.path.name}. Layouts: "
                        + "; ".join(self.layouts))

    def json_layout(self, layout):
        return self._json_by_name.get(layout.name)

    def slots(self, layout):
        """Every fillable placeholder on the layout, named from the JSON when there is one."""
        jl = self.json_layout(layout)
        jph = {}
        if jl:
            for p in jl.get("placeholders", []):
                jph[int(p["idx"])] = p
        out, used = [], set()
        for ph in layout.placeholders:
            pf = ph.placeholder_format
            if pf.type in (PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE,
                           PP_PLACEHOLDER.HEADER):
                continue
            j = jph.get(pf.idx)
            if j:
                name = j.get("name") or j.get("key") or f"ph{pf.idx}"
                t = (j.get("type") or j.get("ph_type") or "").lower()
                kind = {"pic": "pic", "chart": "chart", "tbl": "tbl", "table": "tbl",
                        "media": "media"}.get(t, _kind(pf.type))
                b = j.get("box_in") or j.get("box")
                box = (b["x"], b["y"], b["w"], b["h"]) if isinstance(b, dict) else (tuple(b) if b else _box_of(ph))
                meta = _meta_from_json(j)
            else:
                raw = ph.name or ""
                stripped = re.sub(r"^(?:picture placeholder|placeholder|ph:)\s*", "", raw, flags=re.I)
                if GENERIC_NAME.match(stripped) or not stripped:
                    name = "title" if pf.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE) \
                        and "title" not in used else f"ph{pf.idx}"
                else:
                    name = _slug(stripped)
                kind = _kind(pf.type)
                box = _box_of(ph)
                prompt = ph.text_frame.text if ph.has_text_frame else ""
                meta = _meta_from_prompt(prompt)
            if name in used:
                name = f"{name}_{pf.idx}"
            used.add(name)
            out.append(Slot(name=name, idx=pf.idx, kind=kind, ph_type=str(pf.type).split(" ")[0], box=box,
                            meta=meta))
        return out

    def content_box(self):
        if abs(self.slide_w - 10) < 0.01 and abs(self.slide_h - 5.625) < 0.01:
            return (0.5, 1.42, 9.0, 3.3)   # ATR 16:9 content zone (below the plate, above the footer)
        m = min(self.slide_w, self.slide_h) * 0.06
        return (m, self.slide_h * 0.22, self.slide_w - 2 * m, self.slide_h * 0.66)


def lookup(slots, key):
    """Field name -> Slot (exact name, ph<idx>/idx:<n>, then aliases)."""
    k = key.strip()
    kl = k.lower()
    for s in slots:
        if s.name.lower() == kl:
            return s
    m = re.fullmatch(r"(?:ph|idx:?)(\d+)", kl)
    if m:
        for s in slots:
            if s.idx == int(m.group(1)):
                return s
    if kl in PIC_ALIASES:
        pics = [s for s in slots if s.kind == "pic"]
        return pics[0] if pics else None
    if kl in ("video", "media", "movie"):
        media = [s for s in slots if s.kind == "media"] or [s for s in slots if s.kind == "pic"]
        return media[0] if media else None
    if kl == "chart":
        c = [s for s in slots if s.kind == "chart"] or [s for s in slots if s.kind == "pic"]
        return c[0] if c else None
    if kl == "table":
        c = [s for s in slots if s.kind == "tbl"] or [s for s in slots if s.kind == "pic"]
        return c[0] if c else None
    if kl in BODY_ALIASES:
        texts = [s for s in slots if s.kind == "text" and s.name not in ("title", "eyebrow", "subtitle",
                                                                         "caption", "source", "flag")]
        texts = [s for s in texts if s.box] or texts
        if texts:
            return max(texts, key=lambda s: (s.box[2] * s.box[3]) if s.box else 0)
    if kl in ("subtitle", "subhead"):
        for s in slots:
            if s.name in ("subtitle", "subhead"):
                return s
    if kl == "title":
        for s in slots:
            if s.ph_type in ("TITLE", "CENTER_TITLE"):
                return s
    if kl == "caption":
        for s in slots:
            if s.name in ("caption", "source", "q3cap", "caption1"):
                return s
    return None


# --------------------------------------------------------------------------- build context

@dataclass
class Ctx:
    tpl: Template
    spec_dir: Path
    allow_placeholders: bool = False
    verbose: bool = False
    warnings: list = field(default_factory=list)
    refusals: list = field(default_factory=list)
    infos: list = field(default_factory=list)
    agenda: list = field(default_factory=list)

    def warn(self, loc, msg):
        self.warnings.append(f"{loc}: {msg}")

    def refuse(self, loc, msg):
        self.refusals.append(f"{loc}: {msg}")


def resolve_path(p, spec_dir):
    if not str(p or "").strip():
        raise SpecError("empty file path")
    cand = Path(str(p)).expanduser()
    if cand.is_absolute():
        if cand.is_file():
            return cand
        raise SpecError(f"file not found: {p}")
    for base in (Path.cwd(), spec_dir, SKILL):
        q = base / cand
        if q.is_file():
            return q.resolve()
    raise SpecError(f"file not found: {p} (looked in the current folder, {spec_dir} and {SKILL})")


# --------------------------------------------------------------------------- text

def to_paragraphs(value):
    """Spec text -> [(level0, text)]: strings split on newlines (leading Tabs = deeper levels);
    lists nest one level per nested list; dicts {"text", "level"}."""
    out = []

    def add_str(s, lvl):
        for line in str(s).split("\n"):
            extra = len(line) - len(line.lstrip("\t"))
            out.append((lvl + extra, line.lstrip("\t")))

    def walk(items, lvl):
        for it in items:
            if isinstance(it, (list, tuple)):
                walk(it, lvl + 1)
            elif isinstance(it, dict):
                out.append((int(it.get("level", lvl + 1)) - 1 if "level" in it else lvl, str(it.get("text", ""))))
            elif it is not None:
                add_str(it, lvl)

    if isinstance(value, (list, tuple)):
        walk(value, 0)
    elif value is not None:
        add_str(value, 0)
    return out


def _runs(text):
    """Inline markup -> runs: '**bold**' and '^{sup}' (affiliation marks, units).
    'a **b** c^{1}' -> [('a ', False, False), ('b', True, False), (' c', False, False), ('1', False, True)]."""
    parts = re.split(r"(\*\*[^*]+\*\*|\^\{[^}]*\})", text)
    out = []
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            out.append((part[2:-2], True, False))
        elif part.startswith("^{") and part.endswith("}"):
            out.append((part[2:-1], False, True))
        else:
            out.append((part, False, False))
    return out or [("", False, False)]


def plain(text):
    return re.sub(r"\^\{([^}]*)\}", r"\1", re.sub(r"\*\*([^*]+)\*\*", r"\1", text))


def fill_text(shape, paras, size_pt=None):
    """Replace the text of a placeholder, keeping any formatting the shape already has per level
    (cloned showcase slides) and otherwise inheriting everything from the layout (run.text only)."""
    tf = shape.text_frame
    txBody = tf._txBody
    old_paras = []   # (level, pPr, rPr) of the text being replaced, in order
    for p in txBody.findall(qn("a:p")):
        pPr = p.find(qn("a:pPr"))
        lvl = int(pPr.get("lvl", "0")) if pPr is not None else 0
        r = p.find(qn("a:r"))
        rPr = r.find(qn("a:rPr")) if r is not None else None
        old_paras.append((lvl, copy.deepcopy(pPr) if pPr is not None else None,
                          copy.deepcopy(rPr) if rPr is not None else None))
        txBody.remove(p)

    def template_for(k, lvl):
        # same position and level first (a callout's label / numeral / text), then the last of that level
        if k < len(old_paras) and old_paras[k][0] == lvl and (old_paras[k][1] is not None or old_paras[k][2] is not None):
            return old_paras[k][1:]
        for ol, opPr, orPr in reversed(old_paras):
            if ol == lvl and (opPr is not None or orPr is not None):
                return opPr, orPr
        return None

    for k, (lvl, text) in enumerate(paras):
        p = etree.SubElement(txBody, qn("a:p"))
        tmpl = template_for(k, lvl)
        if tmpl and tmpl[0] is not None:
            pPr = copy.deepcopy(tmpl[0])
            p.append(pPr)
        else:
            pPr = None
        if lvl > 0:
            if pPr is None:
                pPr = etree.SubElement(p, qn("a:pPr"))
            pPr.set("lvl", str(lvl))
        elif pPr is not None and "lvl" in pPr.attrib:
            del pPr.attrib["lvl"]
        for chunk, bold, sup in _runs(text):
            r = etree.SubElement(p, qn("a:r"))
            if tmpl and tmpl[1] is not None:
                rPr = copy.deepcopy(tmpl[1])
                for a in ("dirty", "err", "baseline"):
                    rPr.attrib.pop(a, None)
                r.append(rPr)
            else:
                rPr = None
            if sup:
                if rPr is None:
                    rPr = etree.Element(qn("a:rPr"))
                    r.insert(0, rPr)
                rPr.set("baseline", "30000")
            if bold or size_pt:
                if rPr is None:
                    rPr = etree.SubElement(r, qn("a:rPr"))
                    r.remove(rPr)
                    r.insert(0, rPr)
                rPr.set("lang", rPr.get("lang", "en-US"))
                if bold:
                    rPr.set("b", "1")
                if size_pt:
                    rPr.set("sz", str(int(round(size_pt * 100))))
            t = etree.SubElement(r, qn("a:t"))
            t.text = chunk
    if not paras:
        etree.SubElement(txBody, qn("a:p"))


def check_budget(ctx, loc, slot, paras):
    """Warnings for text over the layout's measured budget; returns the fallback size to apply, if any."""
    m = slot.meta
    mc, ml, mp, cpl = m.get("mc"), m.get("ml"), m.get("mp"), m.get("cpl")
    texts = [plain(t) for _, t in paras]
    fallback = None
    if not texts:
        return None
    levels = m.get("levels")
    deepest = max(l for l, _ in paras) + 1
    if levels and deepest > levels:
        ctx.warn(loc, f"uses level {deepest}; this placeholder styles {levels} level(s) "
                      f"({m.get('purpose', '')[:80]})")
    if not mp and levels and levels > 1:
        # a label with a level-2 detail line (timeline stations): the budget is the label's
        texts = [plain(t) for l, t in paras if l == 0] or texts
    if mp:
        limit = max(mp, m.get("one_line") or 0)
        if len(texts) > limit:
            ctx.warn(loc, f"{len(texts)} paragraphs; the layout holds {limit} ({slot.budget_str()}). "
                          "Cut, or split the slide")
        per = mc
        two_line_ok = bool(m.get("two_line")) or "two-line" in m.get("purpose", "")
        if mc and cpl and two_line_ok:
            per = max(mc, 2 * cpl)
        long = [t for t in texts if per and len(t) > per]
        for t in long[:3]:
            ctx.warn(loc, f"paragraph of {len(t)} characters over the {per}-character budget: \"{t[:50]}...\"")
        if cpl and not CAN_MEASURE:
            est = sum(max(1, math.ceil(len(t) / cpl)) for t in texts)
            cap = max(m.get("one_line") or 0, 2 * (m.get("two_line") or 0)) or (mp * (ml or 1))
            if est > cap:
                ctx.warn(loc, f"about {est} lines of text; the box holds about {cap} "
                              f"({cpl} characters per line). Render and check, or cut")
    else:
        total = sum(len(t) for t in texts)
        if mc and total > mc:
            fb = m.get("fallback_pt")
            size = m.get("size_pt")
            cap_fb = math.ceil(mc * size / fb) if (fb and size) else None
            if fb and size and total <= cap_fb:
                fallback = fb
                ctx.warn(loc, f"{total} characters over the {mc}-character budget: set to the template's "
                              f"{fb} pt fallback (holds about {cap_fb}). A shorter title is better")
            elif fb and size:
                fallback = fb
                ctx.warn(loc, f"{total} characters: too long even at the {fb} pt fallback "
                              f"(about {cap_fb}). Shorten it"
                         + (" or use 'ATR - Title + Content (Two-line title)'" if slot.name == "title"
                            and ctx.tpl.slide_h < 6 else ""))
            else:
                ctx.warn(loc, f"{total} characters over the {mc}-character budget ({m.get('purpose', '')[:90]})")
        if ml and not CAN_MEASURE:
            if cpl:
                est = sum(max(1, math.ceil(len(t) / cpl)) for t in texts)
            else:
                est = len(texts)
            if est > ml and not fallback:
                ctx.warn(loc, f"about {est} lines; the box holds {ml}. Cut words")
    return fallback


# --------------------------------------------------------------------------- measurement (brand fonts)

FONT_DIR = SKILL / "assets" / "fonts"
CAN_MEASURE = Image is not None and (FONT_DIR / "source-sans-3" / "SourceSans3-Regular.ttf").exists()
FONT_FILES = {   # (family, bold) -> file; weight-named families ignore bold (they have no bold style)
    ("source sans 3", False): "source-sans-3/SourceSans3-Regular.ttf",
    ("source sans 3", True): "source-sans-3/SourceSans3-Bold.ttf",
    ("source sans 3 black", None): "source-sans-3/SourceSans3-Black.ttf",
    ("source sans 3 semibold", None): "source-sans-3/SourceSans3-Semibold.ttf",
    ("source sans 3 light", None): "source-sans-3/SourceSans3-Light.ttf",
    ("roboto slab", False): "roboto-slab/RobotoSlab-Regular.ttf",
    ("roboto slab", True): "roboto-slab/RobotoSlab-Bold.ttf",
    ("roboto slab semibold", None): "roboto-slab/RobotoSlab-SemiBold.ttf",
    ("roboto slab light", None): "roboto-slab/RobotoSlab-Light.ttf",
    ("source code pro", False): "source-code-pro/SourceCodePro-Regular.ttf",
    ("source code pro", True): "source-code-pro/SourceCodePro-Bold.ttf",
    ("source code pro semibold", None): "source-code-pro/SourceCodePro-Semibold.ttf",
}
_FONT_CACHE = {}


def _ttf(family, bold, size_pt):
    """(ImageFont at 10x the point size, width factor). Arial and unknown faces are approximated with
    Source Sans 3 widened by the measured Arial difference (references/typography.md 8)."""
    if Image is None:
        return None, 1.0
    from PIL import ImageFont
    fam = (family or FONT).strip().lower()
    factor = 1.0
    if fam.startswith("+m") or fam in ("", "inherit"):
        fam = "source sans 3"
    key = (fam, None) if (fam, None) in FONT_FILES else (fam, bool(bold))
    if key not in FONT_FILES:
        factor = 1.09 if bold else 1.08   # Arial / Helvetica / anything else: Source Sans 3 plus Arial's extra width
        key = ("source sans 3", bool(bold))
    path = FONT_DIR / FONT_FILES[key]
    ck = (str(path), round(size_pt * 10))
    if ck not in _FONT_CACHE:
        if not path.exists():
            return None, 1.0
        _FONT_CACHE[ck] = ImageFont.truetype(str(path), max(1, round(size_pt * 10)))
    return _FONT_CACHE[ck], factor


def _level_style(ph_el, level):
    """Resolve size, bold, face, caps, spacing, indent and line spacing for one level of a placeholder from
    its own lstStyle, falling back to ATR body defaults."""
    st = {"sz": 18.0, "b": False, "face": FONT, "cap": False, "spc": 0.0, "marL": 0, "lnPct": 1.0, "lnPts": None,
          "bef": 0.0, "aft": 0.0}
    if ph_el is None:
        return st
    lst = ph_el.find(".//" + qn("a:lstStyle"))
    pPr = lst.find(qn(f"a:lvl{level + 1}pPr")) if lst is not None else None
    if pPr is None and lst is not None:
        pPr = lst.find(qn("a:lvl1pPr"))
    if pPr is None:
        return st
    st["marL"] = int(pPr.get("marL", "0")) / EMU_IN
    ln = pPr.find(qn("a:lnSpc"))
    if ln is not None:
        pct, pts = ln.find(qn("a:spcPct")), ln.find(qn("a:spcPts"))
        if pct is not None:
            st["lnPct"] = int(pct.get("val")) / 100000
        if pts is not None:
            st["lnPts"] = int(pts.get("val")) / 100
    for tag, k in (("a:spcBef", "bef"), ("a:spcAft", "aft")):
        e = pPr.find(qn(tag))
        if e is not None and e.find(qn("a:spcPts")) is not None:
            st[k] = int(e.find(qn("a:spcPts")).get("val")) / 100
    d = pPr.find(qn("a:defRPr"))
    if d is not None:
        if d.get("sz"):
            st["sz"] = int(d.get("sz")) / 100
        st["b"] = d.get("b") == "1"
        st["cap"] = d.get("cap") == "all"
        st["spc"] = int(d.get("spc", "0")) / 100
        lat = d.find(qn("a:latin"))
        if lat is not None and lat.get("typeface"):
            st["face"] = lat.get("typeface")
    return st


def _wrap_count(text, font, factor, spc_pt, size_pt, width_pt):
    """Lines a paragraph needs at this width (greedy word wrap, like PowerPoint)."""
    if not text.strip():
        return 1
    words = re.split(r"(\s+)", text)
    lines, cur = 1, ""

    def w(s):
        return font.getlength(s) / 10 * factor + spc_pt * len(s)

    for tok in words:
        trial = cur + tok
        if w(trial.rstrip()) <= width_pt or not cur.strip():
            cur = trial
            while w(cur.rstrip()) > width_pt and len(cur) > 1:   # a single word wider than the box
                lines += 1
                cur = cur[int(len(cur) * width_pt / max(w(cur), 1)):]
        else:
            lines += 1
            cur = tok.lstrip()
    return lines


def measure(shape, paras, size_override=None):
    """(lines per paragraph, text height in inches, box text height in inches) for text typed into a
    placeholder, using the layout's own level styles and the brand TTFs. None when fonts are missing."""
    if Image is None or not FONT_DIR.exists():
        return None
    try:
        layout_ph = shape._base_placeholder._element if getattr(shape, "_base_placeholder", None) is not None \
            else None
    except Exception:
        layout_ph = None
    src = layout_ph if layout_ph is not None else shape._element
    bp = src.find(".//" + qn("a:bodyPr"))
    ins = {k: int(bp.get(k)) / EMU_IN if bp is not None and bp.get(k) is not None else d
           for k, d in (("lIns", 0.1), ("rIns", 0.1), ("tIns", 0.05), ("bIns", 0.05))}
    try:
        box_w, box_h = shape.width / EMU_IN, shape.height / EMU_IN
    except (TypeError, AttributeError):
        return None
    counts, height = [], 0.0
    for k, (lvl, text) in enumerate(paras):
        st = _level_style(src, lvl)
        size = size_override or st["sz"]
        runs = _runs(text)
        bold = st["b"] or all(b for t_, b, _s in runs if t_.strip())
        font, factor = _ttf(st["face"], bold, size)
        if font is None:
            return None
        t = plain(text).upper() if st["cap"] else plain(text)
        width_pt = (box_w - ins["lIns"] - ins["rIns"] - st["marL"]) * 72
        n = _wrap_count(t.replace("\t", "    "), font, factor, st["spc"], size, width_pt)
        counts.append(n)
        asc, desc = font.getmetrics()
        natural = (asc + desc) / 10              # the face's natural line height at this size, in pt
        line = st["lnPts"] if st["lnPts"] else natural * st["lnPct"]
        height += n * line + (st["bef"] if k else 0) + (st["aft"] if k < len(paras) - 1 else 0)
    return counts, height / 72, box_h - ins["tIns"] - ins["bIns"]


def check_measured(ctx, loc, slot, shape, paras, size_override):
    """Overflow warnings from the brand-font measurement; returns (lines per paragraph, text h, box h) or None."""
    res = measure(shape, paras, size_override)
    if res is None:
        return None
    counts, h, box_h = res
    ml = slot.meta.get("ml")
    mp = slot.meta.get("mp")
    lv = slot.meta.get("levels")
    label_lines = sum(c for c, (l, _) in zip(counts, paras) if l == 0) if (lv and lv > 1) else sum(counts)
    if ml and not mp and label_lines > ml:
        ctx.warn(loc, f"measured {label_lines} lines with the brand fonts; this box is set for {ml}. "
                      "Cut words (it would wrap into the next element or shrink below the design size)")
        return res
    if h > box_h * 1.02 + 0.02:
        ctx.warn(loc, f"measured {h:.2f} in of text in a {box_h:.2f} in box ({sum(counts)} lines): it overflows. "
                      "Cut words or split the slide")
    return res


# --------------------------------------------------------------------------- fill the frame

# The main text slots of the content layouts. A slide whose body uses under a third of its line budget leaves
# the lower half of the slide empty (references/presentations.md §4.7): an info hint names better layouts.
UNDERFILL_SLOTS = ("body", "left_body", "right_body")
UNDERFILL_SHARE = 1 / 3
NO_FILL_HINT = re.compile(r"^ATR - (?:Title(?: \(Light\))?|Section Divider.*|Statement|Agenda|Thank You|"
                          r"Blank Branded|Content \+ Image|Full-Bleed Image|Video)$")   # a picture carries these


def line_budget(meta):
    """Lines a text slot is built to hold, from its layouts-JSON budget: max_paragraphs x max_lines, else the
    one-line-bullet capacity ("up to 7 one-line bullets"), else max_lines."""
    mp, ml, one = meta.get("mp"), meta.get("ml"), meta.get("one_line")
    if mp and ml:
        return mp * ml
    if mp:
        return max(mp, one or 0)
    return ml or None


def used_lines(meta, paras, res):
    """Lines the text takes: measured with the brand fonts when possible, else estimated from chars per line."""
    if res:
        return sum(res[0])
    cpl = meta.get("cpl") or meta.get("mc")
    return sum(max(1, math.ceil(len(plain(t)) / cpl)) if cpl else 1 for _, t in paras)


def underfill_hint(ctx, loc0, layout, use, visual):
    """Info when the body slots of a content slide hold under a third of their line budget."""
    if not use or NO_FILL_HINT.match(layout.name) or "ATR - Statement" not in ctx.tpl.layouts:
        return
    used, budget = sum(u for _, u, _ in use), sum(b for _, _, b in use)
    if used >= budget * UNDERFILL_SHARE:
        return
    what = "the body holds" if len(use) == 1 else "the columns hold"
    alt = [("ATR - Statement", "one sentence"), ("ATR - Key Numbers", "up to three numbers"),
           ("ATR - Content + Image", "a figure or photo with a short explanation"),
           ("ATR - Three Icon Columns", "three short points")]
    alt = [a for a in alt if a[0] != layout.name and a[0] in ctx.tpl.layouts]
    tip = "; ".join(f"{n} for {w}" for n, w in alt)
    extra = "" if visual else ", or add a figure, photo or icon to the slide"
    ctx.infos.append(f"{loc0}: {what} {used} of about {budget} lines, so the lower part of the slide stays empty. "
                     f"Fill the frame with a layout sized for the content: {tip}{extra} "
                     "(references/presentations.md §4.7).")


def table_height(spec, widths, size, row_h):
    """Estimated table height in inches (rows grow when a cell wraps)."""
    font, factor = _ttf(FONT, False, size)
    if font is None:
        return None
    asc, desc = font.getmetrics()
    line = (asc + desc) / 10 / 72 * 1.0
    rows = ([spec.get("header")] if spec.get("header") else []) + list(spec.get("rows") or [])
    total = 0.0
    for r in rows:
        lines = 1
        for c, v in enumerate(r):
            txt = (STATUS.get(str(v.get("status")), ("", "", ""))[2] if isinstance(v, dict) else str(v or ""))
            wpt = (widths[c] - 0.222) * 72 if c < len(widths) else 100
            lines = max(lines, sum(_wrap_count(x, font, factor, 0, size, wpt) for x in txt.split("\n")))
        total += max(row_h, lines * line + 0.111)
    return total


# --------------------------------------------------------------------------- pictures

def image_info(path):
    """(width_px, height_px, has_transparency)."""
    if Image is None:
        return None, None, False
    with Image.open(path) as im:
        w, h = im.size
        alpha = False
        if im.mode in ("RGBA", "LA", "PA") or (im.mode == "P" and "transparency" in im.info):
            try:
                a = im.convert("RGBA").getchannel("A")
                alpha = a.getextrema()[0] < 250
            except Exception:  # pragma: no cover
                alpha = True
        return w, h, alpha


def set_alt(shape_el, alt, decorative=False):
    cNvPr = shape_el.find(".//" + qn("p:cNvPr"))
    if cNvPr is None:
        return
    cNvPr.set("descr", alt or "")
    if decorative:
        ext_uri = "{C183D7F6-B498-43B3-948B-1728B52AA6E4}"
        lst = cNvPr.find(qn("a:extLst"))
        if lst is None:
            lst = etree.SubElement(cNvPr, qn("a:extLst"))
        if not any(e.get("uri") == ext_uri for e in lst):
            ext = etree.SubElement(lst, qn("a:ext"))
            ext.set("uri", ext_uri)
            dec = etree.SubElement(ext, "{http://schemas.microsoft.com/office/drawing/2017/decorative}decorative")
            dec.set("val", "1")


def _crop_cover(img_w, img_h, box_w, box_h):
    ia, ba = img_w / img_h, box_w / box_h
    if ia > ba:
        c = (1 - ba / ia) / 2
        return c, 0.0, c, 0.0
    c = (1 - ia / ba) / 2
    return 0.0, c, 0.0, c


def _contain(img_w, img_h, box, pad=0.0):
    x, y, w, h = box
    w2, h2 = w - 2 * pad, h - 2 * pad
    s = min(w2 / img_w, h2 / img_h)
    pw, ph = img_w * s, img_h * s
    return (x + (w - pw) / 2, y + (h - ph) / 2, pw, ph)


def add_panel(slide, box, before_el, color=MIST):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(box[0]), emu(box[1]), emu(box[2]), emu(box[3]))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(color)
    shp.line.fill.background()
    shp.shadow.inherit = False
    shp.name = "Figure panel"
    el = shp._element
    el.getparent().remove(el)
    before_el.addprevious(el)
    return shp


def place_picture(ctx, slide, loc, spec, slot_shape=None, box=None, slot=None):
    """Insert a picture into a picture placeholder (crop to fill by default) or into a frame.
    spec: str path or {"path", "alt", "fit": cover|contain, "panel": mist|none, "decorative"}."""
    if isinstance(spec, str):
        spec = {"path": spec}
    path = resolve_path(spec.get("path") or spec.get("image") or "", ctx.spec_dir)
    alt = spec.get("alt")
    decorative = bool(spec.get("decorative")) or alt == ""
    if alt is None and not decorative:
        field_name = loc.rsplit("\u203a", 1)[-1].strip() if "\u203a" in loc else (slot.name if slot else "image")
        ctx.refuse(loc, f"picture {path.name} has no alt text: add \"{field_name}_alt\" (or \"alt\" in the "
                        "picture object; \"\" marks it decorative)")
    w_px, h_px, has_alpha = image_info(path)
    name = (slot.name if slot else "").lower()
    fit = spec.get("fit")
    if not fit:
        fit = "contain" if (has_alpha or name.startswith(("logo", "icon"))) else "cover"
    purpose = (slot.meta.get("purpose", "") if slot else "").lower()
    panel = spec.get("panel")
    if panel is None:
        panel = "mist" if (fit == "contain" and "mist panel first" in purpose) else "none"
    frame = None
    if slot_shape is not None and hasattr(slot_shape, "insert_picture"):
        frame = _box_of(slot_shape) or box
        pic = slot_shape.insert_picture(str(path))
    else:
        frame = box or (_box_of(slot_shape) if slot_shape is not None else None) or ctx.tpl.content_box()
        pic = slide.shapes.add_picture(str(path), emu(frame[0]), emu(frame[1]), emu(frame[2]), emu(frame[3]))
        if slot_shape is not None:   # a non-picture placeholder: the picture takes its place in the z-order
            el = pic._element
            el.getparent().remove(el)
            slot_shape._element.addprevious(el)
            slot_shape._element.getparent().remove(slot_shape._element)
    if w_px and frame:
        if fit == "contain":
            pad = 0.15 if panel == "mist" else 0.0
            x, y, w, h = _contain(w_px, h_px, frame, pad)
            pic.crop_left = pic.crop_right = pic.crop_top = pic.crop_bottom = 0.0
            pic.left, pic.top, pic.width, pic.height = emu(x), emu(y), emu(w), emu(h)
            vis_w, placed_w = w_px, w
        else:
            cl, ct, cr, cb = _crop_cover(w_px, h_px, frame[2], frame[3])
            pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = cl, ct, cr, cb
            lost = max(cl + cr, ct + cb)
            if lost > 0.15:
                ctx.warn(loc, f"{path.name} is cropped by {lost:.0%} to fill the {frame[2]:.2f} x {frame[3]:.2f} in "
                              f"frame (aspect {frame[2] / frame[3]:.2f}). Fine for a photo; for a plot or diagram "
                              "export it at that aspect or pass \"fit\": \"contain\"")
            if slot_shape is None or not hasattr(slot_shape, "insert_picture"):
                pic.left, pic.top, pic.width, pic.height = emu(frame[0]), emu(frame[1]), emu(frame[2]), emu(frame[3])
            vis_w, placed_w = w_px * (1 - cl - cr), frame[2]
        ppi = vis_w / placed_w if placed_w else 999
        if ppi < 100:
            ctx.warn(loc, f"{path.name} is {ppi:.0f} ppi at this size (under 100): it will look soft; "
                          "use a larger original")
        if panel == "mist":
            add_panel(slide, frame, pic._element)
    set_alt(pic._element, alt or "", decorative)
    return pic


def place_media(ctx, slide, loc, spec, slot_shape, slot):
    if isinstance(spec, str):
        spec = {"poster": spec}
    box = _box_of(slot_shape) if slot_shape is not None else ctx.tpl.content_box()
    alt = spec.get("alt")
    if alt is None:
        ctx.refuse(loc, "video/poster has no alt text: add \"alt\"")
    if spec.get("video"):
        video = resolve_path(spec["video"], ctx.spec_dir)
        poster = resolve_path(spec["poster"], ctx.spec_dir) if spec.get("poster") else None
        mime = {".mp4": "video/mp4", ".m4v": "video/mp4", ".mov": "video/quicktime", ".wmv": "video/x-ms-wmv",
                ".avi": "video/x-msvideo"}.get(video.suffix.lower(), "video/unknown")
        mv = slide.shapes.add_movie(str(video), emu(box[0]), emu(box[1]), emu(box[2]), emu(box[3]),
                                    poster_frame_image=str(poster) if poster else None, mime_type=mime)
        el = mv._element
        if slot_shape is not None:
            el.getparent().remove(el)
            slot_shape._element.addprevious(el)
            slot_shape._element.getparent().remove(slot_shape._element)
        set_alt(el, alt or "")
        if not poster:
            ctx.warn(loc, "no poster frame: PowerPoint shows a generic icon until the video plays")
        return mv
    if spec.get("poster"):
        return place_picture(ctx, slide, loc, {"path": spec["poster"], "alt": alt, "fit": "cover"},
                             slot_shape=slot_shape, box=box, slot=slot)
    raise SpecError(f"{loc}: media needs \"video\" and/or \"poster\"")


# --------------------------------------------------------------------------- charts

CHART_TYPES = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED, "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED, "stacked_column": XL_CHART_TYPE.COLUMN_STACKED,
    "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "scatter": XL_CHART_TYPE.XY_SCATTER,
}
C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"


def _c(tag):
    return "{%s}%s" % (C_NS, tag)


def _nofill_sppr(parent_tag="c:spPr"):
    sp = etree.Element(qn(parent_tag))
    etree.SubElement(sp, qn("a:noFill"))
    ln = etree.SubElement(sp, qn("a:ln"))
    etree.SubElement(ln, qn("a:noFill"))
    return sp


def _font(font, size=None, color=None, bold=None, name=FONT):
    font.name = name
    if size:
        font.size = Pt(size)
    if color:
        font.color.rgb = rgb(color)
    if bold is not None:
        font.bold = bold


def build_chart_data(spec, kind):
    series = spec.get("series") or []
    if not series:
        raise SpecError("chart needs \"series\": [{\"name\", \"values\"}]")
    if kind == "scatter":
        data = XyChartData()
        for s in series:
            ser = data.add_series(s.get("name", ""))
            for x, y in s.get("values", []):
                ser.add_data_point(x, y)
        return data
    cats = spec.get("categories") or []
    data = CategoryChartData(number_format=spec.get("number_format", "General"))
    data.categories = cats
    for s in series:
        vals = s.get("values", [])
        if cats and len(vals) != len(cats):
            raise SpecError(f"chart series '{s.get('name')}' has {len(vals)} values for {len(cats)} categories")
        data.add_series(s.get("name", ""), vals)
    return data


def style_chart(ctx, loc, chart, spec, kind):
    series_spec = spec.get("series") or []
    n = len(series_spec)
    colors = [c.lstrip("#").upper() for c in (spec.get("colors") or ATR_SERIES)]
    if n > 8 and kind not in ("pie", "doughnut"):
        ctx.warn(loc, f"{n} series: the ATR palette has 8 slots; fold the rest into 'Other' or use small multiples")
    if kind == "line" and n > 4:
        ctx.warn(loc, f"{n} lines: over 4 lines need markers, dashes and end labels, or small multiples "
                      "(references/data-visualization.md 2.1)")
    if kind in ("pie", "doughnut") and len(spec.get("categories") or []) > 6:
        ctx.warn(loc, "pie with more than 6 slices: use a bar chart")
    cs = chart._chartSpace
    rc = cs.find(_c("roundedCorners"))
    if rc is None:
        rc = etree.Element(_c("roundedCorners"))
        anchor = cs.find(_c("chart"))
        for tag in ("style", "clrMapOvr", "pivotSource", "protection"):
            e = cs.find(_c(tag))
            if e is not None:
                anchor = e
                break
        anchor.addprevious(rc)
    rc.set("val", "0")
    chart.font.name = FONT
    chart.font.size = Pt(14)
    chart.font.color.rgb = rgb(SLATE)
    chart.has_title = False
    plot = chart.plots[0]
    single = n == 1
    stacked = kind.startswith("stacked")
    categories = spec.get("categories") or []
    hl = spec.get("highlight")
    hl_idx = None
    if hl is not None:
        if isinstance(hl, int):
            hl_idx = hl
        elif hl in categories:
            hl_idx = categories.index(hl)
        else:
            ctx.warn(loc, f"highlight '{hl}' is not a category")
        if not single:
            ctx.warn(loc, "highlight works on single-series charts only; ignored")
            hl_idx = None
    # series colors in fixed ATR order
    if kind in ("bar", "column", "stacked_bar", "stacked_column"):
        plot.gap_width = int(spec.get("gap_width", 55 if single else 80))
        plot.overlap = 100 if stacked else (0 if single else -8)
        for i, s in enumerate(plot.series):
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = rgb(colors[0] if (single and hl_idx is not None) else colors[i % len(colors)])
            s.invert_if_negative = False
        if single and hl_idx is not None:
            pt = plot.series[0].points[hl_idx]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = rgb(GOLD)
    elif kind == "line":
        for i, s in enumerate(plot.series):
            col = colors[i % len(colors)]
            s.smooth = False
            s.format.line.color.rgb = rgb(col)
            s.format.line.width = Pt(2.5)    # >= 2.25 pt: gold is never a hairline
            s.marker.style = XL_MARKER_STYLE.CIRCLE
            s.marker.size = 7
            s.marker.format.fill.solid()
            s.marker.format.fill.fore_color.rgb = rgb(col)
            s.marker.format.line.color.rgb = rgb(WHITE)
    elif kind in ("pie", "doughnut"):
        s = plot.series[0]
        for j in range(len(categories)):
            pt = s.points[j]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = rgb(colors[j % len(colors)])
            pt.format.line.color.rgb = rgb(WHITE)
            pt.format.line.width = Pt(1.5)
    elif kind == "scatter":
        for i, s in enumerate(plot.series):
            col = colors[i % len(colors)]
            s.format.line.fill.background()
            s.marker.style = XL_MARKER_STYLE.CIRCLE
            s.marker.size = 8
            s.marker.format.fill.solid()
            s.marker.format.fill.fore_color.rgb = rgb(col)
            s.marker.format.line.color.rgb = rgb(WHITE)
    # legend
    legend = spec.get("legend", "auto")
    show_legend = (n > 1) if legend == "auto" else bool(legend)
    if kind in ("pie", "doughnut") and legend == "auto":
        show_legend = kind == "doughnut"
    chart.has_legend = show_legend
    if show_legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        _font(chart.legend.font, 16, INK)
    # data labels
    labels = spec.get("labels", "auto")
    if labels == "auto":
        labels = ("all" if (single and kind in ("bar", "column")) or kind in ("pie",) else
                  "last" if kind == "line" else "none")
    fmt = spec.get("label_format") or spec.get("number_format")
    if labels == "all":
        plot.has_data_labels = True
        dl = plot.data_labels
        _font(dl.font, 14, INK)
        if fmt:
            dl.number_format = fmt
            dl.number_format_is_linked = False
        if kind in ("bar", "column", "line", "scatter"):
            dl.position = XL_LABEL_POSITION.OUTSIDE_END if kind in ("bar", "column") else XL_LABEL_POSITION.ABOVE
        elif stacked:
            dl.position = XL_LABEL_POSITION.CENTER
            for i, s in enumerate(plot.series):   # inside labels: pick white or ink per fill, never gold on white
                fill = colors[i % len(colors)]
                best = WHITE if contrast(WHITE, fill) >= contrast(INK, fill) else INK
                if contrast(best, fill) < 4.5:
                    ctx.warn(loc, f"labels inside {fill} bars reach only {contrast(best, fill):.1f}:1; "
                                  "label ends or use a legend instead")
                s_dl = s.data_labels
                _font(s_dl.font, 14, best)
                s_dl.position = XL_LABEL_POSITION.CENTER
                s_dl.show_value = True
        elif kind == "pie":
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
            dl.show_category_name = True
            dl.show_percentage = bool(spec.get("percent", True))
            dl.show_value = not bool(spec.get("percent", True))
    elif labels == "last" and kind in ("line", "scatter"):
        for s in plot.series:
            npts = len(list(s.values)) if hasattr(s, "values") else 0
            if not npts:
                continue
            dlbl = s.points[npts - 1].data_label._get_or_add_dLbl()
            txPr = dlbl.find(_c("txPr"))
            defRPr = txPr.find(".//" + qn("a:defRPr"))
            defRPr.set("sz", "1400")
            sf = etree.SubElement(defRPr, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr")).set("val", INK)
            etree.SubElement(defRPr, qn("a:latin")).set("typeface", FONT)
            if fmt:
                nf = etree.Element(_c("numFmt"))
                nf.set("formatCode", fmt)
                nf.set("sourceLinked", "0")
                dlbl.find(_c("idx")).addnext(nf)
            pos = etree.Element(_c("dLblPos"))
            pos.set("val", "r")
            dlbl.find(_c("showLegendKey")).addprevious(pos)
    # axes
    if kind not in ("pie", "doughnut"):
        ca, va = chart.category_axis, chart.value_axis
        caspec, vaspec = spec.get("category_axis") or {}, spec.get("value_axis") or {}
        hide_values = single and labels == "all" and kind in ("bar", "column") and not vaspec.get("visible")
        _font(ca.tick_labels.font, 16 if (single and kind in ("bar", "column")) else 14,
              INK if (single and kind in ("bar", "column")) else SLATE)
        ca.has_major_gridlines = False
        if single and kind == "bar":
            ca.format.line.fill.background()
        else:
            ca.format.line.color.rgb = rgb(AXIS)
            ca.format.line.width = Pt(0.75)
        if kind in ("bar", "stacked_bar"):
            ca.reverse_order = True            # first category on top
            vax = va._element
            cr = vax.find(_c("crosses"))       # keep the value axis at the bottom
            if cr is not None:
                cr.set("val", "max")
        if hide_values:
            va.visible = False
            va.has_major_gridlines = False
        else:
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = rgb(GRID)
            va.major_gridlines.format.line.width = Pt(0.75)
            va.format.line.fill.background()
            _font(va.tick_labels.font, 14, SLATE)
            if fmt and kind != "scatter":
                va.tick_labels.number_format = fmt
                va.tick_labels.number_format_is_linked = False
        vals = [v for s in series_spec for v in (s.get("values") or []) if isinstance(v, (int, float))]
        if "min" in vaspec:
            va.minimum_scale = vaspec["min"]
        elif kind in ("bar", "column", "stacked_bar", "stacked_column") and vals and min(vals) >= 0:
            va.minimum_scale = 0
        if "max" in vaspec:
            va.maximum_scale = vaspec["max"]
        if "major_unit" in vaspec:
            va.major_unit = vaspec["major_unit"]
        for ax, aspec in ((va, vaspec), (ca, caspec)):
            if aspec.get("title"):
                ax.has_title = True
                ax.axis_title.text_frame.text = aspec["title"]
                _font(ax.axis_title.text_frame.paragraphs[0].runs[0].font, 16, INK, bold=False)
    # quiet frame: no border or fill on the chart and plot areas
    ch = cs.find(_c("chart"))
    old = cs.find(_c("spPr"))
    if old is not None:
        cs.remove(old)
    ch.addnext(_nofill_sppr())
    pa = ch.find(_c("plotArea"))
    old = pa.find(_c("spPr"))
    if old is not None:
        pa.remove(old)
    ext = pa.find(_c("extLst"))
    sp = _nofill_sppr()
    if ext is not None:
        ext.addprevious(sp)
    else:
        pa.append(sp)


def place_chart(ctx, slide, loc, spec, slot_shape=None, box=None):
    if not isinstance(spec, dict):
        raise SpecError(f"{loc}: chart must be an object with type, categories and series")
    kind = str(spec.get("type", "bar")).lower().replace("-", "_").replace(" ", "_")
    if kind not in CHART_TYPES:
        raise SpecError(f"{loc}: chart type '{kind}' unknown; use {', '.join(CHART_TYPES)}")
    data = build_chart_data(spec, kind)
    ct = CHART_TYPES[kind]
    if slot_shape is not None and hasattr(slot_shape, "insert_chart"):
        gf = slot_shape.insert_chart(ct, data)
    else:
        frame = box or (_box_of(slot_shape) if slot_shape is not None else None) or ctx.tpl.content_box()
        gf = slide.shapes.add_chart(ct, emu(frame[0]), emu(frame[1]), emu(frame[2]), emu(frame[3]), data)
        if slot_shape is not None:
            el = gf._element
            el.getparent().remove(el)
            slot_shape._element.addprevious(el)
            slot_shape._element.getparent().remove(slot_shape._element)
    style_chart(ctx, loc, gf.chart, spec, kind)
    alt = spec.get("alt")
    if not alt:
        ctx.refuse(loc, "chart has no \"alt\": write what it shows and its takeaway, with the key numbers")
    set_alt(gf._element, alt or "")
    return gf


# --------------------------------------------------------------------------- tables

def _cell_text(cell, value, *, size, color, bold=False, font=FONT):
    tf = cell.text_frame
    txBody = tf._txBody
    for p in txBody.findall(qn("a:p")):
        txBody.remove(p)
    runs = []
    if isinstance(value, dict) and "status" in value:
        key = str(value["status"]).lower().replace(" ", "-").replace("_", "-")
        if key not in STATUS:
            raise SpecError(f"unknown status '{value['status']}'; use {', '.join(STATUS)}")
        glyph, gcol, label = STATUS[key]
        runs = [(glyph, gcol, False, FONT, size + 2), (" " + str(value.get("label", label)), color, bold, font, size)]
    elif isinstance(value, str) and value.lower().startswith("status:"):
        return _cell_text(cell, {"status": value.split(":", 1)[1].strip()}, size=size, color=color, bold=bold,
                          font=font)
    else:
        runs = None
    lines = [None] if runs else str("" if value is None else value).split("\n")
    for line in lines:
        p = etree.SubElement(txBody, qn("a:p"))
        pPr = etree.SubElement(p, qn("a:pPr"))
        pPr.set("marL", "0")
        pPr.set("indent", "0")
        etree.SubElement(pPr, qn("a:buNone"))
        chunks = runs if runs else [(t, color, bold or b, font, size) for t, b, _s in _runs(line)]
        for text, col, bb, fnt, sz in chunks:
            r = etree.SubElement(p, qn("a:r"))
            rPr = etree.SubElement(r, qn("a:rPr"))
            rPr.set("lang", "en-US")
            rPr.set("sz", str(int(sz * 100)))
            rPr.set("b", "1" if bb else "0")
            sf = etree.SubElement(rPr, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr")).set("val", col)
            etree.SubElement(rPr, qn("a:latin")).set("typeface", fnt)
            etree.SubElement(r, qn("a:t")).text = text


def _cell_frame(cell, fill=None):
    tcPr = cell._tc.get_or_add_tcPr()
    for ch in list(tcPr):
        tcPr.remove(ch)
    tcPr.set("marL", "101600")
    tcPr.set("marR", "101600")
    tcPr.set("marT", "50800")
    tcPr.set("marB", "50800")
    tcPr.set("anchor", "ctr")
    for tag in ("a:lnL", "a:lnR", "a:lnT"):
        ln = etree.SubElement(tcPr, qn(tag))
        ln.set("w", "0")
        etree.SubElement(ln, qn("a:noFill"))
    ln = etree.SubElement(tcPr, qn("a:lnB"))
    ln.set("w", "9525")
    sf = etree.SubElement(ln, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr")).set("val", LINE)
    if fill:
        sf = etree.SubElement(tcPr, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", fill)
    else:
        etree.SubElement(tcPr, qn("a:noFill"))


def _cell_plain(v):
    if isinstance(v, dict):
        key = str(v.get("status", "")).lower().replace(" ", "-").replace("_", "-")
        return STATUS.get(key, ("", "", ""))[0] + " " + str(v.get("label") or STATUS.get(key, ("", "", ""))[2])
    return plain(str("" if v is None else v))


def auto_widths(header, rows, ncols, size, total_w):
    """Column widths from the measured unwrapped width of each column's longest line (brand fonts), so short
    columns do not wrap; when everything cannot fit on one line, share the width by those natural widths."""
    nat = []
    for c in range(ncols):
        best = 0.6
        cells = [(header[c] if c < len(header) else "", True)] + [(r[c] if c < len(r) else "", False) for r in rows]
        for v, bold in cells:
            txt = _cell_plain(v)
            mono = isinstance(v, str) and bool(ISO_DATE.match(v.strip() or "x"))
            font, factor = _ttf(MONO if mono else FONT, bold, size) if CAN_MEASURE else (None, 1.0)
            for line in txt.split("\n"):
                w = (font.getlength(line) / 10 * factor / 72) if font else len(line) * size * 0.5 / 72
                best = max(best, w + 0.222 + 0.06)   # cell margins + slack
        nat.append(best)
    tot = sum(nat)
    if tot <= total_w:
        return [n + (total_w - tot) * n / tot for n in nat]
    return [max(0.9, total_w * n / tot) for n in nat]


ISO_DATE = re.compile(r"^\d{4}-\d{2}(?:-\d{2})?$")


def place_table(ctx, slide, loc, spec, slot_shape=None, box=None):
    if not isinstance(spec, dict) or "rows" not in spec:
        raise SpecError(f"{loc}: table must be an object with \"header\" and \"rows\"")
    header = spec.get("header") or []
    rows = spec.get("rows") or []
    ncols = max([len(header)] + [len(r) for r in rows])
    nrows = len(rows) + (1 if header else 0)
    size = float(spec.get("font_pt", 16))
    row_h = float(spec.get("row_height_in", 0.48))
    if len(rows) > 6 and size >= 16 and ctx.tpl.slide_h < 6:
        ctx.warn(loc, f"{len(rows)} body rows: six at most at 16 pt; split the table across slides")
    if size < 14:
        ctx.warn(loc, f"table text {size} pt is under the 14 pt floor")
    if slot_shape is not None and hasattr(slot_shape, "insert_table"):
        frame = _box_of(slot_shape)
        gf = slot_shape.insert_table(nrows, ncols)
    else:
        frame = box or (_box_of(slot_shape) if slot_shape is not None else None) or ctx.tpl.content_box()
        gf = slide.shapes.add_table(nrows, ncols, emu(frame[0]), emu(frame[1]), emu(frame[2]), emu(frame[3]))
        if slot_shape is not None:
            el = gf._element
            el.getparent().remove(el)
            slot_shape._element.addprevious(el)
            slot_shape._element.getparent().remove(slot_shape._element)
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for ch in list(tblPr):             # no built-in table style: explicit ATR formatting, never gold + white
        if ch.tag == qn("a:tableStyleId"):
            tblPr.remove(ch)
    tblPr.set("firstRow", "1" if header else "0")   # header row flag for screen readers
    tblPr.set("bandRow", "0")
    # column widths: given, or weighted by content length
    total_w = frame[2] if frame else 9.0
    widths = spec.get("widths")
    if not widths:
        widths = auto_widths(header, rows, ncols, size, total_w)
    s = sum(widths)
    widths = [w * total_w / s for w in widths]
    for c, w in enumerate(widths):
        tbl.columns[c].width = emu(w)
    mono_cols = set(spec.get("mono_cols") or [])
    r0 = 0
    if header:
        for c in range(ncols):
            cell = tbl.cell(0, c)
            _cell_frame(cell, NAVY)
            _cell_text(cell, header[c] if c < len(header) else "", size=size, color=WHITE, bold=True)
        tbl.rows[0].height = emu(row_h)
        r0 = 1
    for i, row in enumerate(rows):
        for c in range(ncols):
            v = row[c] if c < len(row) else ""
            cell = tbl.cell(r0 + i, c)
            _cell_frame(cell, None)
            mono = c in mono_cols or (isinstance(v, str) and ISO_DATE.match(v.strip() or "x"))
            _cell_text(cell, v, size=size, color=INK, font=MONO if mono else FONT)
        tbl.rows[r0 + i].height = emu(row_h)
    gf.height = emu(row_h * nrows)
    est = table_height(spec, widths, size, row_h) if CAN_MEASURE else None
    avail = frame[3] if frame else ctx.tpl.content_box()[3]
    if est and est > avail + 0.05:
        ctx.warn(loc, f"table measures about {est:.2f} in tall with the brand fonts; the frame is {avail:.2f} in. "
                      "Cut rows or words, widen columns (\"widths\") or split the table")
    if spec.get("alt"):
        set_alt(gf._element, spec["alt"])
    return gf


def fill_existing_table(ctx, loc, gf, spec):
    """Replace cell text in a table on a cloned slide, keeping each cell's formatting; add rows by copying
    the last body row, drop extra rows."""
    tbl = gf.table
    header = spec.get("header")
    rows = spec.get("rows") or []
    trs = tbl._tbl.tr_lst
    start = 1 if (header is not None or tbl.first_row) else 0   # never overwrite a header row with body data
    if header is not None:
        _replace_row(trs[0], header)
    body = trs[start:]
    if not body:
        raise SpecError(f"{loc}: table has no body row to copy")
    while len(body) < len(rows):
        new = copy.deepcopy(body[-1])
        body[-1].addnext(new)
        body.append(new)
    for extra in body[len(rows):]:
        extra.getparent().remove(extra)
    for tr, row in zip(body, rows):
        _replace_row(tr, row)


def _set_rpr_color(rPr, hexstr):
    for old in rPr.findall(qn("a:solidFill")):
        rPr.remove(old)
    sf = etree.Element(qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr")).set("val", hexstr)
    ln = rPr.find(qn("a:ln"))
    if ln is not None:
        ln.addnext(sf)
    else:
        rPr.insert(0, sf)


def _replace_row(tr, values):
    """Cell text keeping each cell's formatting: plain text takes the last run's style (the label style in a
    status cell); a {"status": ...} value writes the glyph in its status color with the first run's style,
    then the label."""
    tcs = tr.findall(qn("a:tc"))
    for tc, v in zip(tcs, values):
        txBody = tc.find(qn("a:txBody"))
        ps = txBody.findall(qn("a:p"))
        first = ps[0] if ps else etree.SubElement(txBody, qn("a:p"))
        runs = first.findall(qn("a:r"))
        rprs = [r.find(qn("a:rPr")) for r in runs]
        r_first = copy.deepcopy(rprs[0]) if rprs and rprs[0] is not None else None
        r_last = copy.deepcopy(rprs[-1]) if rprs and rprs[-1] is not None else None
        pPr = copy.deepcopy(first.find(qn("a:pPr"))) if first.find(qn("a:pPr")) is not None else None
        for p in ps:
            txBody.remove(p)
        if isinstance(v, str) and v.lower().startswith("status:"):
            v = {"status": v.split(":", 1)[1].strip()}
        if isinstance(v, dict) and "status" in v:
            key = str(v["status"]).lower().replace(" ", "-").replace("_", "-")
            if key not in STATUS:
                raise SpecError(f"unknown status '{v['status']}'; use {', '.join(STATUS)}")
            glyph, gcol, label = STATUS[key]
            p = etree.SubElement(txBody, qn("a:p"))
            if pPr is not None:
                p.append(copy.deepcopy(pPr))
            for text, rpr, col in ((glyph, r_first, gcol), ("  " + str(v.get("label", label)), r_last, None)):
                r = etree.SubElement(p, qn("a:r"))
                rp = copy.deepcopy(rpr) if rpr is not None else etree.Element(qn("a:rPr"))
                if col:
                    _set_rpr_color(rp, col)
                r.append(rp)
                etree.SubElement(r, qn("a:t")).text = text
            continue
        for line in str("" if v is None else v).split("\n"):
            p = etree.SubElement(txBody, qn("a:p"))
            if pPr is not None:
                p.append(copy.deepcopy(pPr))
            r = etree.SubElement(p, qn("a:r"))
            if r_last is not None:
                r.append(copy.deepcopy(r_last))
            etree.SubElement(r, qn("a:t")).text = line


# --------------------------------------------------------------------------- slides: clone, remove

def clone_slide(prs, src):
    """Duplicate a template slide (its shapes and their images, charts, media and links) onto a new slide
    with the same layout. Charts are deep-copied so the copy owns its chart and workbook."""
    from pptx.parts.chart import ChartPart
    from pptx.parts.embeddedpackage import EmbeddedXlsxPart
    new = prs.slides.add_slide(src.slide_layout)
    tree = new.shapes._spTree
    for el in list(tree):
        if el.tag not in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            tree.remove(el)
    rid_map = {}
    for rId, rel in list(src.part.rels.items()):
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
            continue
        if rel.is_external:
            rid_map[rId] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            continue
        target = rel.target_part
        if rel.reltype == RT.CHART:
            pkg = target.package
            cp = ChartPart.load(pkg.next_partname(ChartPart.partname_template), target.content_type, pkg,
                                target.blob)
            sub = {}
            for crid, crel in list(target.rels.items()):
                if crel.is_external:
                    sub[crid] = cp.relate_to(crel.target_ref, crel.reltype, is_external=True)
                elif crel.reltype == RT.PACKAGE:
                    sub[crid] = cp.relate_to(EmbeddedXlsxPart.new(crel.target_part.blob, pkg), RT.PACKAGE)
                else:
                    sub[crid] = cp.relate_to(crel.target_part, crel.reltype)
            _remap_rids(cp._element, sub)
            target = cp
        rid_map[rId] = new.part.relate_to(target, rel.reltype)
    ext = tree.find(qn("p:extLst"))
    for el in src.shapes._spTree:
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr"), qn("p:extLst")):
            continue
        dup = copy.deepcopy(el)
        _remap_rids(dup, rid_map)
        if ext is not None:
            ext.addprevious(dup)
        else:
            tree.append(dup)
    src_bg = src._element.cSld.find(qn("p:bg"))
    if src_bg is not None:
        new._element.cSld.insert(0, copy.deepcopy(src_bg))
    return new


def _remap_rids(el, mapping):
    rns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    for node in el.iter():
        for attr, val in list(node.attrib.items()):
            if attr.startswith(rns) and val in mapping:
                node.set(attr, mapping[val])


def remove_slides(prs, keep_ids):
    """Remove every slide whose sldId is not in keep_ids: drop the relationship and the sldId, so the
    slide, its notes and anything only it used are not written (no orphans)."""
    sldIdLst = prs.slides._sldIdLst
    removed = set()
    for sldId in list(sldIdLst):
        if int(sldId.get("id")) in keep_ids:
            continue
        rId = sldId.get(qn("r:id"))
        removed.add(sldId.get("id"))
        prs.part.drop_rel(rId)
        sldIdLst.remove(sldId)
    pres = prs.part._element
    # sections (p14:sectionLst) and custom shows can reference removed slides: drop the stale entries
    for sec_id in pres.iter("{http://schemas.microsoft.com/office/powerpoint/2010/main}sldId"):
        if sec_id.get("id") in removed:
            sec_id.getparent().remove(sec_id)
    cust = pres.find(qn("p:custShowLst"))
    if cust is not None:
        pres.remove(cust)
    return len(removed)


def fix_app_properties(prs):
    """docProps/app.xml still counts the template's slides; set Slides/Notes and drop the stale slide-title
    list (optional elements PowerPoint rebuilds on its next save)."""
    for part in prs.part.package.iter_parts():
        if str(part.partname) != "/docProps/app.xml":
            continue
        try:
            root = etree.fromstring(part.blob)
        except etree.XMLSyntaxError:
            return
        ns = root.nsmap.get(None, "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties")
        n_notes = sum(1 for s in prs.slides if s.has_notes_slide)
        for tag, val in (("Slides", len(prs.slides)), ("Notes", n_notes), ("HiddenSlides", 0)):
            el = root.find("{%s}%s" % (ns, tag))
            if el is not None:
                el.text = str(val)
        for tag in ("HeadingPairs", "TitlesOfParts"):
            el = root.find("{%s}%s" % (ns, tag))
            if el is not None:
                root.remove(el)
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        return


def renumber_parts(prs):
    """Name slide and notes parts slide1..N in deck order (cloned and new slides were appended after the
    showcase, whose parts are no longer written)."""
    for i, slide in enumerate(prs.slides, 1):
        slide.part.partname = PackURI(f"/ppt/slides/slide{i}.xml")
    for i, slide in enumerate(prs.slides, 1):
        if slide.has_notes_slide:
            slide.notes_slide.part.partname = PackURI(f"/ppt/notesSlides/notesSlide{i}.xml")


# --------------------------------------------------------------------------- slide builder

def _shape_by_name(slide, name):
    """First shape with this name (the Selection Pane name), searching inside groups too."""
    for shp in _walk_all(slide.shapes):
        if shp.name == name:
            return shp
    return None


def _walk_all(shapes):
    for shp in shapes:
        yield shp
        if shp.shape_type == 6:
            yield from _walk_all(shp.shapes)


def _delete(shape):
    el = shape._element
    el.getparent().remove(el)


def section_defaults(ctx, sspec, slots, fields):
    """"section": N fills the eyebrow (SECTION 0N  ·  NAME), a divider's number and title and the
    outline from the agenda items, so the numbers always match the agenda plates."""
    n = sspec.get("section")
    if n is None:
        return
    try:
        n = int(n)
    except (TypeError, ValueError):
        raise SpecError(f"section must be a number, got {n!r}")
    name = ctx.agenda[n - 1] if 0 < n <= len(ctx.agenda) else None
    names = {s.name for s in slots}
    if "eyebrow" in names and "eyebrow" not in fields and name:
        eb = next(s for s in slots if s.name == "eyebrow")
        if "no \"SECTION" in eb.meta.get("format", "") or "NAME only" in eb.meta.get("purpose", ""):
            fields["eyebrow"] = plain(name).upper()
        else:
            fields["eyebrow"] = f"SECTION {n:02d}  ·  {plain(name).upper()}"
    if "number" in names and "number" not in fields:
        fields["number"] = f"{n:02d}"
    if "outline" in names and "outline" not in fields and ctx.agenda:
        fields["outline"] = [(f"**{i:02d}\t{plain(t)}**" if i == n else f"{i:02d}\t{plain(t)}")
                             for i, t in enumerate(ctx.agenda, 1)]
    if "title" in names and "title" not in fields and name and "number" in names:
        fields["title"] = plain(name)


def build_slide(ctx, prs, i, sspec, clones):
    tpl = ctx.tpl
    if not isinstance(sspec, dict):
        raise SpecError(f"slide {i}: each slide must be an object")
    fields = dict(sspec.get("fields") or {})
    for k, v in sspec.items():
        if k not in RESERVED:
            fields[k] = v
    if sspec.get("from_slide"):
        n = int(sspec["from_slide"])
        if not 1 <= n <= len(clones):
            raise SpecError(f"slide {i}: from_slide {n} out of range (template has {len(clones)} slides)")
        slide = clone_slide(prs, clones[n - 1])
        layout = slide.slide_layout
    else:
        if "layout" not in sspec:
            raise SpecError(f"slide {i}: needs \"layout\" (or \"from_slide\")")
        layout = tpl.find_layout(sspec["layout"])
        slide = prs.slides.add_slide(layout)
    loc0 = f"slide {i} ({layout.name})"
    slots = tpl.slots(layout)
    by_idx = {s.idx: s for s in slots}
    section_defaults(ctx, sspec, slots, fields)
    phs = {}
    for shp in slide.placeholders:
        phs[shp.placeholder_format.idx] = shp
    filled, deleted = set(), set()
    fill_use, visual = [], False     # (slot, lines used, line budget) of the body slots; a picture/chart/table placed
    alts = {k[:-4]: v for k, v in fields.items() if k.endswith("_alt")}
    for key, value in fields.items():
        if key.endswith("_alt"):
            continue
        slot = lookup(slots, key)
        kl = key.lower()
        loc = f"{loc0} › {key}"
        special = kl in ("chart", "table") or kl in PIC_ALIASES or kl in ("video", "media")
        if slot is None and not special:
            names = ", ".join(s.name for s in slots) or "(no placeholders: use chart, table, image, textboxes)"
            raise SpecError(f"{loc}: no such field. This layout has: {names}")
        shape = phs.get(slot.idx) if slot else None
        if slot and slot.idx in filled:
            raise SpecError(f"{loc}: '{slot.name}' is already filled on this slide")
        if slot and shape is None and not (slot.idx in deleted):
            raise SpecError(f"{loc}: placeholder idx {slot.idx} is missing on the slide")
        box = None
        if isinstance(value, dict) and isinstance(value.get("box"), dict):
            b = value["box"]
            box = (b["x"], b["y"], b["w"], b["h"])
        if value is None or (isinstance(value, str) and value.strip() == ""):
            if shape is not None:
                _delete(shape)
                deleted.add(slot.idx)
            continue
        if kl == "chart" or (slot and slot.kind == "chart"):
            target = shape if box is None else None
            place_chart(ctx, slide, loc, value, slot_shape=target, box=box or (slot.box if slot and target is None else None))
            if shape is not None and target is None:
                _delete(shape)
        elif kl == "table" or (slot and slot.kind == "tbl"):
            target = shape if box is None else None
            place_table(ctx, slide, loc, value, slot_shape=target, box=box)
            if shape is not None and target is None:
                _delete(shape)
        elif kl in ("video", "media", "movie") or (slot and slot.kind == "media"):
            place_media(ctx, slide, loc, value, shape, slot)
        elif slot is None or slot.kind == "pic":
            pspec = value if isinstance(value, dict) else {"path": value}
            if "alt" not in pspec:
                for k2 in (key, slot.name if slot else None):
                    if k2 in alts:
                        pspec = dict(pspec, alt=alts[k2])
                        break
            place_picture(ctx, slide, loc, pspec, slot_shape=shape if box is None else None, box=box, slot=slot)
            if shape is not None and box is not None:
                _delete(shape)
        else:   # text
            if isinstance(value, dict) and "text" in value and len(value) <= 2:
                value = value["text"]
            if isinstance(value, dict):
                raise SpecError(f"{loc}: '{slot.name}' is a text placeholder; give a string or a list")
            paras = to_paragraphs(value)
            fb = check_budget(ctx, loc, slot, paras)
            fill_text(shape, paras, size_pt=fb)
            res = check_measured(ctx, loc, slot, shape, paras, fb)
            budget = line_budget(slot.meta) if slot.name in UNDERFILL_SLOTS else None
            if budget:
                fill_use.append((slot.name, used_lines(slot.meta, paras, res), budget))
        if (slot and slot.kind != "text") or special:
            visual = True
        if slot:
            filled.add(slot.idx)
    keep = set(sspec.get("keep") or [])
    # remove named shapes / fields
    for name in sspec.get("remove") or []:
        s = lookup(slots, name)
        target = phs.get(s.idx) if s else _shape_by_name(slide, name)
        if target is None:
            ctx.warn(loc0, f"remove: no shape or field named '{name}'")
            continue
        if target._element.getparent() is not None:
            _delete(target)
        if s:
            deleted.add(s.idx)
    # text of named (non-placeholder) shapes on cloned slides
    for name, value in (sspec.get("shapes") or {}).items():
        shp = _shape_by_name(slide, name)
        if shp is None or not shp.has_text_frame:
            raise SpecError(f"{loc0}: shapes: no text shape named '{name}'")
        fill_text(shp, to_paragraphs(value))
    for name, tspec in (sspec.get("tables") or {}).items():
        shp = _shape_by_name(slide, name)
        if shp is None or not getattr(shp, "has_table", False):
            raise SpecError(f"{loc0}: tables: no table named '{name}'")
        fill_existing_table(ctx, loc0, shp, tspec)
    for j, ispec in enumerate(sspec.get("images") or []):
        loc = f"{loc0} › images[{j}]"
        if ispec.get("replace"):
            old = _shape_by_name(slide, ispec["replace"])
            if old is None:
                raise SpecError(f"{loc}: no shape named '{ispec['replace']}'")
            pic = place_picture(ctx, slide, loc, ispec, box=_box_of(old))
            old._element.addprevious(pic._element)
            _delete(old)
        else:
            b = ispec.get("box")
            if not b:
                raise SpecError(f"{loc}: needs \"box\" {{x, y, w, h}} in inches or \"replace\"")
            place_picture(ctx, slide, loc, ispec, box=(b["x"], b["y"], b["w"], b["h"]))
    for j, tb in enumerate(sspec.get("textboxes") or []):
        b = tb.get("box")
        if not b:
            raise SpecError(f"{loc0} › textboxes[{j}]: needs \"box\"")
        shp = slide.shapes.add_textbox(emu(b["x"]), emu(b["y"]), emu(b["w"]), emu(b["h"]))
        tf = shp.text_frame
        tf.word_wrap = True
        fill_text(shp, to_paragraphs(tb.get("text", "")))
        for p in tf.paragraphs:
            for r in p.runs:
                _font(r.font, float(tb.get("size", 18)), str(tb.get("color", INK)).lstrip("#"),
                      bold=True if tb.get("bold") or r.font.bold else None, name=tb.get("font", FONT))
        if float(tb.get("size", 18)) < 14 and tpl.slide_h < 8:
            ctx.warn(f"{loc0} › textboxes[{j}]", "text under the 14 pt floor")
    # delete every placeholder left unfilled (no "Click to add" prompts, no empty frames in LibreOffice)
    for shp in list(slide.placeholders):
        pf = shp.placeholder_format
        if pf.type in (PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE,
                       PP_PLACEHOLDER.HEADER):
            continue   # slide numbers and footers are kept (cloned slides carry them)
        idx = pf.idx
        if idx in filled or shp.name in keep or (by_idx.get(idx) and by_idx[idx].name in keep):
            continue
        _delete(shp)
    underfill_hint(ctx, loc0, layout, fill_use, visual or bool(sspec.get("images")))
    # layouts whose numbered plates, stations or panels are fixed art: fewer entries leave empty ones showing
    for tag, prefix, total, what, fix in FIXED_SETS:
        if layout.name.startswith("ATR - ") and tag in layout.name:
            n_set = sum(1 for x in filled if x in by_idx and re.fullmatch(prefix + r"\d", by_idx[x].name))
            if n_set < total:
                ctx.warn(loc0, f"{n_set} of {total} {what} filled: the other {total - n_set} are fixed on the layout "
                               f"and show empty. {fix}")
    # NASA quad: the acknowledgement sentence is required text on the slide
    if "NASA" in layout.name and any(s.name == "ack" for s in slots):
        ack = next(s for s in slots if s.name == "ack")
        if ack.idx not in filled:
            ctx.refuse(loc0, "NASA quad without the acknowledgement sentence: fill \"ack\" with the exact "
                             "GSFC sentence (references/quad-charts.md)")
    notes = sspec.get("speaker_notes", sspec.get("notes"))
    if "notes" in sspec and "speaker_notes" not in sspec and any(sl.name == "notes" for sl in slots) \
            and "notes" not in (sspec.get("fields") or {}):
        ctx.warn(loc0, "slide-level \"notes\" are speaker notes; this layout also has a placeholder called "
                       "\"notes\": put its text in \"fields\": {\"notes\": ...} (or use \"speaker_notes\")")
    if notes:
        text = "\n".join(notes) if isinstance(notes, list) else str(notes)
        slide.notes_slide.notes_text_frame.text = text
    if ctx.verbose:
        print(f"  {loc0}: filled {sorted(by_idx[x].name for x in filled if x in by_idx)}")
    return slide


# --------------------------------------------------------------------------- deck-level checks

def iter_text(shapes):
    for shp in shapes:
        if shp.shape_type is not None and getattr(shp, "shapes", None) is not None and shp.shape_type == 6:
            yield from iter_text(shp.shapes)
            continue
        if getattr(shp, "has_text_frame", False) and shp.has_text_frame:
            for p in shp.text_frame.paragraphs:
                yield shp.name, "".join(r.text for r in p.runs)
        if getattr(shp, "has_table", False) and shp.has_table:
            for r, row in enumerate(shp.table.rows):
                for c, cell in enumerate(row.cells):
                    yield f"{shp.name} r{r + 1}c{c + 1}", cell.text


def deck_checks(ctx, prs):
    titles = {}
    for n, slide in enumerate(prs.slides, 1):
        loc = f"slide {n} ({slide.slide_layout.name})"
        for where, text in iter_text(slide.shapes):
            if not text.strip():
                continue
            for rx, label in PH_RULES:
                m = rx.search(text)
                if m:
                    msg = f"leftover {label} {m.group(0)!r} in {where}"
                    (ctx.warn if ctx.allow_placeholders else ctx.refuse)(loc, msg)
            for cid, rx, sev, msg, fix in WORDING_RULES:
                m = rx.search(text)
                if m and (sev == "error" or cid in REFUSE_IDS):
                    ctx.refuse(loc, f"{msg} Found {m.group(0)!r} in {where}. {fix}")
                elif m and sev == "warning":
                    ctx.warn(loc, f"{msg} Found {m.group(0)!r} in {where}. {fix}")
        if slide.has_notes_slide:
            nt = slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame else ""
            for rx, label in PH_RULES[:1]:
                m = rx.search(nt)
                if m:
                    (ctx.warn if ctx.allow_placeholders else ctx.refuse)(
                        f"{loc} notes", f"leftover {label} {m.group(0)!r}")
            for cid, rx, sev, msg, fix in WORDING_RULES:
                m = rx.search(nt)
                if m and sev in ("error", "warning"):
                    ctx.warn(f"{loc} notes", f"{msg} Found {m.group(0)!r}.")
        # alt text on pictures, charts, media
        for shp in _walk(slide.shapes):
            el = shp._element
            tag = etree.QName(el).localname
            if tag == "pic" or (tag == "graphicFrame" and el.find(".//" + _c("chart")) is not None):
                cNvPr = el.find(".//" + qn("p:cNvPr"))
                deco = cNvPr is not None and cNvPr.find(".//{http://schemas.microsoft.com/office/drawing/2017/"
                                                        "decorative}decorative") is not None
                if cNvPr is not None and cNvPr.get("descr") is None and not deco:
                    ctx.refuse(loc, f"{shp.name} has no alt text")
        # titles
        t = slide_title(slide)
        if t:
            titles.setdefault(t.lower(), []).append(n)
        elif ctx.tpl.slide_h < 8:
            ctx.infos.append(f"{loc}: no slide title (screen readers navigate by titles; "
                             "references/accessibility.md 3.1)")
    for t, ns in titles.items():
        if len(ns) > 1 and len(prs.slides) > 1:
            ctx.warn(f"slides {', '.join(map(str, ns))}", f"share the title \"{t}\"; give each slide a unique title")


def slide_title(slide):
    """Text of the slide's title placeholder, found by type (ATR layouts give titles a non-zero idx, which
    python-pptx's shapes.title does not look for)."""
    for shp in slide.placeholders:
        if shp.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE) and shp.has_text_frame:
            return plain(shp.text_frame.text).replace("\v", " ").strip()
    return ""


def _walk(shapes):
    for shp in shapes:
        if shp.shape_type == 6:
            yield from _walk(shp.shapes)
        else:
            yield shp


def verify_package(path):
    """Re-read the saved file: every part reachable from the package relationships, every sldId with a
    slide relationship, no duplicate zip entries. Returns a list of problems."""
    problems = []
    z = zipfile.ZipFile(path)
    names = [i.filename for i in z.infolist()]
    if len(names) != len(set(names)):
        problems.append("duplicate zip entries")
    names_set = set(names)

    def rels_of(part):
        d, b = posixpath.split(part)
        return posixpath.join(d, "_rels", b + ".rels")

    reach, todo = set(), [""]
    while todo:
        part = todo.pop()
        rp = "_rels/.rels" if part == "" else rels_of(part)
        if rp not in names_set:
            continue
        root = etree.fromstring(z.read(rp))
        for rel in root:
            if rel.get("TargetMode") == "External":
                continue
            tgt = rel.get("Target")
            base = posixpath.dirname(part)
            full = tgt.lstrip("/") if tgt.startswith("/") else posixpath.normpath(posixpath.join(base, tgt))
            if full not in names_set:
                problems.append(f"{rp} points to a missing part {full}")
                continue
            if full not in reach:
                reach.add(full)
                todo.append(full)
    for n in names:
        if n.endswith(".rels") or n == "[Content_Types].xml" or n.endswith("/"):
            continue
        if n not in reach:
            problems.append(f"orphan part {n}")
    pres = etree.fromstring(z.read("ppt/presentation.xml"))
    prels = etree.fromstring(z.read("ppt/_rels/presentation.xml.rels"))
    rids = {r.get("Id"): r.get("Type") for r in prels}
    for s in pres.iter(qn("p:sldId")):
        rid = s.get(qn("r:id"))
        if not rids.get(rid, "").endswith("/slide"):
            problems.append(f"sldId {s.get('id')} has no slide relationship ({rid})")
    return problems


# --------------------------------------------------------------------------- listing

def list_layouts(tpl: Template, as_json=False):
    out = {"template": str(tpl.path), "layouts_json": str(tpl.spec_path) if tpl.spec_path else None,
           "slide_in": [round(tpl.slide_w, 3), round(tpl.slide_h, 3)], "layouts": [], "showcase_slides": []}
    for i, layout in enumerate(tpl.prs.slide_layouts):
        jl = tpl.json_layout(layout) or {}
        entry = {"index": i, "name": layout.name, "purpose": jl.get("purpose", ""), "when": jl.get("when", ""),
                 "fields": []}
        for s in tpl.slots(layout):
            entry["fields"].append({"field": s.name, "idx": s.idx, "kind": s.kind, "budget": s.budget_str(),
                                    "box_in": [round(v, 3) for v in s.box] if s.box else None,
                                    "about": (s.meta.get("purpose") or s.meta.get("prompt") or "")[:160]})
        out["layouts"].append(entry)
    for n, slide in enumerate(tpl.prs.slides, 1):
        out["showcase_slides"].append({"from_slide": n, "layout": slide.slide_layout.name,
                                       "named_shapes": [s.name for s in slide.shapes if not s.is_placeholder]})
    if as_json:
        print(json.dumps(out, indent=1, ensure_ascii=False))
        return
    print(f"{tpl.path.name}  ({tpl.slide_w:g} x {tpl.slide_h:g} in)  field names from: "
          f"{tpl.spec_path.name if tpl.spec_path else 'placeholder introspection (no layouts JSON)'}")
    for L in out["layouts"]:
        print(f"\n[{L['index']}] {L['name']}")
        for key in ("purpose", "when"):
            if L[key]:
                print(textwrap.fill(L[key], 110, initial_indent="    ", subsequent_indent="    "))
        if not L["fields"]:
            print("    (no placeholders: chart, table, image and textboxes go in the content area)")
        for f in L["fields"]:
            about = f["about"].replace("\n", " / ")
            print(f"    {f['field']:<16} idx {f['idx']:<4} {f['kind']:<6} {f['budget']:<30} {about[:95]}")
    print("\nShowcase slides (clone one with \"from_slide\": N; its named shapes take \"shapes\"/\"tables\"/"
          "\"images\" edits):")
    for s in out["showcase_slides"]:
        named = ", ".join(s["named_shapes"][:8]) + (" ..." if len(s["named_shapes"]) > 8 else "")
        print(f"    {s['from_slide']:>2}  {s['layout']}" + (f"   [{named}]" if named else ""))
    print("\nAliases: image/photo/figure -> first picture slot; chart/table -> chart/table slot, else the picture "
          "slot's frame, else the content area; video -> media slot; body -> main text slot; ph<idx> -> any "
          "placeholder. <field>_alt = alt text (\"\" = decorative).")


# --------------------------------------------------------------------------- render + check

def render(out_path: Path, render_dir: Path, dpi: int):
    script = HERE / "render_office.sh"
    render_dir.mkdir(parents=True, exist_ok=True)
    r = subprocess.run(["bash", str(script), str(out_path), str(render_dir), str(dpi)], capture_output=True,
                       text=True)
    if r.returncode != 0:
        print(r.stdout[-2000:], r.stderr[-2000:], file=sys.stderr)
        return None
    pages = sorted(render_dir.glob(f"{out_path.stem}-page-*.png"),
                   key=lambda p: int(re.search(r"-page-(\d+)\.png$", p.name).group(1)))
    sheet = None
    if pages and Image is not None:
        sheet = render_dir / f"{out_path.stem}-contact.png"
        contact_sheet(pages, sheet)
    return pages, sheet


def contact_sheet(pages, out, cols=3, width=2400, gap=24, bg=(128, 128, 128)):
    ims = [Image.open(p).convert("RGB") for p in pages]
    cols = min(cols, len(ims))
    cw = (width - gap * (cols + 1)) // cols
    ch = int(cw * ims[0].height / ims[0].width)
    rows = math.ceil(len(ims) / cols)
    sheet = Image.new("RGB", (width, rows * (ch + gap) + gap), bg)
    for k, im in enumerate(ims):
        r, c = divmod(k, cols)
        sheet.paste(im.resize((cw, ch)), (gap + c * (cw + gap), gap + r * (ch + gap)))
    sheet.save(out)


def run_brand_check(out_path: Path):
    env = dict(os.environ, PYTHONDONTWRITEBYTECODE="1")
    r = subprocess.run([sys.executable, str(HERE / "brand_check.py"), str(out_path)], capture_output=True, text=True,
                       env=env)
    print(r.stdout)
    if r.stderr.strip():
        print(r.stderr, file=sys.stderr)
    return r.returncode


# --------------------------------------------------------------------------- main

def load_spec(path: Path):
    text = path.read_text(encoding="utf-8")
    if path.suffix.lower() in (".yml", ".yaml"):
        try:
            import yaml
        except ImportError:
            raise SpecError("YAML specs need PyYAML (pip install pyyaml); or write the spec as JSON")
        return yaml.safe_load(text)
    try:
        return json.loads(text)
    except json.JSONDecodeError as exc:
        raise SpecError(f"{path}: invalid JSON: {exc}")


def collect_agenda(tpl, slides):
    for s in slides:
        if not isinstance(s, dict) or s.get("from_slide"):
            continue
        lay = str(s.get("layout", ""))
        if "agenda" in lay.lower():
            f = dict(s.get("fields") or {})
            f.update({k: v for k, v in s.items() if k not in RESERVED})
            items = []
            for k in range(1, 10):
                v = f.get(f"item{k}")
                if v is None:
                    break
                items.append(plain(v if isinstance(v, str) else " ".join(map(str, v))))
            return items
    return []


def section_consistency(ctx, prs):
    """Eyebrow and divider numbers vs the agenda plates."""
    if not ctx.agenda:
        return
    for n, slide in enumerate(prs.slides, 1):
        slots = {s.idx: s.name for s in ctx.tpl.slots(slide.slide_layout)}
        vals = {}
        for shp in slide.placeholders:
            nm = slots.get(shp.placeholder_format.idx)
            if nm and shp.has_text_frame:
                vals[nm] = plain(shp.text_frame.text).strip()
        eb = vals.get("eyebrow", "")
        m = re.match(r"SECTION (\d+)\s+·\s+(.+)$", eb)
        if m:
            k = int(m.group(1))
            if k > len(ctx.agenda):
                ctx.warn(f"slide {n}", f"eyebrow says section {k:02d} but the agenda has {len(ctx.agenda)} items")
            elif m.group(2).strip().upper() != ctx.agenda[k - 1].upper():
                ctx.warn(f"slide {n}", f"eyebrow '{eb}' does not match agenda item {k:02d} "
                                       f"'{ctx.agenda[k - 1]}' (use \"section\": {k})")
        if "number" in vals and vals["number"].isdigit() and "Divider" in slide.slide_layout.name:
            k = int(vals["number"])
            if k > len(ctx.agenda):
                ctx.warn(f"slide {n}", f"divider {k:02d} but the agenda has {len(ctx.agenda)} items")


EXAMPLES = """\
examples (run from the skill root, atr-lab-design/):
  python3 scripts/new_deck.py --list-templates
  python3 scripts/new_deck.py --list-layouts presentation          # every layout, field, budget
  python3 scripts/new_deck.py --list-layouts quad --json
  python3 scripts/new_deck.py scripts/examples/decks/minimal.json -o out/minimal.pptx
  python3 scripts/new_deck.py scripts/examples/decks/sponsor-intro.json -o out/sponsor-intro.pptx --render --check
  python3 scripts/new_deck.py scripts/examples/decks/research-talk-skeleton.json -o out/draft.pptx \\
      --allow-placeholders                                        # a draft that still has [placeholders]
  python3 scripts/new_deck.py talk.yaml --strict                  # warnings (budgets, ppi) fail the build

a minimal spec:
  {"template": "presentation", "output": "hello.pptx", "slides": [
     {"layout": "ATR - Title", "title": "Physical AI at Kent State", "subtitle": "Lab overview",
      "presenter": "Jong-Hoon Kim, associate professor of computer science"},
     {"layout": "ATR - Title + Content", "title": "What the lab does",
      "body": ["Telepresence robotics", ["and tele-embodiment"], "Physical AI"]},
     {"layout": "ATR - Thank You", "title": "Thank you", "presenter": "Jong-Hoon Kim",
      "email": "jkim72@kent.edu"}]}

Full guide: references/presentations.md (layout catalog, deck recipes, writing rules, QA).
"""


def main(argv=None):
    ap = argparse.ArgumentParser(
        prog="new_deck.py", formatter_class=argparse.RawDescriptionHelpFormatter,
        description="Build an ATR Lab deck (or poster, quad chart, flyer...) from a JSON/YAML spec on the "
                    "ATR templates: showcase slides removed, layouts filled, unused placeholders deleted, "
                    "checked.", epilog=EXAMPLES)
    ap.add_argument("spec", nargs="?", help="deck spec (.json, or .yaml/.yml with PyYAML)")
    ap.add_argument("-o", "--output", help="output .pptx (overrides the spec's \"output\")")
    ap.add_argument("--template", help="override the spec's template (alias or .pptx path)")
    ap.add_argument("--layouts", help="layouts JSON for a custom template (default: found next to the template)")
    ap.add_argument("--list-layouts", metavar="TEMPLATE", help="print layouts, fields, budgets and showcase slides")
    ap.add_argument("--list-templates", action="store_true", help="print the template aliases")
    ap.add_argument("--json", action="store_true", help="with --list-layouts: machine-readable output")
    ap.add_argument("--allow-placeholders", action="store_true",
                    help="build a draft that still contains [bracketed] placeholders (warned, not refused)")
    ap.add_argument("--render", action="store_true", help="render PNGs + a contact sheet (scripts/render_office.sh)")
    ap.add_argument("--render-dir", help="render folder (default: <output folder>/<name>-render)")
    ap.add_argument("--dpi", type=int, default=110, help="render resolution (default 110)")
    ap.add_argument("--check", action="store_true", help="run scripts/brand_check.py on the result")
    ap.add_argument("--strict", action="store_true", help="exit 3 when there are warnings")
    ap.add_argument("-v", "--verbose", action="store_true", help="print the field mapping of every slide")
    a = ap.parse_args(argv)

    try:
        if a.list_templates:
            for k in sorted(TEMPLATE_ALIASES):
                print(f"  {k:<20} assets/templates/{TEMPLATE_ALIASES[k]}")
            return 0
        if a.list_layouts:
            list_layouts(Template(a.list_layouts, a.layouts), a.json)
            return 0
        if not a.spec:
            ap.print_help()
            return 1
        spec_path = Path(a.spec).expanduser().resolve()
        if not spec_path.exists():
            raise SpecError(f"spec not found: {a.spec}")
        spec = load_spec(spec_path)
        if not isinstance(spec, dict) or not isinstance(spec.get("slides"), list) or not spec["slides"]:
            raise SpecError("the spec needs a non-empty \"slides\" list")
        tpl = Template(a.template or spec.get("template", "presentation"), a.layouts or spec.get("layouts"))
        out = Path(a.output or spec.get("output") or (spec_path.stem + ".pptx")).expanduser()
        if out.suffix.lower() != ".pptx":
            raise SpecError("output must be a .pptx file")
        out = out.resolve()
        ctx = Ctx(tpl=tpl, spec_dir=spec_path.parent, allow_placeholders=a.allow_placeholders, verbose=a.verbose)
        ctx.agenda = collect_agenda(tpl, spec["slides"])
        prs = tpl.prs
        originals = list(prs.slides)
        original_ids = {int(s.slide_id) for s in originals}
        for i, sspec in enumerate(spec["slides"], 1):
            build_slide(ctx, prs, i, sspec, originals)
        new_ids = {int(s.slide_id) for s in prs.slides} - original_ids
        removed = remove_slides(prs, new_ids)
        renumber_parts(prs)
        fix_app_properties(prs)
        cp = prs.core_properties
        first_title = slide_title(prs.slides[0]) if len(prs.slides) else ""
        cp.title = spec.get("title") or first_title.replace("\n", " ")
        cp.subject = spec.get("subject", "")
        cp.author = spec.get("author", "")
        cp.last_modified_by = spec.get("author", "")
        cp.keywords = spec.get("keywords", "")
        cp.comments = ""
        cp.revision = 1
        now = _dt.datetime.now(_dt.timezone.utc).replace(tzinfo=None, microsecond=0)
        cp.created = now
        cp.modified = now
        deck_checks(ctx, prs)
        section_consistency(ctx, prs)
        if any(l.name in ("ATR - Title", "ATR - Title (Light)", "ATR - Thank You")
               for l in (s.slide_layout for s in prs.slides)):
            ctx.infos.append("The Kent State wordmark on the Title and Thank You layouts is a working copy of the "
                             "old Stacked raster (colors corrected): replace it with the official file from "
                             "kent.edu/brand/logos before public use (references/logo-system.md).")
    except SpecError as exc:
        print(f"new_deck.py: {exc}", file=sys.stderr)
        return 1

    for w in ctx.warnings:
        print(f"WARNING  {w}")
    for w in ctx.infos:
        print(f"info     {w}")
    if ctx.refusals:
        print("\nREFUSED (nothing written):", file=sys.stderr)
        for r in ctx.refusals:
            print(f"  {r}", file=sys.stderr)
        if any("leftover" in r for r in ctx.refusals):
            print("  Replace the placeholders, or pass --allow-placeholders for a draft.", file=sys.stderr)
        return 2
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    problems = verify_package(out)
    status = 0
    print(f"wrote {out}  ({len(prs.slides)} slides from {tpl.path.name}; {removed} showcase slides removed; "
          f"{len(ctx.warnings)} warning(s))")
    if problems:
        status = 3
        print("PACKAGE PROBLEMS:", *problems, sep="\n  ", file=sys.stderr)
    if a.render:
        rdir = Path(a.render_dir).expanduser() if a.render_dir else out.parent / f"{out.stem}-render"
        res = render(out, rdir, a.dpi)
        if res is None:
            status = 3
            print("render failed (see above); is LibreOffice installed?", file=sys.stderr)
        else:
            pages, sheet = res
            print(f"rendered {len(pages)} page(s) to {rdir}" + (f"; contact sheet {sheet}" if sheet else "")
                  + ". Look at every slide.")
    if a.check:
        rc = run_brand_check(out)
        if rc != 0:
            status = 3
    if a.strict and ctx.warnings:
        status = 3
    return status


if __name__ == "__main__":
    sys.exit(main())
