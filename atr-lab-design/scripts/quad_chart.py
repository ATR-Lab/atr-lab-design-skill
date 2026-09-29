#!/usr/bin/env python3
"""
quad_chart.py: build a finished ATR Lab quad chart (.pptx) from a JSON spec, and check quad charts.
=====================================================================================================

Builds on the shipped template `assets/templates/ATR-Quad-Chart-Template.pptx` and its placeholder map
`assets/templates/quad-layouts.json`. For every quad in the spec the script DUPLICATES the template's showcase
slide for that type (never New Slide, so the NASA acknowledgement sentence stays real slide text), fills the
placeholders, deletes every placeholder it did not fill, removes the template's own slides, and then checks the
result: NASA rules, measured text fit and character budgets, leftover prompts and placeholders, alt text,
milestone glyph colors. Rules and writing guidance: references/quad-charts.md.

Usage
-----
    python3 scripts/quad_chart.py spec.json -o quad.pptx                 # build + built-in checks
    python3 scripts/quad_chart.py spec.json -o quad.pptx --check         # ... + brand_check.py on the result
    python3 scripts/quad_chart.py spec.json -o quad.pptx --render [DIR]  # ... + PDF and PNGs (render_office.sh)
    python3 scripts/quad_chart.py --check existing.pptx                  # check a quad built or edited by hand
    python3 scripts/quad_chart.py --list-fields [nasa|status|weekly]     # spec fields, targets and budgets

Options: --allow-placeholders (build a draft that still contains [bracketed] unknowns; reported as warnings),
--strict (warnings fail too), --quiet (only warnings, failures and the result line), --dpi N (render).

Exit status: 0 pass; 1 a check failed (the .pptx is still written so you can render and see the problem) or
the spec was refused (missing required field, [placeholder] without --allow-placeholders, a banned string);
2 usage or file error.

The spec (JSON; one quad, or {"quads": [...]} for a multi-slide deck; paths are relative to the spec file)
---------------------------------------------------------------------------------------------------------
Common keys: "type" ("nasa" | "status" | "weekly"), "title", "notes" (speaker notes: string or list of
paragraphs; GSFC rule 12 puts methods detail and acronym definitions here), "example_tag" (true or a string:
marks demo content visibly, "EXAMPLE - illustrative content" chip on NASA quads, "Example" footer marking on the
others). Text fields take a string, a list of strings (one paragraph or bullet each) or items
{"text": "...", "level": 1, "url": "https://..."}; **double asterisks** make a bold run.

nasa    title, citation (string, or {"authors","year","venue","doi"}), pi, program ("NASA EPSCoR"), date
        ("2026" or "September 2026": AP spells out a month without a day) or meta (list of lines), background, analysis, figures (1-2), caption, significance,
        grant ("80NSSC26K1234"). The acknowledgement is NASA's sentence, built from grant and program.
status  title, meta (3 lines) or pi/poc/sponsor/award/period, objective, image {"path","alt"} (optional; without
        it the objective text takes the full quadrant), approach, milestones (1-5 rows of
        {"milestone","target","status"}; status = complete | on-track | at-risk | late | not-started), impact,
        sponsor_logo {"path","alt"} (optional, NSF), marking ("Internal"), version_date (default today).
weekly  title, meta (3 lines) or week_of/student/advisor, accomplishments, path_forward, risks, week_notes,
        media ({"title","label","url"}), marking, version_date.

A figure is {"image": "fig.png", "alt": "..."}, {"chart": {...}, "alt": "..."} or {"callout": {...}}; a chart is
    {"type": "bar" | "column" | "line", "categories": ["50 ms", "250 ms", "500 ms"],
     "series": [{"name": "...", "values": [4, 15, 31]}], "value_title": "Fewer grasp failures (%)",
     "category_title": "Delay", "label_format": "0\\"%\\"", "highlight": [2], "min": 0, "max": 40,
     "value_ticks": false, "flag": "Illustrative data"}
Native charts are set in Arial 14 pt black (NASA figure text); axis titles carry the units.
A callout is the figure for a result that is one number (a chart with a single bar is a number, not a comparison):
    {"callout": {"value": "31%", "label": "fewer grasp failures", "context": "adaptive vs fixed haptic gain at
     500 ms, n = 24", "flag": "Illustrative data", "panel": true}, "alt": "..."}
The value is set large (Arial bold navy on the NASA layout, Roboto Slab bold navy on the brand layouts), the
label and context 14 pt (black on NASA; Source Sans 3 ink and slate on the brand layouts), on a mist panel that
fills the slot ("panel": false for none). Alone it takes both Results slots (value at the left, label and
context beside it); next to a chart or an image it takes the narrower left slot and stacks. A status quad
takes a callout in place of "image" ({"image": {"callout": {...}}}).

Needs python-pptx, lxml and fontTools (scripts/requirements.txt, core block). Text fit is measured with the
real font files: the brand fonts from assets/fonts/, and Arial from the system (or Liberation Sans, which has
Arial's metrics). Without Arial the widths are estimated and the report says so.
"""
from __future__ import annotations

import sys

sys.dont_write_bytecode = True

import argparse
import copy
import datetime as _dt
import json
import math
import os
import re
import subprocess
from pathlib import Path

try:
    from lxml import etree
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION, XL_TICK_MARK
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches, Pt
except ImportError as exc:  # pragma: no cover
    sys.stderr.write(f"quad_chart.py needs python-pptx and lxml ({exc}).\n"
                     "  python3 -m venv .venv && .venv/bin/pip install -r scripts/requirements.txt\n")
    sys.exit(2)

__version__ = "1.0"

SKILL = Path(__file__).resolve().parent.parent
TEMPLATE = SKILL / "assets" / "templates" / "ATR-Quad-Chart-Template.pptx"
LAYOUTS_JSON = SKILL / "assets" / "templates" / "quad-layouts.json"
FONT_DIR = SKILL / "assets" / "fonts"
BRAND_CHECK = SKILL / "scripts" / "brand_check.py"
RENDER = SKILL / "scripts" / "render_office.sh"

EMU_IN = 914400
TYPES = {  # type -> layout name, index of its showcase slide in the template
    "nasa": ("ATR - NASA Research Quad", 0),
    "status": ("ATR - Project Status Quad", 1),
    "weekly": ("ATR - Weekly Summary", 2),
}
SAMPLE_SLIDE = 3  # template slide 4: the NASA sample (sample-tag chip, "Illustrative data" flag)
NASA_HEADINGS = ["Background or Science Question", "Analysis", "Results", "Significance", "Acknowledgements"]
NASA_ACK = ("This work was supported by the National Aeronautics and Space Administration under "
            "Grant/Contract/Agreement No. {grant} and was part of the NASA {program} program.")
ACK_RE = re.compile(r"This work was supported by the National Aeronautics and Space Administration under "
                    r"Grant/Contract/Agreement No\. (?P<grant>.+?) and was part of the (?P<program>.+?) program\.")
GSFC_URL = "https://cce-signin.gsfc.nasa.gov/online_help_docs/quadchart_help.html"

NAVY, GOLD, INK, SLATE, BLACK, MIST = "003976", "EFAB00", "1B2533", "4A5868", "000000", "F3F6FA"
AXIS, GRID = "B9C2CD", "E6EBF1"
EMPTY_PANEL_FRACTION = 0.40  # a bullet panel under this share of its character and line budgets looks empty
# Callout type profiles: (value face, value bold, value color), (label face, bold, color), (context face, bold,
# color), line pitch in inches for the 14 pt lines. NASA: Arial, value navy, figure text black (GSFC rule 7).
CALLOUT_PROFILES = {
    "nasa": {"value": ("Arial", True, NAVY), "label": ("Arial", True, BLACK), "context": ("Arial", False, BLACK),
             "pitch": 16 / 72},
    "brand": {"value": ("Roboto Slab", True, NAVY), "label": ("Source Sans 3 Semibold", False, INK),
              "context": ("Source Sans 3", False, SLATE), "pitch": 18 / 72},
}
CALLOUT_VALUE_SIZES = [60, 54, 48, 44, 40, 36, 32, 28, 24]
CATEGORICAL = ["003976", "EFAB00", "2C8ECD", "B63B35", "059583", "DC7533", "7D4DAD", "47A34E"]
# a status word may be typed in several ways
STATUS_ALIASES = {"complete": "complete", "completed": "complete", "done": "complete",
                  "on-track": "on-track", "on track": "on-track", "ontrack": "on-track",
                  "at-risk": "at-risk", "at risk": "at-risk", "risk": "at-risk",
                  "late": "late", "behind": "late", "overdue": "late",
                  "not-started": "not-started", "not started": "not-started", "planned": "not-started"}
# strings that brand_check.py treats as errors: refuse them before building (see references/voice-and-copy.md)
BANNED = [(re.compile(r"@atr_kent\b", re.I), "@atr_kent does not exist; use @atrlab_kent (X) or @atr_lab (Instagram)"),
          (re.compile(r"College of Arts (and|&) Sciences", re.I), "Computer Science is in the College of Sciences and Humanities"),
          (re.compile(r"Advanced Tele-Robotics", re.I), "the name is Advanced Telerobotics Research Lab")]
PLACEHOLDER_RE = re.compile(r"\[(?!\d+\])[^\[\]\n]{1,80}\]")  # [bracketed text]; citations like [3] are ignored
MAX_CHARS_FAIL = 1.15  # over 115 % of a character budget fails; over 100 % warns

# Spec fields per type: (spec key, placeholder key or target, what it is, required). --list-fields prints this.
FIELDS = {
    "nasa": [
        ("title", "title", "attention-grabbing finding, 2 lines max (Arial bold 20)", True),
        ("citation", "citation", "short citation + DOI, one line: Author et al. (Year), Journal, doi:10.xxxx/...", True),
        ("pi / date  or  meta", "meta", "header right, 2 lines: PI; NASA program  ·  date", False),
        ("background", "q1", "1  Background or Science Question: why it matters, the question (bullets)", True),
        ("analysis", "q2", "2  Analysis: what was done and learned + 'NASA resources used: ...' (bullets)", True),
        ("figures", "q3fig1/q3fig2", "3  Results: 1-2 figures (image + alt, native chart from data, or a callout: one big number + label)", True),
        ("caption", "q3cap", "3  Results caption, one line, black: quantity, units, n", False),
        ("significance", "q4", "4  Significance: 1-3 accomplishments and why they matter (bullets)", True),
        ("grant + program", "ack", "5  Acknowledgements: NASA's exact sentence, built for you (program also shows in the header)", True),
        ("notes", "notes pane", "methods detail, acronyms, full reference (GSFC rule 12)", False),
    ],
    "status": [
        ("title", "title", "project title as a goal, 2 lines max (bold 24)", True),
        ("meta  or  pi / poc / sponsor / award / period", "meta", "3 right-aligned lines", False),
        ("objective", "q1", "1  Objective and Description (bullets)", True),
        ("image", "q1img", "1  photo, render or diagram of the system {path, alt}, or a callout {callout: {value, label}}; optional", False),
        ("approach", "q2", "2  Technical Approach (bullets)", True),
        ("milestones", "Milestone table", "3  1-5 rows {milestone, target, status}", True),
        ("impact", "q4", "4  Impact and Deliverables; last line transition or next steps (bullets)", True),
        ("sponsor_logo", "header right", "NSF full-color logo {path, alt}; optional", False),
        ("marking / version_date", "footer_right", "e.g. Internal  ·  2026-09-29 (25 characters)", False),
        ("notes", "notes pane", "speaker notes", False),
    ],
    "weekly": [
        ("title", "title", "[Project or student]: weekly summary, 2 lines max", True),
        ("meta  or  week_of / student / advisor", "meta", "3 right-aligned lines", False),
        ("accomplishments", "acc", "1  outcomes finished this week (bullets)", True),
        ("path_forward", "path", "2  next steps with dates, decisions needed (bullets)", True),
        ("risks", "risks", "3  blockers and asks (bullets)", True),
        ("week_notes", "notes", "4  reading, ideas, questions (bullets)", False),
        ("media", "media", "5  {title, label, url}: the label becomes the link", False),
        ("marking / version_date", "footer_right", "e.g. Internal  ·  2026-09-29", False),
        ("notes", "notes pane", "speaker notes", False),
    ],
}


class SpecError(Exception):
    """The spec cannot be built as written (missing field, placeholder, banned string, bad value)."""


# ----------------------------------------------------------------------------------------------------------
# Findings
# ----------------------------------------------------------------------------------------------------------
class Findings:
    def __init__(self):
        self.items = []  # (severity, where, message)

    def add(self, sev, where, msg):
        self.items.append((sev, where, msg))

    def fail(self, where, msg): self.add("FAIL", where, msg)
    def warn(self, where, msg): self.add("warn", where, msg)
    def info(self, where, msg): self.add("info", where, msg)

    def count(self, sev):
        return sum(1 for s, _, _ in self.items if s == sev)

    def print(self, quiet=False):
        order = {"FAIL": 0, "warn": 1, "info": 2}
        for sev, where, msg in sorted(self.items, key=lambda x: order[x[0]]):
            if quiet and sev == "info":
                continue
            print(f"  {sev:4}  {where}: {msg}")


# ----------------------------------------------------------------------------------------------------------
# Font metrics: real advance widths (fontTools), else average advances
# ----------------------------------------------------------------------------------------------------------
_ARIAL = {
    False: ["/System/Library/Fonts/Supplemental/Arial.ttf", "/Library/Fonts/Arial.ttf", "~/Library/Fonts/Arial.ttf",
            "C:/Windows/Fonts/arial.ttf", "/usr/share/fonts/truetype/msttcorefonts/Arial.ttf",
            "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
            "/usr/share/fonts/liberation-sans/LiberationSans-Regular.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"],
    True: ["/System/Library/Fonts/Supplemental/Arial Bold.ttf", "/Library/Fonts/Arial Bold.ttf", "~/Library/Fonts/Arial Bold.ttf",
           "C:/Windows/Fonts/arialbd.ttf", "/usr/share/fonts/truetype/msttcorefonts/Arial_Bold.ttf",
           "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
           "/usr/share/fonts/liberation-sans/LiberationSans-Bold.ttf",
           "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"],
}
_AVG_EM = {"Arial": (0.447, 0.476), "Source Sans 3": (0.43, 0.46), "Source Code Pro": (0.60, 0.60),
           "Roboto Slab": (0.50, 0.53)}


class Metrics:
    def __init__(self):
        self.cache = {}
        self.estimated = set()

    def _path(self, face, bold):
        f = face.replace(" Semibold", "").replace(" SemiBold", "").replace(" Black", "").replace(" Light", "")
        if f == "Arial":
            extra = [os.path.join(d, n) for d in os.environ.get("ATR_FONT_DIRS", "").split(os.pathsep) if d
                     for n in (["Arial Bold.ttf", "arialbd.ttf", "LiberationSans-Bold.ttf"] if bold
                               else ["Arial.ttf", "arial.ttf", "LiberationSans-Regular.ttf"])]
            for p in extra + _ARIAL[bold]:
                p = os.path.expanduser(p)
                if os.path.exists(p):
                    return p
            return None
        weight = "Bold" if bold else "Regular"
        if face.endswith("Semibold") or face.endswith("SemiBold"):
            weight = "Semibold"
        if face.endswith("Black"):
            weight = "Black"
        if face.endswith("Light"):
            weight = "Light"
        files = {"Source Sans 3": FONT_DIR / "source-sans-3" / f"SourceSans3-{weight}.ttf",
                 "Source Code Pro": FONT_DIR / "source-code-pro" / f"SourceCodePro-{weight}.ttf",
                 "Roboto Slab": FONT_DIR / "roboto-slab" / f"RobotoSlab-{'SemiBold' if weight == 'Semibold' else weight}.ttf"}
        p = files.get(f)
        return str(p) if p and p.exists() else None

    def _font(self, face, bold):
        key = (face, bool(bold))
        if key not in self.cache:
            p = self._path(face, bold)
            font = None
            if p:
                try:
                    from fontTools.ttLib import TTFont
                    tt = TTFont(p, lazy=True)
                    font = (tt.getBestCmap(), tt["hmtx"].metrics, tt["head"].unitsPerEm)
                except Exception:  # fontTools missing or unreadable font: estimate
                    font = None
            if font is None:
                self.estimated.add(face)
            self.cache[key] = font
        return self.cache[key]

    def width_in(self, text, face, bold, size_pt):
        base = next((k for k in _AVG_EM if face.startswith(k)), "Arial")
        f = self._font(base if base != "Source Sans 3" else face, bold)
        if f is None:
            return len(text) * _AVG_EM[base][1 if bold else 0] * size_pt / 72
        cmap, hmtx, upm = f
        avg = _AVG_EM[base][1 if bold else 0] * upm
        units = 0
        for ch in text:
            g = cmap.get(ord(ch))
            units += hmtx[g][0] if g in hmtx else avg
        return units / upm * size_pt / 72

    def wrap(self, text, face, bold, size_pt, width_in):
        """Greedy word wrap the way PowerPoint does it (break at spaces; an over-long word breaks by characters)."""
        lines, cur = [], ""
        for word in text.split(" "):
            trial = word if not cur else cur + " " + word
            if self.width_in(trial, face, bold, size_pt) <= width_in + 1e-6:
                cur = trial
                continue
            if cur:
                lines.append(cur)
            cur = word
            while self.width_in(cur, face, bold, size_pt) > width_in and len(cur) > 1:  # break a long token
                n = len(cur)
                while n > 1 and self.width_in(cur[:n], face, bold, size_pt) > width_in:
                    n -= 1
                lines.append(cur[:n])
                cur = cur[n:]
        lines.append(cur)
        return lines


METRICS = Metrics()


# ----------------------------------------------------------------------------------------------------------
# Spec helpers
# ----------------------------------------------------------------------------------------------------------
class Para:
    def __init__(self, text, level=0, url=None):
        self.text, self.level, self.url = str(text), int(level), url

    def segments(self):
        """'**bold**' markup -> [(text, bold)]."""
        parts = re.split(r"(\*\*[^*]+\*\*)", self.text)
        return [(p[2:-2], True) if p.startswith("**") and p.endswith("**") else (p, False) for p in parts if p]

    def plain(self):
        return "".join(t for t, _ in self.segments())


def paras(value, where):
    if value is None:
        return []
    if isinstance(value, str):
        value = [v for v in value.split("\n")]
    if not isinstance(value, list):
        raise SpecError(f"{where}: expected a string or a list, got {type(value).__name__}")
    out = []
    for it in value:
        if isinstance(it, str):
            out.append(Para(it))
        elif isinstance(it, dict) and "text" in it:
            lvl = int(it.get("level", 0))
            if not 0 <= lvl <= 4:
                raise SpecError(f"{where}: level must be 0-4, got {lvl}")
            out.append(Para(it["text"], lvl, it.get("url")))
        else:
            raise SpecError(f"{where}: each item is a string or {{\"text\": ..., \"level\": n}}")
    return [p for p in out if p.plain().strip() or len(out) == 1]


def all_strings(obj, path=""):
    """Every string in a spec, with its JSON path (for the placeholder and banned-string scans)."""
    if isinstance(obj, str):
        yield path, obj
    elif isinstance(obj, dict):
        for k, v in obj.items():
            if not str(k).startswith("_"):
                yield from all_strings(v, f"{path}.{k}" if path else str(k))
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            yield from all_strings(v, f"{path}[{i}]")


def scan_spec(spec, allow_placeholders, findings, label):
    for path, s in all_strings(spec):
        if path.endswith(("path", "image", "url")) or ".path" in path:
            continue
        for rx, why in BANNED:
            if rx.search(s):
                raise SpecError(f"{label} {path}: {why}")
        m = PLACEHOLDER_RE.search(s)
        if m and not allow_placeholders:  # with --allow-placeholders the slide and notes check lists each one
            raise SpecError(f'{label} {path}: placeholder "{m.group(0)}" (replace it, or pass --allow-placeholders for a draft)')


def need(spec, key, label):
    v = spec.get(key)
    if v is None or (isinstance(v, (str, list)) and not v):
        raise SpecError(f"{label}: \"{key}\" is required for a {spec.get('type')} quad")
    return v


def resolve(spec_dir, p):
    q = Path(os.path.expanduser(str(p)))
    return q if q.is_absolute() else (spec_dir / q)


# ----------------------------------------------------------------------------------------------------------
# XML helpers
# ----------------------------------------------------------------------------------------------------------
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"a": A, "p": P, "c": C, "r": R}


def set_xfrm(el, x, y, w, h):
    """Set a shape's position and size (inches) on its spPr/xfrm, creating the xfrm if needed."""
    sppr = el.find(qn("p:spPr"))
    xfrm = sppr.find(qn("a:xfrm"))
    if xfrm is None:
        xfrm = etree.Element(qn("a:xfrm"))
        sppr.insert(0, xfrm)
        etree.SubElement(xfrm, qn("a:off"))
        etree.SubElement(xfrm, qn("a:ext"))
    xfrm.find(qn("a:off")).set("x", str(int(round(x * EMU_IN))))
    xfrm.find(qn("a:off")).set("y", str(int(round(y * EMU_IN))))
    xfrm.find(qn("a:ext")).set("cx", str(int(round(w * EMU_IN))))
    xfrm.find(qn("a:ext")).set("cy", str(int(round(h * EMU_IN))))


def renumber_ids(slide):
    n = 2
    for el in slide.shapes._spTree.iter(qn("p:cNvPr")):
        if el.getparent().getparent() is slide.shapes._spTree and el.getparent().tag == qn("p:nvGrpSpPr"):
            el.set("id", "1")
            continue
        el.set("id", str(n))
        n += 1


def layout_lvl_ppr(layout_el, level):
    if layout_el is None:
        return None
    lst = layout_el.find(".//" + qn("a:lstStyle"))
    return lst.find(qn(f"a:lvl{level + 1}pPr")) if lst is not None else None


def write_paragraphs(shape_el, items, layout_el=None):
    """Replace the text of a placeholder, keeping the showcase slide's explicit paragraph and run formatting (the
    first paragraph is the prototype). Level > 0 takes marL, indent and bullet character from the layout list style."""
    tx = shape_el.find(qn("p:txBody"))
    old = tx.findall(qn("a:p"))
    proto = old[0]
    ppr0 = proto.find(qn("a:pPr"))
    r0 = proto.find(qn("a:r"))
    rpr0 = r0.find(qn("a:rPr")) if r0 is not None else None
    end0 = proto.find(qn("a:endParaRPr"))
    for p in old:
        tx.remove(p)
    for it in items:
        p = etree.SubElement(tx, qn("a:p"))
        if ppr0 is not None:
            ppr = copy.deepcopy(ppr0)
            if it.level:
                ppr.set("lvl", str(it.level))
                lv = layout_lvl_ppr(layout_el, it.level)
                if lv is not None:
                    for att in ("marL", "indent"):
                        if lv.get(att) is not None:
                            ppr.set(att, lv.get(att))
                    bu, cur = lv.find(qn("a:buChar")), ppr.find(qn("a:buChar"))
                    if bu is not None and cur is not None:
                        cur.set("char", bu.get("char"))
            p.append(ppr)
        for text, bold in it.segments():
            r = etree.SubElement(p, qn("a:r"))
            rpr = copy.deepcopy(rpr0) if rpr0 is not None else etree.Element(qn("a:rPr"), lang="en-US")
            if bold:
                rpr.set("b", "1")
            r.append(rpr)
            t = etree.SubElement(r, qn("a:t"))
            t.text = text
        if end0 is not None:
            p.append(copy.deepcopy(end0))


def set_run_text(tc, texts):
    """Put texts into the first len(texts) runs of a table cell's first paragraph (extra runs are removed)."""
    p = tc.find(".//" + qn("a:p"))
    runs = p.findall(qn("a:r"))
    for i, r in enumerate(runs):
        if i < len(texts):
            r.find(qn("a:t")).text = texts[i]
        else:
            p.remove(r)
    return runs[:len(texts)]


def run_color(r, hexcol):
    rpr = r.find(qn("a:rPr"))
    fill = rpr.find(qn("a:solidFill"))
    if fill is None:
        fill = etree.SubElement(rpr, qn("a:solidFill"))
        rpr.insert(0, fill)
    for c in list(fill):
        fill.remove(c)
    etree.SubElement(fill, qn("a:srgbClr"), val=hexcol)


# ----------------------------------------------------------------------------------------------------------
# Building
# ----------------------------------------------------------------------------------------------------------
class Builder:
    def __init__(self, lay, allow_placeholders=False):
        self.lay = lay
        self.allow = allow_placeholders
        self.prs = Presentation(str(TEMPLATE))
        self.template_slides = list(self.prs.slides)
        self.original_ids = list(self.prs.slides._sldIdLst)
        self.findings = Findings()
        self.layouts = {L["name"]: L for L in lay["layouts"]}

    # -- slide duplication (the template's own guidance: duplicate the showcase slide, never New Slide) --------
    def duplicate(self, src):
        new = self.prs.slides.add_slide(src.slide_layout)
        tree = new.shapes._spTree
        for el in list(tree)[2:]:
            tree.remove(el)
        for el in list(src.shapes._spTree)[2:]:
            dup = copy.deepcopy(el)
            if etree.ElementBase.xpath(dup, ".//@r:embed|.//@r:id|.//@r:link", namespaces=NS):
                raise RuntimeError("showcase slide shape carries a relationship; the template changed, update quad_chart.py")
            tree.append(dup)
        return new

    def sample_shape(self, name):
        for sh in self.template_slides[SAMPLE_SLIDE].shapes:
            if sh.name == name:
                return copy.deepcopy(sh._element)
        raise RuntimeError(f'template slide 4 has no shape "{name}"; the template changed, update quad_chart.py')

    def build(self, quads, spec_dir):
        for n, q in enumerate(quads, 1):
            label = f"quad {n}" if len(quads) > 1 else "spec"
            t = q.get("type")
            if t not in TYPES:
                raise SpecError(f'{label}: "type" must be one of {", ".join(TYPES)} (got {t!r})')
            scan_spec(q, self.allow, self.findings, label)
            name, idx = TYPES[t]
            slide = self.duplicate(self.template_slides[idx])
            ctx = SlideCtx(self, slide, self.layouts[name], q, spec_dir, label)
            getattr(ctx, f"build_{t}")()
            ctx.delete_unfilled()
            ctx.write_notes()
            renumber_ids(slide)
        # drop the template's four slides (and with them the sample chart and its workbook)
        lst = self.prs.slides._sldIdLst
        for sid in self.original_ids:
            self.prs.part.drop_rel(sid.rId)
            lst.remove(sid)
        self.prs.part.rename_slide_parts([sid.rId for sid in lst])  # slide1.xml ... in deck order
        first = quads[0]
        self.prs.core_properties.title = str(first.get("title", ""))[:250]
        self.prs.core_properties.subject = "ATR Lab quad chart"
        return self.prs


class SlideCtx:
    def __init__(self, builder, slide, L, spec, spec_dir, label):
        self.b, self.slide, self.L, self.q, self.dir, self.label = builder, slide, L, spec, spec_dir, label
        self.rec = {p["key"]: p for p in L["placeholders"]}
        self.layout_el = {ph.placeholder_format.idx: ph._element for ph in slide.slide_layout.placeholders}
        self.filled = set()
        self.F = builder.findings

    # -- placeholder access ------------------------------------------------------------------------------
    def ph(self, key):
        idx = self.rec[key]["idx"]
        for sh in self.slide.placeholders:
            if sh.placeholder_format.idx == idx:
                return sh
        return None

    def fill(self, key, items):
        sh = self.ph(key)
        if sh is None:
            raise RuntimeError(f"placeholder {key} missing on the showcase slide")
        if not items:
            return
        write_paragraphs(sh._element, items, self.layout_el.get(self.rec[key]["idx"]))
        self.filled.add(key)
        return sh

    def delete(self, key):
        sh = self.ph(key)
        if sh is not None:
            sh._element.getparent().remove(sh._element)

    def delete_unfilled(self):
        for key, rec in self.rec.items():
            if key not in self.filled:
                self.delete(key)

    def write_notes(self):
        notes = paras(self.q.get("notes"), f"{self.label} notes")
        lines = [p.plain() for p in notes]
        if self.q.get("example_tag"):
            lines.insert(0, "EXAMPLE: built from an example spec. The content is illustrative, not an ATR Lab result; "
                            "do not cite it or submit it.")
        if lines:
            self.slide.notes_slide.notes_text_frame.text = "\n".join(lines)

    # -- shared pieces -----------------------------------------------------------------------------------
    def meta_lines(self, keys_fmt):
        if self.q.get("meta") is not None:
            return paras(self.q["meta"], f"{self.label} meta")
        out = []
        for fmt, keys in keys_fmt:
            vals = [self.q.get(k) for k in keys]
            if all(v not in (None, "") for v in vals):
                out.append(Para(fmt.format(*vals)))
            elif any(v not in (None, "") for v in vals):
                out.append(Para("  ·  ".join(str(v) for v in vals if v not in (None, ""))))
        return out

    def marking(self):
        date = self.q.get("version_date") or _dt.date.today().isoformat()
        mark = self.q.get("marking")
        if mark is None:
            mark = "Example" if self.q.get("example_tag") else "Internal"
        text = mark if "·" in str(mark) else f"{mark} · {date}"
        self.fill("footer_right", [Para(text)])

    def insert_after(self, anchor_el, el):
        anchor_el.addnext(el)

    def picture_into(self, key, path, alt, box, name):
        """Insert an image into a picture placeholder, uncropped and fitted inside box (x, y, w, h in inches)."""
        sh = self.ph(key)
        path = resolve(self.dir, path)
        if not path.exists():
            raise SpecError(f"{self.label}: image not found: {path}")
        if not alt or not str(alt).strip():
            raise SpecError(f"{self.label}: image {path.name} needs \"alt\" text (what the figure shows and its main values)")
        pic = sh.insert_picture(str(path))
        pic.crop_left = pic.crop_right = pic.crop_top = pic.crop_bottom = 0
        iw, ih = pic.image.size
        x, y, w, h = box
        s = min(w / iw, h / ih)
        nw, nh = iw * s, ih * s
        pic.left, pic.top = Emu(int((x + (w - nw) / 2) * EMU_IN)), Emu(int((y + (h - nh) / 2) * EMU_IN))
        pic.width, pic.height = Emu(int(nw * EMU_IN)), Emu(int(nh * EMU_IN))
        cnv = pic._element.find(".//" + qn("p:cNvPr"))
        cnv.set("descr", str(alt).strip())
        cnv.set("name", name)
        dpi = iw / nw
        if dpi < 150:
            self.F.warn(f"{self.label} {name}", f"{path.name} is {iw} x {ih} px, {dpi:.0f} ppi at its placed size; "
                                                f"export at 300 ppi at {nw:.2f} x {nh:.2f} in so 14 pt text stays sharp")
        self.filled.add(key)
        return pic

    # -- NASA ------------------------------------------------------------------------------------------------
    def build_nasa(self):
        q, lab = self.q, self.label
        title = self.fill("title", paras(need(q, "title", lab), "title"))
        cit = need(q, "citation", lab)
        if isinstance(cit, dict):
            for k in ("authors", "year", "venue", "doi"):
                if not cit.get(k):
                    raise SpecError(f'{lab}: citation needs "{k}" (or give the citation as one string)')
            doi = re.sub(r"^(https?://(dx\.)?doi\.org/|doi:\s*)", "", str(cit["doi"]).strip(), flags=re.I)
            cit = f'{cit["authors"]} ({cit["year"]}), {cit["venue"]}, doi:{doi}'
        self.fill("citation", [Para(cit)])
        grant = str(need(q, "grant", lab)).strip()
        program = str(need(q, "program", lab)).strip()
        program_core = re.sub(r"^NASA\s+", "", program)
        q.setdefault("_program_display", f"NASA {program_core}")
        meta = self.meta_lines([("{}", ["pi"]), ("{}  ·  {}", ["_program_display", "date"])])
        if not q.get("meta") and not q.get("pi"):
            self.F.warn(f"{lab} meta", "no PI line (give \"pi\")")
        if q.get("example_tag"):
            tag_text = q["example_tag"] if isinstance(q["example_tag"], str) else "EXAMPLE - illustrative content"
            chip, tag = self.b.sample_shape("Sample tag chip"), self.b.sample_shape("Sample tag")
            tag.find(".//" + qn("a:t")).text = tag_text
            title._element.addnext(chip)
            chip.addnext(tag)
            meta_el = self.ph("meta")._element
            set_xfrm(meta_el, 6.45, 0.54, 3.05, 0.23)  # as on the template's sample slide
            if len(meta) > 1:
                self.F.info(f"{lab} meta", f'example tag takes the first meta line; "{meta[0].plain()}" moved to the notes')
                q.setdefault("notes", [])
                if isinstance(q["notes"], str):
                    q["notes"] = [q["notes"]]
                q["notes"] = list(q["notes"]) + [f"Header: {meta[0].plain()}"]
                meta = meta[-1:]
        self.fill("meta", meta)
        self.fill("q1", paras(need(q, "background", lab), "background"))
        self.fill("q2", paras(need(q, "analysis", lab), "analysis"))
        self.nasa_figures(need(q, "figures", lab))
        if q.get("caption"):
            self.fill("q3cap", paras(q["caption"], "caption"))
        else:
            self.F.warn(f"{lab} caption", "no caption: name the quantity, units and n under the figure")
        self.fill("q4", paras(need(q, "significance", lab), "significance"))
        self.fill("ack", [Para(NASA_ACK.format(grant=grant, program=program_core))])

    def nasa_figures(self, figs):
        lab = self.label
        if not isinstance(figs, list) or not 1 <= len(figs) <= 2:
            raise SpecError(f"{lab}: figures must be a list of 1 or 2 figures (GSFC rule 9)")
        r1, r2 = self.rec["q3fig1"]["box"], self.rec["q3fig2"]["box"]
        for i, fig in enumerate(figs):
            if not isinstance(fig, dict) or not any(fig.get(k) for k in ("image", "chart", "callout")):
                raise SpecError(f"{lab}: figure {i + 1} must be {{\"image\": ...}}, {{\"chart\": ...}} or {{\"callout\": ...}}")
        kinds = ["callout" if f.get("callout") else "chart" if f.get("chart") else "image" for f in figs]
        full_w = r2["x"] + r2["w"] - r1["x"]
        gap = r2["x"] - (r1["x"] + r1["w"])
        if len(figs) == 1:
            boxes = [(r1["x"], r1["y"], full_w, r1["h"])]
            self.delete("q3fig2")
        elif kinds.count("callout") == 1:
            # a callout beside a chart or an image: the number needs less width than a plot with 14 pt axis text
            cw = round((full_w - gap) * 0.45, 3)
            ow = full_w - gap - cw
            w1 = cw if kinds[0] == "callout" else ow
            boxes = [(r1["x"], r1["y"], w1, r1["h"]), (r1["x"] + w1 + gap, r2["y"], full_w - gap - w1, r2["h"])]
        else:
            boxes = [(r1["x"], r1["y"], r1["w"], r1["h"]), (r2["x"], r2["y"], r2["w"], r2["h"])]
        for i, (fig, box) in enumerate(zip(figs, boxes)):
            key = f"q3fig{i + 1}"
            name = f"Figure {i + 1}"
            if fig.get("image"):
                self.picture_into(key, fig["image"], fig.get("alt"), box, name)
                self.F.info(f"{lab} {name}", "image figure: check by eye that axes are labeled with units, any color "
                                             "bar is labeled, and its text is Arial >= 14 pt black at the placed size")
            elif fig.get("chart"):
                x, y, w, h = box
                if len(figs) == 1:  # the chart may use the gap between the slots and the caption (as the sample does)
                    y, h = y - 0.03, h + 0.06
                if w < 3:
                    beside_callout = "callout" in kinds
                    (self.F.info if beside_callout else self.F.warn)(
                        f"{lab} {name}", f"a native chart in a {w:.2f} in slot leaves little room for 14 pt axis text: "
                        + ("keep it to two or three categories" if beside_callout else "one chart across both slots reads better"))
                frame = add_chart(self.slide, fig["chart"], (x, y, w, h), fig.get("alt"), name, self.F, lab)
                slot = self.ph(key)._element
                slot.addprevious(frame)
                flag = fig["chart"].get("flag")
                if flag:
                    frame.addnext(self.flag_shape(flag, x, y, w))
                slot.getparent().remove(slot)
                self.filled.add(key)
            else:
                self.callout_into(key, fig, box, name, "nasa")

    def flag_shape(self, flag, x, y, w):
        """The template's 'Illustrative data' flag, placed at the top right of a figure box."""
        fl = self.b.sample_shape("Illustrative data flag")
        fl.find(".//" + qn("a:t")).text = str(flag)
        fw = min(w, METRICS.width_in(str(flag), "Arial", False, 14) + 0.05)
        set_xfrm(fl, x + w - fw, y + 0.02, fw, 0.23)
        return fl

    def callout_into(self, key, fig, box, name, profile):
        """Replace a picture slot with a callout figure (big number + label + context) built by add_callout."""
        co = fig["callout"]
        if not isinstance(co, dict):
            raise SpecError(f"{self.label} {name}: \"callout\" must be an object with \"value\" and \"label\"")
        els = add_callout(self.slide, co, box, fig.get("alt"), name, self.F, self.label, profile)
        slot = self.ph(key)._element
        for el in els:
            slot.addprevious(el)
        flag = co.get("flag")
        if flag:
            x, y, w, h = box
            slot.addprevious(self.flag_shape(flag, x, y, w))
        slot.getparent().remove(slot)
        self.filled.add(key)

    # -- Project status ----------------------------------------------------------------------------------
    def build_status(self):
        q, lab = self.q, self.label
        self.fill("title", paras(need(q, "title", lab), "title"))
        meta = self.meta_lines([("PI: {}  ·  POC: {}", ["pi", "poc"]), ("{}  ·  Award {}", ["sponsor", "award"]),
                                ("{}", ["period"])])
        meta_sh = self.fill("meta", meta)
        if q.get("sponsor_logo"):
            lg = q["sponsor_logo"]
            path = resolve(self.dir, lg.get("path") or lg.get("image") or "")
            if not path.is_file():
                raise SpecError(f"{lab}: sponsor_logo image not found: {path}")
            if not lg.get("alt"):
                raise SpecError(f'{lab}: sponsor_logo needs "alt" (e.g. "U.S. National Science Foundation logo")')
            box = self.rec["sponsor"]["box"]
            pic = self.slide.shapes.add_picture(str(path), Inches(box["x"]), Inches(box["y"]))
            iw, ih = pic.image.size
            s = min(box["w"] / iw, box["h"] / ih)
            pic.width, pic.height = Emu(int(iw * s * EMU_IN)), Emu(int(ih * s * EMU_IN))
            pic.left = Emu(int((box["x"] + box["w"] - iw * s) * EMU_IN))
            cnv = pic._element.find(".//" + qn("p:cNvPr"))
            cnv.set("descr", lg["alt"])
            cnv.set("name", "Sponsor logo")
            meta_sh._element.addnext(pic._element)
            mb = self.rec["meta"]["box"]
            set_xfrm(meta_sh._element, mb["x"], mb["y"], mb["w"], mb["h"])  # the layout's 2.70 in beside the slot
            if iw * s < 0.625:
                self.F.fail(f"{lab} sponsor_logo", f"placed {iw * s:.2f} in wide; the NSF logo minimum is 0.625 in")
        self.fill("q1", paras(need(q, "objective", lab), "objective"))
        if q.get("image"):
            img = q["image"]
            if isinstance(img, str):
                img = {"path": img}
            b = self.rec["q1img"]["box"]
            if img.get("callout"):
                self.callout_into("q1img", img, (b["x"], b["y"], b["w"], b["h"]), "Objective callout", "brand")
            else:
                self.picture_into("q1img", img.get("path") or img.get("image"), img.get("alt"), (b["x"], b["y"], b["w"], b["h"]), "Objective image")
        else:
            q2 = self.rec["q2"]["box"]
            self.ph("q1").width = Inches(q2["w"])  # no picture: the objective text takes the whole quadrant
        self.fill("q2", paras(need(q, "approach", lab), "approach"))
        self.milestone_table(need(q, "milestones", lab))
        self.fill("q4", paras(need(q, "impact", lab), "impact"))
        self.marking()

    def milestone_table(self, rows):
        lab = self.label
        if not isinstance(rows, list) or not 1 <= len(rows) <= 5:
            raise SpecError(f"{lab}: milestones must be a list of 1 to 5 rows (more do not fit above the footer; "
                            "split the schedule or use a Gantt figure)")
        glyphs = self.b.lay["milestone_glyphs"]
        frame = next((sh for sh in self.slide.shapes if sh.name == "Milestone table"), None)
        tbl = frame._element.find(".//" + qn("a:tbl"))
        trs = tbl.findall(qn("a:tr"))
        mid, last = trs[1], trs[-1]
        for tr in trs[1:]:
            tbl.remove(tr)
        for i, row in enumerate(rows):
            if isinstance(row, str):
                parts = [s.strip() for s in row.split("|")]
                if len(parts) != 3:
                    raise SpecError(f'{lab}: milestone row "{row}" must be "name | target | status"')
                row = dict(zip(("milestone", "target", "status"), parts))
            state = STATUS_ALIASES.get(str(row.get("status", "")).strip().lower())
            if state is None:
                raise SpecError(f'{lab}: milestone "{row.get("milestone")}" status {row.get("status")!r}: use one of '
                                "complete, on-track, at-risk, late, not-started")
            g = glyphs[state]
            tr = copy.deepcopy(last if i == len(rows) - 1 else mid)
            cells = tr.findall(qn("a:tc"))
            set_run_text(cells[0], [str(row.get("milestone", "")).strip()])
            set_run_text(cells[1], [str(row.get("target", "")).strip()])
            label = str(row.get("label") or g["label"])
            marker, word = set_run_text(cells[2], [g["char"], "  " + label])
            run_color(marker, g["glyph_color"])
            run_color(word, INK)
            tbl.append(tr)
        ext = frame._element.find(qn("p:xfrm")).find(qn("a:ext"))
        ext.set("cy", str(sum(int(tr.get("h")) for tr in tbl.findall(qn("a:tr")))))

    # -- Weekly summary --------------------------------------------------------------------------------------
    def build_weekly(self):
        q, lab = self.q, self.label
        self.fill("title", paras(need(q, "title", lab), "title"))
        meta = self.meta_lines([("Week of {}", ["week_of"]), ("{}  ·  {}", ["student", "advisor"])])
        if meta and not q.get("meta"):
            meta.append(Para("ATR Lab  ·  Kent State University"))
        self.fill("meta", meta)
        self.fill("acc", paras(need(q, "accomplishments", lab), "accomplishments"))
        self.fill("path", paras(need(q, "path_forward", lab), "path_forward"))
        self.fill("risks", paras(need(q, "risks", lab), "risks"))
        if q.get("week_notes"):
            self.fill("notes", paras(q["week_notes"], "week_notes"))
        else:
            self.F.warn(f"{lab} week_notes", 'section 4 "Notes" is empty; its label still prints (write "None this week." if that is true)')
        media = q.get("media")
        if media:
            if isinstance(media, dict):
                items = [Para(media.get("title", ""))]
                if media.get("label"):
                    items.append(Para(media["label"], url=media.get("url")))
                elif media.get("url"):
                    self.F.warn(f"{lab} media", "a URL without a label: add a short label; the URL becomes its link")
            else:
                items = paras(media, "media")
            sh = self.fill("media", items)
            for i, it in enumerate(items):
                if it.url:
                    sh.text_frame.paragraphs[i].runs[-1].hyperlink.address = it.url
        else:
            self.F.warn(f"{lab} media", 'section 5 "Media of the week" is empty; its panel and label still print')
        self.marking()


# ----------------------------------------------------------------------------------------------------------
# Callout figures (one big number + label + context) for a result that is a single value
# ----------------------------------------------------------------------------------------------------------
def _style_run(run, face, size, bold, color):
    f = run.font
    f.name, f.size, f.bold = face, Pt(size), bool(bold)
    f.color.rgb = RGBColor.from_string(color)
    rpr = run._r.find(qn("a:rPr"))
    for tag in ("a:ea", "a:cs"):  # the same face for every script, so nothing falls back to the theme font
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", face)


def _textbox(slide, name, box, anchor, paragraphs, pitch_in=None):
    """A zero-inset text box at box (inches) whose paragraphs are [(text, face, size, bold, color), ...]."""
    x, y, w, h = box
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    body = tf._txBody.find(qn("a:bodyPr"))
    body.set("anchor", anchor)  # "t" or "ctr"
    for att in ("lIns", "tIns", "rIns", "bIns"):
        body.set(att, "0")
    body.set("wrap", "square")
    for child in list(body):  # no autofit: the sizes were measured to fit
        body.remove(child)
    for i, (text, face, size, bold, color) in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if pitch_in and size <= 14:
            p.line_spacing = Pt(pitch_in * 72)
        r = p.add_run()
        r.text = text
        _style_run(r, face, size, bold, color)
    return tb


def _lines_for(text, face, bold, size, width_in):
    return len(METRICS.wrap(text, face, bold, size, width_in)) if text else 0


_UNIT_RE = re.compile(r"(?<=\d) (?=(?:ms|s|min|h|Hz|kHz|m|cm|mm|km|kg|g|N|W|V|A|dB|px|pt|fps|participants|trials|"
                      r"students|robots)\b)")


def glue_units(text):
    """A number never loses its unit at a line break: '500 ms' becomes '500\u00a0ms' (also 'n = 24')."""
    text = _UNIT_RE.sub("\u00a0", text)
    return re.sub(r"\b([nN]) = (\d)", "\\1\u00a0=\u00a0\\2", text)


def add_callout(slide, co, box, alt, name, F, lab, profile):
    """Build a callout figure in box (x, y, w, h inches): a mist panel, the value in large type and, beside or
    under it, a bold label with an optional context line (14 pt). Returns the shape elements, detached, in
    reading order (panel, value, label), for the caller to insert where the figure slot was. The value size
    steps down from 60 pt until the text fits; when the label and context cannot fit at 14 pt the build fails
    (the NASA floor is 14 pt, so the fix is fewer words, never smaller type)."""
    prof = CALLOUT_PROFILES[profile]
    value = str(co.get("value", "")).strip()
    label = glue_units(str(co.get("label", "")).strip())
    context = glue_units(str(co.get("context", "")).strip())
    if not value:
        raise SpecError(f'{lab} {name}: a callout needs "value" (the number, e.g. "31%")')
    if not label:
        raise SpecError(f'{lab} {name}: a callout needs "label" (what the number is, e.g. "fewer grasp failures")')
    x, y, w, h = box
    inset = 0.07
    pitch = prof["pitch"]
    vf, vb, vc = prof["value"]
    lf, lb, lc = prof["label"]
    cf, cb, cc = prof["context"]
    flag_room = 0.25 if co.get("flag") else 0.0
    inner_w, inner_h = w - 2 * inset, h - 2 * inset
    side_by_side = w >= 3.0
    chosen = None
    if side_by_side:
        # value at the left; label and context in a column beside it, both centred on the panel's height
        for size in CALLOUT_VALUE_SIZES:
            vw = METRICS.width_in(value, vf, vb, size)
            vh = size * 1.15 / 72
            if vw > inner_w * 0.5 or size * 0.95 / 72 > inner_h:
                continue
            col_x = x + inset + vw + 0.22
            col_w = x + w - inset - col_x
            n = _lines_for(label, lf, lb, 14, col_w) + _lines_for(context, cf, cb, 14, col_w)
            if n * pitch <= inner_h - flag_room:
                chosen = (size, vw, vh, col_x, col_w, n)
                break
        if chosen is None:
            size = CALLOUT_VALUE_SIZES[-1]
            vw = METRICS.width_in(value, vf, vb, size)
            col_x = x + inset + vw + 0.22
            col_w = max(0.5, x + w - inset - col_x)
            n = _lines_for(label, lf, lb, 14, col_w) + _lines_for(context, cf, cb, 14, col_w)
            chosen = (size, vw, size * 1.15 / 72, col_x, col_w, n)
            F.fail(f"{lab} {name}", f"callout text needs {n} lines of 14 pt beside the value in a {h:.2f} in slot: "
                                    "shorten the label or the context (the caption can carry n and the unit)")
        size, vw, vh, col_x, col_w, n = chosen
        els_boxes = {
            "value": ((x + inset, y + inset, vw + 0.05, inner_h), "ctr"),
            "label": ((col_x, y + inset + flag_room, col_w, inner_h - flag_room), "ctr"),
        }
    else:
        # stacked: value on top, label and context under it
        n = _lines_for(label, lf, lb, 14, inner_w) + _lines_for(context, cf, cb, 14, inner_w)
        for size in CALLOUT_VALUE_SIZES:
            vw = METRICS.width_in(value, vf, vb, size)
            vh = size * 0.95 / 72  # the numerals' ascent (no descenders in a number), not the whole line box
            if vw <= inner_w and vh + 0.02 + n * pitch <= inner_h - flag_room:
                chosen = (size, vw, vh, n)
                break
        if chosen is None:
            size = CALLOUT_VALUE_SIZES[-1]
            chosen = (size, METRICS.width_in(value, vf, vb, size), size * 0.95 / 72, n)
            room = int((inner_h - flag_room - 28 * 0.95 / 72 - 0.02) / pitch)
            F.fail(f"{lab} {name}", f"callout does not fit a {w:.2f} x {h:.2f} in slot: {n} lines of 14 pt under the value, "
                                    f"room for {room} (about {int(room * inner_w / 0.095)} characters): shorten the label, "
                                    "drop the context (the caption carries n and the unit), or give the callout both slots")
        size, vw, vh, n = chosen
        els_boxes = {
            "value": ((x + inset, y + inset + flag_room, inner_w, size * 1.2 / 72), "t"),
            "label": ((x + inset, y + inset + flag_room + vh + 0.02, inner_w, max(pitch, inner_h - flag_room - vh - 0.02)), "t"),
        }
    if size < 28:
        F.warn(f"{lab} {name}", f'callout value "{value}" set at {size} pt to fit: shorten it (write "31%" rather than '
                                '"31 percent"; move the unit into the label)')
    els = []
    if co.get("panel", True):
        panel = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))  # 1 = rectangle
        panel.name = f"{name} callout panel"
        panel.fill.solid()
        panel.fill.fore_color.rgb = RGBColor.from_string(MIST)
        panel.line.fill.background()
        panel.shadow.inherit = False
        tf = panel.text_frame
        tf.text = ""
        cnv = panel._element.find(".//" + qn("p:cNvPr"))
        if alt:
            cnv.set("descr", str(alt).strip())
        els.append(panel._element)
    vbox, vanchor = els_boxes["value"]
    val_tb = _textbox(slide, f"{name} callout", vbox, vanchor, [(value, vf, size, vb, vc)])
    paras_lbl = [(label, lf, 14, lb, lc)]
    if context:
        paras_lbl.append((context, cf, 14, cb, cc))
    lbox, lanchor = els_boxes["label"]
    lbl_tb = _textbox(slide, f"{name} callout label", lbox, lanchor, paras_lbl, pitch_in=pitch)
    for tb in (val_tb, lbl_tb):
        els.append(tb._element)
    for el in els:
        el.getparent().remove(el)
    if not alt:
        F.info(f"{lab} {name}", "callout: the number and label are real text, so no alt text is needed; "
                                "\"alt\" is stored on the panel if you give one")
    return els


# ----------------------------------------------------------------------------------------------------------
# Native charts (NASA figure text: Arial 14 pt black)
# ----------------------------------------------------------------------------------------------------------
def _font(font, size=14, bold=False, color=BLACK):
    font.name, font.size, font.bold = "Arial", Pt(size), bold
    font.color.rgb = RGBColor.from_string(color)


LINE_IN = 0.26  # one line of 14 pt chart text with its padding, in inches


def _fmt_value(v, fmt):
    """Approximate Excel number formats of the kind used on data labels: General, 0, 0.0, 0"%", 0.0" ms"."""
    m = re.match(r'^0(?:\.(0+))?(.*)$', str(fmt or "General"))
    if m:
        return f"{v:.{len(m.group(1) or '')}f}" + m.group(2).replace('"', "").replace("\\", "")
    return f"{v:g}"


def _manual_layout(parent, x, y, w=None, h=None, inner=False):
    """Give a chart element (plotArea, title, legend) a manual position; fractions of the chart frame."""
    lay = parent.find(qn("c:layout"))
    if lay is None:
        lay = etree.Element(qn("c:layout"))
        if parent.tag == qn("c:plotArea"):
            parent.insert(0, lay)
        else:
            tx = parent.find(qn("c:tx")) if parent.tag == qn("c:title") else parent.find(qn("c:legendPos"))
            (tx.addnext(lay) if tx is not None else parent.insert(0, lay))
    for c in list(lay):
        lay.remove(c)
    ml = etree.SubElement(lay, qn("c:manualLayout"))
    if inner:
        etree.SubElement(ml, qn("c:layoutTarget"), val="inner")
    etree.SubElement(ml, qn("c:xMode"), val="edge")
    etree.SubElement(ml, qn("c:yMode"), val="edge")
    etree.SubElement(ml, qn("c:x"), val=f"{max(0.0, x):.4f}")
    etree.SubElement(ml, qn("c:y"), val=f"{max(0.0, y):.4f}")
    if w is not None:
        etree.SubElement(ml, qn("c:w"), val=f"{max(0.05, w):.4f}")
        etree.SubElement(ml, qn("c:h"), val=f"{max(0.05, h):.4f}")


def _no_wrap(title_el):
    """Axis titles wrap at the axis length in LibreOffice (and in narrow PowerPoint charts): join the words with
    no-break spaces so a title stays on its one reserved line."""
    for t in title_el.iter(qn("a:t")):
        t.text = (t.text or "").replace(" ", "\u00a0")
    body = title_el.find(".//" + qn("c:rich") + "/" + qn("a:bodyPr"))
    if body is not None:
        body.set("wrap", "none")


def _chart_label(chart, text):
    """The unit label of a vertical axis, set as the chart title and later placed top left. LibreOffice wraps an
    unrotated axis title at the axis length; a chart title keeps its line."""
    chart.has_title = True
    chart.chart_title.text_frame.text = text
    _font(chart.chart_title.text_frame.paragraphs[0].runs[0].font)
    title = chart._chartSpace.find(qn("c:chart")).find(qn("c:title"))
    ov = title.find(qn("c:overlay"))
    if ov is None:
        ov = etree.SubElement(title, qn("c:overlay"))
        title.find(qn("c:tx")).addnext(ov)
    ov.set("val", "0")
    return title


def _layout_chart(chart, ch, kind, cats, series, size, labels, ticks, multi):
    """Place the plot area by hand so 14 pt axis text fits a quadrant: measured label widths on the left, one text
    line per title, and a top row for the flag, the legend or a horizontal axis title. Renderers then agree."""
    W, H = size
    cs = chart._chartSpace
    plot_area = cs.find(".//" + qn("c:plotArea"))
    ca_el, va_el = cs.find(".//" + qn("c:catAx")), cs.find(".//" + qn("c:valAx"))
    skip = etree.Element(qn("c:tickLblSkip"), val="1")  # never let a renderer drop a category label
    nm = ca_el.find(qn("c:noMultiLvlLbl"))
    nm.addprevious(skip) if nm is not None else ca_el.append(skip)
    fmt = ch.get("label_format", "General")
    vals = [float(v) for s in series for v in s["values"]]
    width = lambda s: METRICS.width_in(str(s), "Arial", False, 14)  # noqa: E731
    dlw = max(width(_fmt_value(v, fmt)) for v in vals) if labels else 0.0
    flag_w = width(ch["flag"]) + 0.1 if ch.get("flag") else 0.0
    legend_w = sum(width(s.get("name", "")) + 0.45 for s in series) if multi else 0.0
    top_items = [w_ for w_ in (legend_w, flag_w) if w_]
    vt, ct = ch["value_title"], ch.get("category_title")
    cat_title_el = ca_el.find(qn("c:title"))
    val_title_el = va_el.find(qn("c:title"))
    for el in (cat_title_el, val_title_el, cs.find(qn("c:chart")).find(qn("c:title"))):
        if el is not None:
            _no_wrap(el)
    if kind == "bar":
        catw = max(width(c) for c in cats)
        left, right = 0.05 + catw + 0.1, 0.06 + (dlw + 0.1 if labels else 0.1)
        bottom = 0.04 + LINE_IN + (LINE_IN if ticks else 0)
        rows = 1 if top_items and sum(top_items) + 0.2 <= W else len(top_items)
        if flag_w and not multi:  # the flag needs no row of its own when the top bar and its label stop short of it
            lo = float(ch.get("min", 0))
            hi = float(ch.get("max", max(vals) * 1.2 if max(vals) > 0 else 1))
            first = float(series[0]["values"][0])
            bar_end = left + (W - left - right) * max(0.0, (first - lo) / (hi - lo)) + dlw + 0.12
            if bar_end < W - flag_w:
                rows = 0
        top = 0.04 + LINE_IN * rows
        if ct:
            if width(ct) + 0.1 <= H - top - bottom:
                left += LINE_IN  # rotated title beside the category labels
            else:  # too long to stand on its side: it becomes the label above the plot, top left
                if rows == 0 or sum(top_items) + width(ct) + 0.3 > W:
                    top += LINE_IN
                cat_title_el.getparent().remove(cat_title_el)
                _manual_layout(_chart_label(chart, ct), 0.02 / W, 0.02 / H)
    else:
        tickw = max(width(_fmt_value(v, "General")) for v in vals + [float(ch.get("max", max(vals)))]) if ticks else 0.0
        left = 0.05 + (tickw + 0.12 if ticks else 0.08)
        right = 0.12 + (dlw / 2 if labels else 0)
        items = [width(vt) + 0.1] + top_items
        rows = 1 if sum(items) + 0.3 <= W else 2
        top = 0.04 + LINE_IN * rows + (LINE_IN if labels else 0.12)  # headroom: labels or the top tick label
        bottom = 0.04 + LINE_IN + (LINE_IN if ct else 0)
        _manual_layout(cs.find(qn("c:chart")).find(qn("c:title")), 0.02 / W, 0.02 / H)
    if multi:
        leg = cs.find(".//" + qn("c:legend"))
        _manual_layout(leg, (0.02 + (width(vt) + 0.3 if kind != "bar" and rows == 1 else 0)) / W,
                       (0.02 + (LINE_IN if kind != "bar" and rows == 2 else 0)) / H, legend_w / W, LINE_IN / H)
    _manual_layout(plot_area, left / W, top / H, (W - left - right) / W, (H - top - bottom) / H, inner=True)
    return W - left - right, H - top - bottom


def add_chart(slide, ch, box, alt, name, F, lab):
    kind = ch.get("type", "bar")
    xl = {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
          "line": XL_CHART_TYPE.LINE_MARKERS}.get(kind)
    if xl is None:
        raise SpecError(f'{lab} {name}: chart type must be bar, column or line (got "{kind}")')
    cats = [str(c) for c in ch.get("categories", [])]
    series = ch.get("series") or []
    if not cats or not series:
        raise SpecError(f"{lab} {name}: a chart needs categories and at least one series")
    for s in series:
        if len(s.get("values", [])) != len(cats):
            raise SpecError(f'{lab} {name}: series "{s.get("name")}" has {len(s.get("values", []))} values for {len(cats)} categories')
    if len(series) > 4:
        F.warn(f"{lab} {name}", f"{len(series)} series in a quadrant figure; keep one to three")
    if len(cats) == 1 and len(series) == 1:
        F.warn(f"{lab} {name}", "a chart with one data point is a number, not a comparison, and leaves the panel mostly "
                                'empty: use a callout figure instead ({"callout": {"value": "31%", "label": "..."}}), '
                                "or add the other conditions from the paper (never invent them)")
    vt = ch.get("value_title")
    if not vt:
        raise SpecError(f'{lab} {name}: "value_title" is required: label the value axis with its unit, e.g. "Latency (ms)" (GSFC rule 9)')
    if not re.search(r"\(.+\)", vt):
        F.warn(f"{lab} {name}", f'value axis title "{vt}" has no unit in parentheses (GSFC rule 9: units of measurement)')
    ct = ch.get("category_title")
    if not ct and not any(re.search(r"\d\s*[A-Za-z%°µ]", c) for c in cats):
        F.warn(f"{lab} {name}", "the category axis has no title and its labels carry no unit: add category_title")
    rev = kind == "bar"  # horizontal bars list the first category at the top
    data = CategoryChartData()
    data.categories = cats[::-1] if rev else cats
    for s in series:
        vals = [float(v) for v in s["values"]]
        data.add_series(str(s.get("name", "Series")), vals[::-1] if rev else vals)
    x, y, w, h = box
    gf = slide.shapes.add_chart(xl, Inches(x), Inches(y), Inches(w), Inches(h), data)
    chart = gf.chart
    _font(chart.font)
    chart.has_title = False
    multi = len(series) > 1
    chart.has_legend = multi
    if multi:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        _font(chart.legend.font)
    plot = chart.plots[0]
    plot.vary_by_categories = False
    if kind != "line":
        plot.gap_width = int(ch.get("gap_width", 55))
        if multi:
            plot.overlap = 0
    highlight = ch.get("highlight", [])
    highlight = [highlight] if isinstance(highlight, int) else list(highlight)
    for si, s in enumerate(plot.series):
        col = RGBColor.from_string(str(series[si].get("color", CATEGORICAL[si % len(CATEGORICAL)])).lstrip("#"))
        if kind == "line":
            s.format.line.color.rgb = col
            s.format.line.width = Pt(2.5)
            s.smooth = False
            s.marker.format.fill.solid()
            s.marker.format.fill.fore_color.rgb = col
        else:
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = col
            if not multi:
                for hi in highlight:
                    pi = (len(cats) - 1 - hi) if rev else hi
                    pt = s.points[pi].format
                    pt.fill.solid()
                    pt.fill.fore_color.rgb = RGBColor.from_string(GOLD)
    labels = ch.get("data_labels", True)
    if labels:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = ch.get("label_format", "General")
        dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.ABOVE if kind == "line" else XL_LABEL_POSITION.OUTSIDE_END
        _font(dl.font)
        body = dl._element.find(qn("c:txPr") + "/" + qn("a:bodyPr"))
        if body is not None:
            body.set("wrap", "none")  # a value label never breaks ("15 / %") in a narrow slot
    elif highlight:
        F.warn(f"{lab} {name}", "a gold highlight without data labels: gold is 2.0:1 on white, label the values (relief rule)")
    va, ca = chart.value_axis, chart.category_axis
    if "min" in ch:
        va.minimum_scale = float(ch["min"])
    elif kind != "line":
        va.minimum_scale = 0
    if "max" in ch:
        va.maximum_scale = float(ch["max"])
    va.has_major_gridlines = bool(ch.get("gridlines", False))
    if va.has_major_gridlines:
        va.major_gridlines.format.line.color.rgb = RGBColor.from_string(GRID)
        va.major_gridlines.format.line.width = Pt(0.75)
    va.format.line.fill.background()
    ticks = ch.get("value_ticks", not labels)
    va.major_tick_mark = XL_TICK_MARK.NONE
    va.tick_label_position = XL_TICK_LABEL_POSITION.NEXT_TO_AXIS if ticks else XL_TICK_LABEL_POSITION.NONE
    _font(va.tick_labels.font)
    if kind == "bar":  # the value axis runs along the bottom: a normal axis title
        va.has_title = True
        va.axis_title.text_frame.text = vt
        _font(va.axis_title.text_frame.paragraphs[0].runs[0].font)
    else:  # a vertical value axis: its label goes above the plot, top left (a rotated title outgrows the plot)
        _chart_label(chart, vt)
    ca.format.line.color.rgb = RGBColor.from_string(AXIS)
    ca.major_tick_mark = XL_TICK_MARK.NONE
    _font(ca.tick_labels.font)
    if ct:
        ca.has_title = True
        ca.axis_title.text_frame.text = ct
        _font(ca.axis_title.text_frame.paragraphs[0].runs[0].font)
    pw, ph = _layout_chart(chart, ch, kind, cats, series, (w, h), labels, ticks, multi)
    if ph < 0.55 or pw < 1.5:
        F.warn(f"{lab} {name}", f"the plot area is only {pw:.2f} x {ph:.2f} in after its 14 pt text: drop the category title "
                                "(put the unit in the labels), label series directly instead of a legend, or use a "
                                "horizontal bar chart")
    if not alt:
        fmt_unit = "%" if "%" in str(ch.get("label_format", "")) else ""
        parts = []
        for s in series:
            vals = ", ".join(f"{c} {v:g}{fmt_unit}" for c, v in zip(cats, [float(v) for v in s["values"]]))
            parts.append(f'{s.get("name", "Series")}: {vals}')
        alt = f"{kind.capitalize()} chart of {vt}" + (f" by {ct}" if ct else "") + ". " + "; ".join(parts) + "."
        F.info(f"{lab} {name}", f'alt text composed from the data: "{alt}"; write your own "alt" for the message')
    cnv = gf._element.find(".//" + qn("p:cNvPr"))
    cnv.set("descr", str(alt).strip())
    cnv.set("name", f"{name} chart")
    el = gf._element
    el.getparent().remove(el)
    return el


# ----------------------------------------------------------------------------------------------------------
# Checking (also runs on quads built or edited by hand)
# ----------------------------------------------------------------------------------------------------------
def _defrpr_chain(run_el, p_el, shape_el, layout_el, master_style):
    ppr = p_el.find(qn("a:pPr"))
    lvl = int(ppr.get("lvl", 0)) if ppr is not None else 0
    chain = []
    rpr = run_el.find(qn("a:rPr")) if run_el is not None else p_el.find(qn("a:endParaRPr"))
    if rpr is not None:
        chain.append(rpr)
    for el in (shape_el, layout_el):
        if el is None:
            continue
        lst = el.find(".//" + qn("a:lstStyle"))
        if lst is not None:
            lv = lst.find(qn(f"a:lvl{lvl + 1}pPr"))
            if lv is not None and lv.find(qn("a:defRPr")) is not None:
                chain.append(lv.find(qn("a:defRPr")))
    if master_style is not None:
        lv = master_style.find(qn(f"a:lvl{lvl + 1}pPr"))
        if lv is not None and lv.find(qn("a:defRPr")) is not None:
            chain.append(lv.find(qn("a:defRPr")))
    return chain, lvl


def eff_run(run_el, p_el, shape_el, layout_el, master_style):
    chain, lvl = _defrpr_chain(run_el, p_el, shape_el, layout_el, master_style)
    size = face = color = bold = None
    for d in chain:
        if size is None and d.get("sz"):
            size = int(d.get("sz")) / 100
        if bold is None and d.get("b") is not None:
            bold = d.get("b") in ("1", "true")
        if face is None and d.find(qn("a:latin")) is not None:
            face = d.find(qn("a:latin")).get("typeface")
        if color is None and d.find(qn("a:solidFill") + "/" + qn("a:srgbClr")) is not None:
            color = d.find(qn("a:solidFill") + "/" + qn("a:srgbClr")).get("val").upper()
    if face in ("+mn-lt", "+mj-lt", None):
        face = "Source Sans 3" if face else face
    return {"size": size or 18.0, "face": face or "Source Sans 3", "color": color, "bold": bool(bold), "lvl": lvl}


def _ppr_attr(p_el, shape_el, layout_el, lvl, att):
    ppr = p_el.find(qn("a:pPr"))
    if ppr is not None and ppr.get(att) is not None:
        return int(ppr.get(att))
    for el in (shape_el, layout_el):
        if el is None:
            continue
        lst = el.find(".//" + qn("a:lstStyle"))
        lv = lst.find(qn(f"a:lvl{lvl + 1}pPr")) if lst is not None else None
        if lv is not None and lv.get(att) is not None:
            return int(lv.get(att))
    return 0


def _line_pitch(p_el, shape_el, layout_el, lvl, size):
    for src in [p_el.find(qn("a:pPr"))] + [
            (el.find(".//" + qn("a:lstStyle")).find(qn(f"a:lvl{lvl + 1}pPr")) if el is not None and el.find(".//" + qn("a:lstStyle")) is not None else None)
            for el in (shape_el, layout_el)]:
        if src is None:
            continue
        pts = src.find(qn("a:lnSpc") + "/" + qn("a:spcPts"))
        if pts is not None:
            return int(pts.get("val")) / 100
        pct = src.find(qn("a:lnSpc") + "/" + qn("a:spcPct"))
        if pct is not None:
            return size * 1.2 * int(pct.get("val")) / 100000
    return size * 1.2


def _chart_text_ok(chart_xml, where, F, nasa):
    """NASA: every chart text style Arial, >= 14 pt, black; labeled value axis with a unit."""
    def deleted_axis(el):
        for anc in el.iterancestors():
            if anc.tag in (qn("c:valAx"), qn("c:catAx"), qn("c:dateAx"), qn("c:serAx")):
                d = anc.find(qn("c:delete"))
                return d is not None and d.get("val") in ("1", "true")
        return False
    root = chart_xml
    top = root.find(qn("c:txPr"))
    if nasa and top is None:
        F.fail(where, "chart has no default text style: set all chart text to Arial 14 pt black")
    for d in root.iter(qn("a:defRPr"), qn("a:rPr")):
        if deleted_axis(d):
            continue
        sz = d.get("sz")
        if sz and int(sz) < 1400:
            F.fail(where, f"chart text at {int(sz) / 100:g} pt (NASA and brand floor: 14 pt)")
        if nasa:
            lat = d.find(qn("a:latin"))
            if lat is not None and lat.get("typeface") != "Arial":
                F.fail(where, f'chart text in {lat.get("typeface")}: NASA quads use Arial for all text')
            col = d.find(qn("a:solidFill") + "/" + qn("a:srgbClr"))
            if col is not None and col.get("val").upper() != BLACK:
                F.warn(where, f'chart text #{col.get("val")}: NASA figure text is black')
    ch_title = root.find(qn("c:chart")).find(qn("c:title"))
    label_above = "".join(x.text or "" for x in ch_title.iter(qn("a:t"))) if ch_title is not None else ""
    for ax in root.iter(qn("c:valAx")):
        d = ax.find(qn("c:delete"))
        if d is not None and d.get("val") in ("1", "true"):
            if nasa:
                F.warn(where, "value axis deleted: label the values directly and name the quantity and unit in the caption")
            continue
        t = ax.find(qn("c:title"))
        txt = "".join(x.text or "" for x in t.iter(qn("a:t"))) if t is not None else ""
        if not txt and ax.find(qn("c:axPos")) is not None and ax.find(qn("c:axPos")).get("val") in ("l", "r"):
            txt = label_above  # quad_chart.py labels a vertical value axis above the plot (chart title, top left)
        if not txt:
            (F.fail if nasa else F.warn)(where, "value axis has no title: label it with the unit, e.g. \"Latency (ms)\" (GSFC rule 9)")
        elif not re.search(r"\(.+\)", txt):
            F.warn(where, f'value axis title "{txt}" has no unit in parentheses')
    for ax in root.iter(qn("c:catAx")):
        d = ax.find(qn("c:delete"))
        if d is not None and d.get("val") in ("1", "true"):
            continue
        t = ax.find(qn("c:title"))
        if t is None and not label_above:
            cats = [v.text or "" for v in ax.getparent().iter(qn("c:v"))]
            if not any(re.search(r"\d\s*[A-Za-z%°µ]", c) for c in cats):
                F.warn(where, "category axis has no title and its labels carry no unit")


def check_presentation(path, lay, findings=None, notes_expected=True):
    F = findings or Findings()
    prs = Presentation(str(path))
    layouts = {L["name"]: L for L in lay["layouts"]}
    glyph_by_char = {g["char"]: (k, g) for k, g in lay["milestone_glyphs"].items()}
    prompts = {p["prompt"].strip() for L in lay["layouts"] for p in L["placeholders"] if p.get("prompt")}
    n_quads = 0
    for n, slide in enumerate(prs.slides, 1):
        L = layouts.get(slide.slide_layout.name)
        if L is None:
            F.info(f"slide {n}", f'layout "{slide.slide_layout.name}" is not an ATR quad layout; skipped')
            continue
        n_quads += 1
        nasa = L["name"].startswith("ATR - NASA")
        recs = {p["idx"]: p for p in L["placeholders"]}
        lay_ph = {ph.placeholder_format.idx: ph._element for ph in slide.slide_layout.placeholders}
        master = slide.slide_layout.slide_master.element
        styles = {"title": master.find(".//" + qn("p:titleStyle")), "body": master.find(".//" + qn("p:bodyStyle")),
                  "other": master.find(".//" + qn("p:otherStyle"))}
        slide_text = []
        figures = 0
        # ---- NASA headings live on the layout: verify them there ----
        if nasa:
            lay_texts = [sh.text_frame.text.strip() for sh in slide.slide_layout.shapes
                         if sh.has_text_frame and not sh.is_placeholder]
            pos = [lay_texts.index(h) if h in lay_texts else -1 for h in NASA_HEADINGS]
            missing = [h for h, p_ in zip(NASA_HEADINGS, pos) if p_ < 0]
            if missing:
                F.fail(f"slide {n} layout", f"NASA headings missing or reworded: {', '.join(missing)} (GSFC rule 5)")
        for sh in slide.shapes:
            where = f"slide {n} › {sh.name}"
            el = sh._element
            idx = sh.placeholder_format.idx if sh.is_placeholder else None
            rec = recs.get(idx)
            layout_el = lay_ph.get(idx)
            # ---- pictures and charts: alt text, figure count ----
            if sh.shape_type is not None and sh.shape_type in (13,) or el.tag == qn("p:pic"):
                cnv = el.find(".//" + qn("p:cNvPr"))
                descr = (cnv.get("descr") or "").strip()
                if not descr:
                    F.fail(where, "picture without alt text (descr)")
                elif re.fullmatch(r"(image|picture|img|photo|figure)[ _-]?\d*(\.\w+)?|.*\.(png|jpe?g|gif|svg|tiff?)", descr, re.I):
                    F.warn(where, f'alt text "{descr}" is a file or shape name; say what the picture shows')
                if nasa and sh.top is not None and 2.73 * EMU_IN < sh.top < 4.64 * EMU_IN and sh.left < 5.0 * EMU_IN:
                    figures += 1
                continue
            if getattr(sh, "has_chart", False) and sh.has_chart:
                cnv = el.find(".//" + qn("p:cNvPr"))
                if not (cnv.get("descr") or "").strip():
                    F.fail(where, "chart without alt text (descr)")
                _chart_text_ok(sh.chart._chartSpace, where, F, nasa)
                if nasa and 2.73 * EMU_IN < sh.top + sh.height / 2 < 4.64 * EMU_IN and sh.left < 5.0 * EMU_IN:
                    figures += 1
                continue
            if getattr(sh, "has_table", False) and sh.has_table:
                check_table(sh, where, F, glyph_by_char, nasa)
                slide_text.append("\n".join(c.text for r_ in sh.table.rows for c in r_.cells))
                continue
            if not sh.has_text_frame:
                continue
            text = sh.text_frame.text
            slide_text.append(text)
            if nasa and re.fullmatch(r"Figure \d+ callout", sh.name or "") and 2.73 * EMU_IN < sh.top < 4.64 * EMU_IN \
                    and sh.left < 5.0 * EMU_IN:
                figures += 1  # a callout figure (big number + label) built by add_callout
            ph_type = sh.placeholder_format.type if sh.is_placeholder else None
            master_style = styles["title"] if (ph_type is not None and "TITLE" in str(ph_type)) else (
                styles["body"] if sh.is_placeholder else styles["other"])
            # ---- leftovers ----
            if sh.is_placeholder and rec is not None and rec["ph_type"] != "pic":
                if not text.strip():
                    F.warn(where, "empty placeholder left on the slide (delete it: it prints nothing but shows its prompt while editing)")
                    continue
                if text.strip() in prompts:
                    F.fail(where, "the template prompt is still there: replace it or delete the placeholder")
            if sh.is_placeholder and rec is not None and rec["ph_type"] == "pic":
                F.warn(where, "empty picture slot left on the slide (insert the figure or delete the slot)")
                continue
            found = list(dict.fromkeys(PLACEHOLDER_RE.findall(text)))
            if found:
                F.warn(where, "placeholder" + ("s" if len(found) > 1 else "") + " left in: " + ", ".join(found))
            # ---- runs: floor, NASA font and color ----
            measured = []
            for p in sh._element.iter(qn("a:p")):
                runs = p.findall(qn("a:r")) + p.findall(qn("a:fld"))
                ptxt = "".join((r.find(qn("a:t")).text or "") for r in runs if r.find(qn("a:t")) is not None)
                props = None
                for r in runs:
                    t = r.find(qn("a:t"))
                    if t is None or not (t.text or "").strip():
                        continue
                    e = eff_run(r, p, sh._element, layout_el, master_style)
                    props = props or e
                    if e["size"] < 14:
                        F.fail(where, f'"{t.text[:30]}" at {e["size"]:g} pt (14 pt floor)')
                    if nasa:
                        if e["face"] != "Arial":
                            F.fail(where, f'"{t.text[:30]}" in {e["face"]}: NASA quads use Arial for all text (GSFC rule 7)')
                        want = BLACK if (rec and rec["key"] == "q3cap") or sh.name.startswith(("Figure", "Illustrative")) else None
                        callout_value = sh.name.endswith(" callout") and e["color"] == NAVY  # the big number is navy
                        if want and e["color"] and e["color"] != BLACK and not callout_value:
                            F.warn(where, f"figure text #{e['color']}: NASA figure text is black")
                        if not want and e["color"] not in (NAVY, BLACK, None) and "tag" not in sh.name.lower():
                            F.warn(where, f"main text #{e['color']}: NASA main text is blue (ATR Navy #003976)")
                if props is not None:
                    measured.append((ptxt, props, p))
            # ---- fit and budgets (placeholders with a record) ----
            if rec is not None and measured:
                check_fit(sh, rec, measured, layout_el, where, F)
            # ---- NASA fields ----
            if nasa and rec is not None:
                if rec["key"] == "citation" and not re.search(r"doi|10\.\d{4,}", text, re.I):
                    F.warn(where, "no DOI in the citation line (GSFC rule 10)")
                if rec["key"] == "q2" and "nasa" not in text.lower():
                    F.warn(where, "Analysis does not name the NASA resources used (GSFC rule 8)")
        body = "\n".join(slide_text)
        if nasa:
            flat = re.sub(r"\s+", " ", body)
            m = ACK_RE.search(flat)
            if not m:
                F.fail(f"slide {n}", "the NASA acknowledgement sentence is missing or reworded (GSFC rule 6, "
                                     f"{GSFC_URL}): " + NASA_ACK.format(grant="[No.]", program="[program]"))
            else:
                if PLACEHOLDER_RE.search(m.group("grant")) or m.group("grant").strip() in ("xxxx", "< xxxx >"):
                    F.warn(f"slide {n} › acknowledgement", "grant number still a placeholder")
                prog = m.group("program")
                if PLACEHOLDER_RE.search(prog):
                    F.warn(f"slide {n} › acknowledgement", "program name still a placeholder")
                elif not prog.startswith("NASA "):
                    F.warn(f"slide {n} › acknowledgement", f'NASA\'s wording is "...part of the NASA <program> program": '
                                                          f'write "the NASA {prog} program"')
            if figures == 0:
                F.fail(f"slide {n}", "no figure in Results: GSFC rule 9 asks for 1 to 2 figures")
            elif figures > 2:
                F.fail(f"slide {n}", f"{figures} figures in Results: GSFC rule 9 asks for 1 to 2")
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
            if not notes and notes_expected:
                F.warn(f"slide {n} notes", "no notes: use them for methods detail and acronyms (GSFC rule 12)")
            if re.search(r"HOW TO START|duplicate this slide", notes, re.I):
                F.warn(f"slide {n} notes", "the template's instructions are still in the notes; replace them with your own")
        notes_all = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        left = sorted(set(PLACEHOLDER_RE.findall(notes_all)))
        if left:
            F.warn(f"slide {n} notes", f"{len(left)} placeholder(s) left: " + ", ".join(left)[:300])
        for m in re.finditer(r"\bKSU\b", body):
            F.info(f"slide {n}", 'write "Kent State University" or "Kent State", not "KSU", in running copy')
            break
    if n_quads == 0:
        F.fail(str(path), "no slide uses an ATR quad layout (build from assets/templates/ATR-Quad-Chart-Template.pptx)")
    nasa_count = sum(1 for s in prs.slides if s.slide_layout.name.startswith("ATR - NASA"))
    if nasa_count and len(prs.slides) > 1:
        F.info(str(Path(path).name), "the GSFC upload takes one quad per publication: save each NASA quad as its own one-slide file")
    return F, prs


def check_table(sh, where, F, glyph_by_char, nasa):
    tbl = sh._element.find(".//" + qn("a:tbl"))
    widths = [int(g.get("w")) / EMU_IN for g in tbl.find(qn("a:tblGrid")).findall(qn("a:gridCol"))]
    total_h = 0.0
    for ri, tr in enumerate(tbl.findall(qn("a:tr"))):
        row_lines = 1
        row_pitch = 15.5
        for ci, tc in enumerate(tr.findall(qn("a:tc"))):
            tcpr = tc.find(qn("a:tcPr"))
            ml = int(tcpr.get("marL", 91440)) if tcpr is not None else 91440
            mr = int(tcpr.get("marR", 91440)) if tcpr is not None else 91440
            width = widths[ci] - (ml + mr) / EMU_IN
            for p in tc.iter(qn("a:p")):
                runs = p.findall(qn("a:r"))
                txt = "".join(r.find(qn("a:t")).text or "" for r in runs)
                if not txt.strip():
                    continue
                e = eff_run(runs[0], p, None, None, None)
                if e["size"] < 14:
                    F.fail(f"{where} r{ri + 1}c{ci + 1}", f'"{txt[:30]}" at {e["size"]:g} pt (14 pt floor)')
                face = "Arial" if e["face"].startswith("Source Sans") else e["face"]
                lines = len(METRICS.wrap(txt, face, e["bold"], e["size"], width))
                pitch = _line_pitch(p, None, None, 0, e["size"])
                row_lines, row_pitch = max(row_lines, lines), pitch
                if lines > 1 and ri > 0:
                    F.warn(f"{where} r{ri + 1}c{ci + 1}", f'"{txt[:40]}" wraps to {lines} lines (with the Arial fallback); '
                                                         "shorten it so each milestone stays on one row")
                for r in runs:
                    t = (r.find(qn("a:t")).text or "").strip()
                    if t in glyph_by_char:
                        state, g = glyph_by_char[t]
                        col = eff_run(r, p, None, None, None)["color"]
                        if col and col != g["glyph_color"].upper():
                            F.fail(f"{where} r{ri + 1}c{ci + 1}", f'{g["label"]} marker "{t}" is #{col}; a typed marker is text: '
                                                                 f'use #{g["glyph_color"]} (color.md §6)')
                        if g["label"].lower() not in txt.lower() and not re.search(r"[A-Za-z]{3,}", txt.replace(t, "")):
                            F.fail(f"{where} r{ri + 1}c{ci + 1}", "status marker without its written label (never color or shape alone)")
        total_h += max(int(tr.get("h")) / EMU_IN, row_lines * row_pitch / 72 + 0.024)
    top = sh.top / EMU_IN
    if top + total_h > 5.02 + 1e-3:
        F.fail(where, f"the table runs to y {top + total_h:.2f} in, into the footer (limit 5.02 in): fewer or shorter rows")


def check_fit(sh, rec, measured, layout_el, where, F):
    """Measured wrap against the box, plus the character budget from quad-layouts.json."""
    key = rec["key"]
    w_in, h_in = sh.width / EMU_IN, sh.height / EMU_IN
    lines = lines_fb = 0
    chars = 0
    widest = 0.0
    pitch = None
    real_face = None
    for text, e, p in measured:
        indent = _ppr_attr(p, sh._element, layout_el, e["lvl"], "marL") / EMU_IN if rec.get("bullets") or e["lvl"] else 0
        width = w_in - indent
        face = e["face"]
        real_face = face
        fb = "Arial" if face.startswith("Source Sans") else face
        chars += len(text)
        pitch = pitch or _line_pitch(p, sh._element, layout_el, e["lvl"], e["size"])
        if rec.get("nowrap"):
            lines += 1
            lines_fb += 1
            widest = max(widest, METRICS.width_in(text, fb, e["bold"], e["size"]))
            continue
        lines += len(METRICS.wrap(text, face, e["bold"], e["size"], width))
        lines_fb += len(METRICS.wrap(text, fb, e["bold"], e["size"], width))
    max_lines = max(1, int(math.floor(h_in * 72 / pitch + 1e-6))) if pitch else rec.get("max_lines", 1)
    json_w = rec["box"]["w"]
    budget = rec.get("max_chars")
    if budget and abs(w_in - json_w) > 0.01:
        indent = 14 / 72 if rec.get("bullets") else 0
        budget = int(budget * (w_in - indent) / (json_w - indent))
    role = rec.get("role", key)
    label = f"{where} ({role})"
    if rec.get("nowrap"):
        if widest > w_in + 1e-3:
            F.fail(label, f"{widest:.2f} in wide in a {w_in:.2f} in one-line box (shorten it)")
    elif lines > max_lines:
        F.fail(label, f"{lines} lines in a box that holds {max_lines}: cut words (never shrink type below the layout size)")
    elif lines_fb > max_lines and real_face and real_face.startswith("Source Sans"):
        F.warn(label, f"fits in Source Sans 3 but needs {lines_fb} lines with the Arial fallback (box holds {max_lines}); "
                      "trim a few words so Google Slides and machines without the brand fonts do not overflow")
    if budget:
        if chars > budget * MAX_CHARS_FAIL:
            F.fail(label, f"{chars} characters, budget {budget} (+{chars / budget - 1:.0%}); shorten it")
        elif chars > budget:
            F.warn(label, f"{chars} characters, over the budget of {budget}; it fits as measured, but has no slack")
        elif (rec.get("bullets") and max_lines >= 3 and chars < budget * EMPTY_PANEL_FRACTION
              and lines <= max_lines * EMPTY_PANEL_FRACTION):
            F.warn(label, f"{chars} of {budget} characters ({chars / budget:.0%}) on {lines} of {max_lines} lines: the "
                          "panel will look empty; add a second finding or a one-line 'why it matters', never filler")
    F.info(label, f"{chars}/{budget or '-'} chars, {lines}/{max_lines} lines")


# ----------------------------------------------------------------------------------------------------------
# brand_check and render
# ----------------------------------------------------------------------------------------------------------
def run_brand_check(path, F, quiet):
    if not BRAND_CHECK.exists():
        F.warn("brand_check", f"{BRAND_CHECK} not found; skipped")
        return
    env = dict(os.environ, PYTHONDONTWRITEBYTECODE="1")
    r = subprocess.run([sys.executable, str(BRAND_CHECK), str(path), "--json"], capture_output=True, text=True, env=env)
    try:
        rep = json.loads(r.stdout)
    except ValueError:
        F.fail("brand_check", f"could not run: {(r.stderr or r.stdout)[-400:]}")
        return
    rep = rep[0] if isinstance(rep, list) else rep
    s = rep.get("summary", {})
    print(f"brand_check: {s.get('error', 0)} error(s), {s.get('warning', 0)} warning(s), {s.get('info', 0)} info")
    for f in rep.get("findings", []):
        sev = f.get("severity")
        if f.get("check") == "placeholder":  # the quad check above already lists every placeholder, notes included
            continue
        msg = f'[{f.get("check")}] {f.get("message")}' + (f' Fix: {f.get("fix")}' if f.get("fix") and sev != "info" else "")
        loc = f"brand_check {f.get('location') or ''}".strip()
        if sev == "error":
            F.fail(loc, msg)
        elif sev == "warning":
            F.warn(loc, msg)
        elif not quiet:
            F.info(loc, msg)


def run_render(path, outdir, dpi):
    outdir = Path(outdir) if outdir else path.parent / f"render-{path.stem}"
    r = subprocess.run(["bash", str(RENDER), str(path), str(outdir), str(dpi)], capture_output=True, text=True)
    if r.returncode != 0:
        print(f"render failed: {(r.stderr or r.stdout).strip()[-600:]}")
        return []
    pngs = [ln for ln in r.stdout.splitlines() if ln.endswith(".png")]
    print(f"rendered: {outdir / (path.stem + '.pdf')}")
    for p in pngs:
        print(f"  {p}")
    print("  Look at every page: the check measures text, only a render shows crops, overlaps and chart layout.")
    return pngs


# ----------------------------------------------------------------------------------------------------------
# CLI
# ----------------------------------------------------------------------------------------------------------
def list_fields(lay, which):
    for t in ([which] if which else list(TYPES)):
        name, _ = TYPES[t]
        L = next(x for x in lay["layouts"] if x["name"] == name)
        recs = {p["key"]: p for p in L["placeholders"]}
        print(f'{t}  (layout "{name}")')
        for key, target, what, req in FIELDS[t]:
            r = recs.get(target)
            budget = f'{r["max_chars"]} chars / {r["max_lines"]} lines' if r and r.get("max_chars") else ""
            print(f'  {key:34} {"required" if req else "optional":8}  {what}' + (f"  [{budget}]" if budget else ""))
        print()
    print("Every type: type (required), notes (speaker notes), example_tag (true or text: marks demo content).")
    print("Text fields: a string, a list of strings, or {\"text\", \"level\" 0-4, \"url\"}; **bold** runs.")
    print("Figures: {\"image\": path, \"alt\"}, {\"chart\": {...}, \"alt\"} or, for a result that is one number,")
    print("  {\"callout\": {\"value\": \"31%\", \"label\": \"fewer grasp failures\", \"context\": \"... n = 24\"}}.")
    print("Milestone status words: complete, on-track, at-risk, late, not-started (shape + label + color).")
    print("Examples: scripts/examples/quads/.  Rules: references/quad-charts.md.")


def main(argv=None):
    ap = argparse.ArgumentParser(description="Build ATR Lab quad charts from a JSON spec, or check a quad chart .pptx.",
                                 formatter_class=argparse.RawDescriptionHelpFormatter,
                                 epilog="Examples: scripts/examples/quads/*.json   Rules: references/quad-charts.md")
    ap.add_argument("input", nargs="?", help="spec .json to build, or a .pptx to check (with --check)")
    ap.add_argument("-o", "--out", help="output .pptx (default: <spec stem>.pptx in the current directory)")
    ap.add_argument("--check", action="store_true", help="also run brand_check.py (on a .pptx input: check only)")
    ap.add_argument("--render", nargs="?", const="", default=None, metavar="DIR",
                    help="render PDF + PNGs with scripts/render_office.sh (default DIR: render-<stem> beside the output)")
    ap.add_argument("--dpi", type=int, default=110, help="render resolution (default 110)")
    ap.add_argument("--allow-placeholders", action="store_true", help="build a draft that still has [bracketed] unknowns")
    ap.add_argument("--strict", action="store_true", help="warnings fail too")
    ap.add_argument("--quiet", action="store_true", help="hide info lines")
    ap.add_argument("--list-fields", nargs="?", const="", metavar="TYPE", help="print spec fields and budgets")
    ap.add_argument("--version", action="version", version=__version__)
    a = ap.parse_args(argv)

    try:
        lay = json.loads(LAYOUTS_JSON.read_text(encoding="utf-8"))
    except OSError as exc:
        print(f"quad_chart.py: cannot read {LAYOUTS_JSON}: {exc}", file=sys.stderr)
        return 2
    if a.list_fields is not None:
        if a.list_fields and a.list_fields not in TYPES:
            print(f"unknown type {a.list_fields!r}: use {', '.join(TYPES)}", file=sys.stderr)
            return 2
        list_fields(lay, a.list_fields or None)
        return 0
    if not a.input:
        ap.print_usage()
        return 2
    src = Path(a.input)
    if not src.exists():
        print(f"quad_chart.py: no such file: {src}", file=sys.stderr)
        return 2

    F = Findings()
    if src.suffix.lower() in (".pptx", ".potx"):  # check an existing quad
        out = src
        print(f"quad_chart: checking {out}")
    else:
        if not TEMPLATE.exists():
            print(f"quad_chart.py: template missing: {TEMPLATE}", file=sys.stderr)
            return 2
        try:
            spec = json.loads(src.read_text(encoding="utf-8"))
        except (OSError, ValueError) as exc:
            print(f"quad_chart.py: cannot read the spec {src}: {exc}", file=sys.stderr)
            return 2
        quads = spec["quads"] if isinstance(spec, dict) and "quads" in spec else [spec]
        out = Path(a.out) if a.out else Path.cwd() / f"{src.stem}.pptx"
        b = Builder(lay, a.allow_placeholders)
        try:
            prs = b.build(quads, src.resolve().parent)
        except SpecError as exc:
            print(f"quad_chart: spec refused: {exc}")
            print("RESULT: FAIL (nothing written)")
            return 1
        out.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(out))
        F.items.extend(b.findings.items)
        kinds = ", ".join(q.get("type", "?") for q in quads)
        print(f"quad_chart: wrote {out} ({len(quads)} slide{'s' if len(quads) > 1 else ''}: {kinds})")
    try:
        check_presentation(out, lay, F)
    except Exception as exc:  # a file python-pptx cannot open is a failure, not a crash
        F.fail(str(out), f"could not be checked: {exc}")
    if METRICS.estimated:
        F.info("metrics", f"no font file for {', '.join(sorted(METRICS.estimated))}: widths estimated from average advances "
                          "(install Arial or Liberation Sans, or set ATR_FONT_DIRS)")
    if a.check:
        run_brand_check(out, F, a.quiet)
    if a.render is not None and out.suffix.lower() == ".pptx":
        run_render(out, a.render or None, a.dpi)
    F.print(a.quiet)
    fails, warns = F.count("FAIL"), F.count("warn")
    bad = fails or (a.strict and warns)
    print(f"RESULT: {'FAIL' if bad else 'PASS'} ({fails} failure(s), {warns} warning(s))")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
