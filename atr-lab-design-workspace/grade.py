#!/usr/bin/env python3
"""Programmatic grader for atr-lab-design iteration runs.

Usage: python grade.py <iteration-dir>
Writes <run>/grading.json for every eval-*/{with_skill,without_skill} run, using the
skill-creator schema (expectations: text / passed / evidence). Judgment-only assertions
are added later by a human/LLM grader pass (merged, not overwritten).
"""
import json, os, re, subprocess, sys, glob
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SKILL = ROOT / "atr-lab-design"
PY = sys.executable
FORBIDDEN = ["@atr_kent", "College of Arts and Sciences", "Room 236", "330-672-9060",
             "Tele-Robotics", "Mathematics and Computer Science Building"]
BRAND_FONTS = {"Source Sans 3", "Source Sans 3 Semibold", "Source Sans 3 Black", "Source Sans 3 Light",
               "Source Sans 3 Medium", "Roboto Slab", "Roboto Slab SemiBold", "Roboto Slab Light",
               "Roboto Slab Medium", "Roboto Slab ExtraBold", "Source Code Pro", "Source Code Pro Medium",
               "Source Code Pro Semibold", "Arial"}
PALETTE = ["003976", "EFAB00", "1B2533", "4A5868", "8A6100", "F3F6FA", "D6DEE8", "FFFFFF",
           "00295F", "2C8ECD", "FFD702", "96A0A5", "B5B8B5"]


def brand_check(path, extra=()):
    env = dict(os.environ, PYTHONDONTWRITEBYTECODE="1")
    r = subprocess.run([PY, str(SKILL / "scripts/brand_check.py"), str(path), "--json", *extra],
                       capture_output=True, text=True, env=env)
    try:
        d = json.loads(r.stdout)
    except Exception:
        return None
    return d


def pptx_info(path):
    from pptx import Presentation
    from pptx.util import Pt
    prs = Presentation(str(path))
    texts, fonts, sizes, notes = [], set(), [], 0
    layouts = []
    for s in prs.slides:
        layouts.append(s.slide_layout.name)
        # static (non-placeholder) text on the layout is part of what the reader sees
        for lsh in s.slide_layout.shapes:
            if not lsh.is_placeholder and getattr(lsh, "has_text_frame", False) and lsh.has_text_frame:
                t = lsh.text_frame.text.strip()
                if t:
                    texts.append(t)
        if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
            notes += 1
        for sh in s.shapes:
            stack = [sh]
            while stack:
                x = stack.pop()
                if hasattr(x, "shapes"):
                    stack.extend(x.shapes)
                if getattr(x, "has_text_frame", False) and x.has_text_frame:
                    for p in x.text_frame.paragraphs:
                        for r in p.runs:
                            if r.text.strip():
                                texts.append(r.text)
                                if r.font.name:
                                    fonts.add(r.font.name)
                                if r.font.size:
                                    sizes.append(r.font.size.pt)
                if getattr(x, "has_table", False) and x.has_table:
                    for row in x.table.rows:
                        for c in row.cells:
                            texts.append(c.text)
    return {"n": len(prs.slides), "text": "\n".join(texts), "fonts": fonts, "sizes": sizes,
            "notes": notes, "layouts": layouts}


def exp(text, passed, evidence):
    return {"text": text, "passed": bool(passed), "evidence": evidence}


def forbidden_in(txt):
    return [f for f in FORBIDDEN if f.lower() in txt.lower()]


def all_text(outdir):
    buf = []
    for f in Path(outdir).rglob("*"):
        if f.suffix in (".md", ".txt"):
            buf.append(f.read_text(errors="ignore"))
        if f.suffix == ".pptx" and "DRAFT" not in f.name:
            try:
                buf.append(pptx_info(f)["text"])
            except Exception:
                pass
    return "\n".join(buf)


def grade_deck(out):
    e = []
    f = out / "sponsor-intro.pptx"
    e.append(exp("sponsor-intro.pptx exists", f.exists(), str(f)))
    if not f.exists():
        return e
    info = pptx_info(f)
    e.append(exp("Deck has 6-8 slides", 6 <= info["n"] <= 8, f"{info['n']} slides"))
    atr_layouts = sum(1 for l in info["layouts"] if l.startswith("ATR - "))
    e.append(exp("Slides are built on ATR template layouts", atr_layouts >= info["n"] - 1, f"layouts: {info['layouts']}"))
    fb = forbidden_in(all_text(out))
    e.append(exp("No forbidden/incorrect facts (@atr_kent, wrong college, Room 236, old phone, wrong name)", not fb, f"found: {fb}"))
    bad_fonts = sorted(x for x in info["fonts"] if x not in BRAND_FONTS)
    e.append(exp("Only brand fonts (or the Arial fallback) are used", not bad_fonts, f"non-brand: {bad_fonts}; all: {sorted(info['fonts'])}"))
    e.append(exp("Speaker notes on at least 75% of slides", info["notes"] >= 0.75 * info["n"], f"{info['notes']}/{info['n']} slides with notes"))
    e.append(exp("Closing/contact uses the verified lab website", "atr.cs.kent.edu" in info["text"], "found" if "atr.cs.kent.edu" in info["text"] else "not found"))
    d = brand_check(f)
    if d:
        errs = [x for x in d["findings"] if x["severity"] == "error"]
        e.append(exp("brand_check reports 0 errors", not errs, f"{d['summary']}; first: {[x['message'][:90] for x in errs[:3]]}"))
        logo = [x for x in d["findings"] if x["check"] in ("logo-athletic", "logo-retired") and x["severity"] == "error"]
        e.append(exp("No athletic Flash logo or retired block logo", not logo, f"{[x['message'][:80] for x in logo]}"))
    return e


def grade_quad(out):
    e = []
    f = out / "haptics-quad.pptx"
    e.append(exp("haptics-quad.pptx exists", f.exists(), str(f)))
    if not f.exists():
        return e
    info = pptx_info(f)
    heads = ["Background or Science Question", "Analysis", "Results", "Significance", "Acknowledgements"]
    missing = [h for h in heads if h.lower() not in info["text"].lower()]
    e.append(exp("All five NASA headings appear verbatim", not missing, f"missing: {missing}"))
    e.append(exp("DOI 10.1109/LRA.2026.1234567 appears on the slide", "10.1109/LRA.2026.1234567" in info["text"], ""))
    ack = all(k in info["text"] for k in ["National Aeronautics and Space Administration", "80NSSC26K1234", "EPSCoR"])
    e.append(exp("NASA acknowledgement sentence with grant number and EPSCoR program", ack, ""))
    small = [s for s in info["sizes"] if s < 14]
    e.append(exp("No explicit text size below 14 pt", not small, f"sizes <14: {sorted(set(small))}"))
    non_arial = sorted(x for x in info["fonts"] if x != "Arial")
    e.append(exp("Explicit fonts on the quad are Arial (NASA guidance)", not non_arial, f"non-Arial: {non_arial}"))
    e.append(exp("Speaker notes present", info["notes"] >= 1, f"{info['notes']} slides with notes"))
    fb = forbidden_in(all_text(out))
    e.append(exp("No forbidden/incorrect facts", not fb, f"found: {fb}"))
    d = brand_check(f)
    if d:
        errs = [x for x in d["findings"] if x["severity"] == "error"]
        e.append(exp("brand_check reports 0 errors", not errs, f"{d['summary']}; first: {[x['message'][:90] for x in errs[:3]]}"))
    return e


def grade_social(out):
    e = []
    pngs = [p for p in out.glob("*.png")]
    e.append(exp("A PNG image was delivered", bool(pngs), str([p.name for p in pngs])))
    if pngs:
        from PIL import Image
        import numpy as np
        im = Image.open(pngs[0]).convert("RGB")
        e.append(exp("Image is 1080 x 1350 (Instagram portrait)", im.size == (1080, 1350), f"{im.size}"))
        a = np.asarray(im.resize((270, 338))).reshape(-1, 3).astype(int)
        pal = np.array([[int(h[i:i + 2], 16) for i in (0, 2, 4)] for h in PALETTE])
        dist = np.sqrt(((a[:, None, :] - pal[None, :, :]) ** 2).sum(-1)).min(1)
        share = float((dist < 30).mean())
        e.append(exp("At least 70% of pixels are within the brand palette", share >= 0.70, f"{share:.0%} near-palette"))
    cap = out / "caption.txt"
    alt = out / "alt-text.txt"
    ctext = cap.read_text() if cap.exists() else ""
    e.append(exp("Caption delivered", len(ctext) > 40, f"{len(ctext)} chars"))
    e.append(exp("Alt text delivered (>= 80 characters)", alt.exists() and len(alt.read_text()) >= 80, f"{len(alt.read_text()) if alt.exists() else 0} chars"))
    ap = bool(re.search(r"a\.m\.", ctext)) and bool(re.search(r"p\.m\.", ctext)) and not re.search(r"\b\d{1,2}\s?(am|pm|AM|PM)\b", ctext)
    e.append(exp("Caption uses AP time style (a.m./p.m.)", ap, ""))
    fb = forbidden_in(all_text(out))
    e.append(exp("No forbidden/incorrect facts (e.g. @atr_kent)", not fb, f"found: {fb}"))
    handles = set(re.findall(r"@[A-Za-z0-9_.]+", ctext))
    ok_h = {"@atr_lab", "@kentstate", "@kentstatecs"}
    e.append(exp("Any lab handle in the caption is the verified Instagram handle @atr_lab", not (handles - ok_h) or all(h.lower() in ok_h for h in handles), f"handles: {sorted(handles)}"))
    return e


def grade_flyer(out):
    e = []
    f = out / "summer-camp-flyer-FIXED.pptx"
    e.append(exp("summer-camp-flyer-FIXED.pptx exists", f.exists(), str(f)))
    resp = (out / "response.md").read_text(errors="ignore") if (out / "response.md").exists() else ""
    checks = {
        "athletic logo": r"athletic|flash|golden flash",
        "retired block logo": r"retired|old (ATR )?logo|block[- ]letter",
        "@atr_kent handle": r"@atr_kent",
        "College of Arts and Sciences": r"arts and sciences",
        "white-on-gold contrast": r"contrast",
        "non-brand fonts": r"comic sans|papyrus|font",
        "Oxford comma": r"oxford|serial comma",
        "ampersand": r"&|ampersand",
        "wrong lab name": r"tele-robotics",
        "wrong contact (room/phone)": r"room 236|672-9060|phone",
        "KSU abbreviation": r"\bKSU\b",
        "time/date style": r"a\.m\.|ordinal|14th|AP style",
        "missing alt text": r"alt text|alt-text",
        "small/low-contrast text": r"size|pt\b|too small|gray|grey",
    }
    hits = [k for k, rx in checks.items() if re.search(rx, resp, re.I)]
    e.append(exp("Response identifies at least 10 of 14 planted defect categories", len(hits) >= 10, f"{len(hits)}/14: {hits}"))
    if f.exists():
        info = pptx_info(f)
        fb = forbidden_in(info["text"])
        e.append(exp("Fixed flyer contains no forbidden/incorrect facts", not fb, f"found: {fb}"))
        bad_fonts = sorted(x for x in info["fonts"] if x not in BRAND_FONTS)
        e.append(exp("Fixed flyer uses only brand fonts (or Arial)", not bad_fonts, f"non-brand: {bad_fonts}"))
        oxford = bool(re.search(r"\w+, \w+(?: \w+)?, and \w+", info["text"]))
        e.append(exp("Fixed flyer has no Oxford comma and no '&' in copy", not oxford and "&" not in info["text"], f"oxford={oxford}, amp={'&' in info['text']}"))
        d = brand_check(f)
        if d:
            errs = [x for x in d["findings"] if x["severity"] == "error"]
            e.append(exp("brand_check on the fixed flyer reports 0 errors", not errs, f"{d['summary']}; first: {[x['message'][:90] for x in errs[:3]]}"))
            logo = [x for x in d["findings"] if x["check"] in ("logo-athletic", "logo-retired") and x["severity"] == "error"]
            e.append(exp("Fixed flyer has no athletic or retired logo", not logo, f"{[x['message'][:80] for x in logo]}"))
    return e


def grade_copy(out):
    e = []
    li, nb = out / "linkedin-post.md", out / "news-blurb.md"
    e.append(exp("Both linkedin-post.md and news-blurb.md exist", li.exists() and nb.exists(), ""))
    t = (li.read_text() if li.exists() else "") + "\n" + (nb.read_text() if nb.exists() else "")
    e.append(exp("Full lab name 'Advanced Telerobotics Research Lab' used", "Advanced Telerobotics Research Lab" in t, ""))
    e.append(exp("'Kent State University' named", "Kent State University" in t, ""))
    fb = forbidden_in(t)
    e.append(exp("No forbidden/incorrect facts", not fb, f"found: {fb}"))
    body = re.sub(r"`[^`]*`", "", re.sub(r"https?://\S+", "", t))
    e.append(exp("No '&' in the copy", "&" not in body, ""))
    e.append(exp("HRI 2026 named", "HRI 2026" in t or "HRI '26" in t, ""))
    e.append(exp("No 'KSU' abbreviation in running copy", not re.search(r"\bKSU\b", t), ""))
    return e


GRADERS = {"sponsor-intro-deck": grade_deck, "nasa-quad-chart": grade_quad, "k12-instagram-post": grade_social,
           "flyer-brand-review": grade_flyer, "vendobot-paper-announcement": grade_copy}

if __name__ == "__main__":
    it = Path(sys.argv[1])
    for ev in sorted(it.glob("eval-*")):
        name = ev.name[5:]
        for cfg in ("with_skill", "without_skill"):
            run = ev / cfg / "run-1" if (ev / cfg / "run-1").exists() else ev / cfg
            out = run / "outputs"
            if not out.exists():
                continue
            exps = GRADERS[name](out)
            gp = run / "grading.json"
            prev = json.loads(gp.read_text()) if gp.exists() else {}
            judged = [x for x in prev.get("expectations", []) if x.get("judgment")]
            exps = exps + judged
            passed = sum(x["passed"] for x in exps)
            g = {"expectations": exps,
                 "summary": {"passed": passed, "failed": len(exps) - passed, "total": len(exps),
                             "pass_rate": round(passed / len(exps), 3) if exps else 0}}
            gp.write_text(json.dumps(g, indent=2))
            print(f"{name:32s} {cfg:14s} {passed}/{len(exps)}")
