#!/usr/bin/env python3
"""Fill test for ATR-Quad-Chart-Template.pptx (python-pptx).
Three decks, each built from the layouts through the placeholder API (as a student's PowerPoint would), then measured
with the real font metrics (measure.py) and rendered with build/tools/render.sh so overflow and crops are visible:
  fill-budget.pptx  every layout filled exactly to the documented character budget (must fit: exit 1 otherwise), with a
                    level-2 and a level-3 bullet in one placeholder per layout (the list styles carry levels 1-5);
  fill-over.pptx    every layout overfilled to 150 % of the budget (must overflow somewhere: exit 1 otherwise);
  fill-edge.pptx    real-world edge cases (must fit): a real-length DOI citation, a two-line title, one wide figure with
                    figure 2 deleted (no orphan frame may remain), a 3:1 plot cropped by a 1.34:1 slot, a sponsor logo in
                    its slot, typed milestone rows in quadrant 3, a footer marking at the budget, a NASA slide with BOTH
                    figure slots deleted and a level-3 bullet, and the acknowledgement typed on every NASA slide.
Text is typed with NO run formatting, so the render also proves that the layout list styles carry the font, size, color,
spacing and bullets. Picture slots get a generated 3:1 test plot with axes, ticks and labels, so the crop that PowerPoint
applies to fill a slot is visible in the render; the sponsor slot gets a generated square test image.
The slide-number placeholder is cloned from the layout onto every new slide (python-pptx never clones it; PowerPoint does
when Header & Footer > Slide number is on), and the footer marking gets wrap="none" as PowerPoint inherits it.
Every NASA slide is asserted to carry the acknowledgement sentence as slide text.
Usage: $SCRATCH/venv/bin/python fill_test.py [template.pptx] [outdir]
"""
import copy, json, os, subprocess, sys
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import PP_PLACEHOLDER
from PIL import Image, ImageDraw, ImageFont

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import measure                                                           # wrap() / width_in() with the real TTF metrics

ROOT = '/Users/marcodotio/Developer/atr-lab-design-skill'
FL = '/private/tmp/claude-501/-Users-marcodotio-Developer-atr-lab-design-skill/80375296-6f20-4da8-a711-62359143039a/scratchpad/fontlib/'
TEMPLATE = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, 'atr-lab-design/assets/templates/ATR-Quad-Chart-Template.pptx')
OUTDIR = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, 'build/qa/quad')
JSON_PATH = os.path.join(os.path.dirname(TEMPLATE), 'quad-layouts.json')
RENDER = os.path.join(ROOT, 'build/tools/render.sh')
INK, NAVY, GOLD, SLATE = (27, 37, 51), (0, 57, 118), (239, 171, 0), (74, 88, 104)
WORDS = ('robot operator latency network gesture headset study result task error delay control autonomy signal '
         'camera arm rover model data field trial metric baseline overlay display feedback loop safety').split()
NASA_ACK = measure.NASA_ACK
INDENT = 14 / 72


def make_plot(path):
    """3:1 test plot (1500 x 500 px) with axes, tick labels and axis titles at the edges, so any slot crop shows."""
    W, H = 1500, 500
    im = Image.new('RGB', (W, H), 'white'); d = ImageDraw.Draw(im)
    f = ImageFont.truetype(FL + 'SourceSans3-Regular.ttf', 34); fb = ImageFont.truetype(FL + 'SourceSans3-Bold.ttf', 34)
    x0, y0, x1, y1 = 170, 50, 1440, 400
    d.line([(x0, y0), (x0, y1), (x1, y1)], fill=INK, width=3)
    vals = [12, 18, 25, 33, 41, 55, 62]; labels = ['0', '0.5', '1', '1.5', '2', '2.5', '3']
    n = len(vals); bw = (x1 - x0) / n * 0.6
    for i, v in enumerate(vals):
        cx = x0 + (x1 - x0) * (i + 0.5) / n; h = (y1 - y0) * v / 70
        d.rectangle([cx - bw / 2, y1 - h, cx + bw / 2, y1], fill=NAVY if i < n - 1 else GOLD)
        d.text((cx, y1 + 8), labels[i], fill=INK, font=f, anchor='ma')
    for t in (0, 35, 70):
        y = y1 - (y1 - y0) * t / 70
        d.line([(x0 - 10, y), (x0, y)], fill=INK, width=3); d.text((x0 - 16, y), str(t), fill=INK, font=f, anchor='rm')
    d.text(((x0 + x1) / 2, y1 + 52), 'Delay (s)', fill=INK, font=fb, anchor='ma')
    lab = Image.new('RGBA', (300, 50), (255, 255, 255, 0))
    ImageDraw.Draw(lab).text((150, 25), 'Path error (%)', fill=INK, font=fb, anchor='mm')
    lab = lab.rotate(90, expand=True); im.paste(lab, (30, int((y0 + y1) / 2 - 150)), lab)
    d.text((x1, y0 - 40), 'TEST PLOT 3:1 (fill test)', fill=SLATE, font=f, anchor='ra')
    im.save(path)


def make_logo(path):
    """Square test image for the sponsor slot (no real sponsor mark is shipped with the template)."""
    im = Image.new('RGB', (400, 400), NAVY); d = ImageDraw.Draw(im)
    fb = ImageFont.truetype(FL + 'SourceSans3-Bold.ttf', 64)
    d.rectangle([20, 20, 379, 379], outline=GOLD, width=8)
    d.text((200, 170), 'TEST', fill='white', font=fb, anchor='mm'); d.text((200, 240), 'LOGO', fill='white', font=fb, anchor='mm')
    im.save(path)


def filler(n_chars, lines, factor):
    """Deterministic pseudo-prose: `lines` paragraphs of at most n_chars * factor / lines characters each
    (a paragraph never exceeds its share, so an at-budget fill is a fair test of the documented budget)."""
    target = int(n_chars * factor)
    per = max(4, target // max(1, lines))
    out = []
    i = 0
    for k in range(lines):
        s = ''
        while True:
            w = WORDS[i % len(WORDS)]
            if len(s) + len(w) + (1 if s else 0) > per:
                break
            s += (' ' if s else '') + w; i += 1
        out.append(s[0].upper() + s[1:])
    return out


def drop_all_slides(prs):
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        prs.part.drop_rel(sldId.rId); sldIdLst.remove(sldId)


def clone_slide_number(slide, layout):
    """Copy the layout's slide-number placeholder (native field) onto the slide, as PowerPoint does for new slides."""
    for ph in layout.placeholders:
        if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            el = copy.deepcopy(ph._element)
            ids = [int(e.get('id')) for e in slide.shapes._spTree.iter() if e.tag.endswith('}cNvPr') and e.get('id')]
            el.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr').set('id', str(max(ids + [1]) + 1))
            slide.shapes._spTree.append(el)


class Deck:
    """A python-pptx deck built from the template's layouts, remembering what was typed where for the fit check."""
    def __init__(self, label):
        self.label = label
        self.spec = json.load(open(JSON_PATH))
        self.prs = Presentation(TEMPLATE)
        drop_all_slides(self.prs)
        self.typed = []            # (slide_no, layout, key, [(text, level)...], ph_record)

    def add(self, name):
        L = next(l for l in self.spec['layouts'] if l['name'] == name)
        layout = self.prs.slide_layouts.get_by_name(name)
        assert layout is not None, name
        slide = self.prs.slides.add_slide(layout)
        clone_slide_number(slide, layout)
        have = {ph.placeholder_format.idx: ph for ph in slide.placeholders}
        by_key = {}
        for ph in L['placeholders']:
            shp = have.get(ph['idx'])
            assert shp is not None, f"{name}: placeholder idx {ph['idx']} ({ph['key']}) not cloned onto the slide"
            by_key[ph['key']] = (shp, ph)
        return slide, by_key, L

    def text(self, phs, key, text, L):
        """Type paragraphs into a placeholder with no run formatting (inherits the layout list style).
        text: str, list of str, or list of (str, level)."""
        shp, ph = phs[key]
        items = text if isinstance(text, list) else [text]
        items = [(t, 0) if isinstance(t, str) else t for t in items]
        tf = shp.text_frame
        tf.paragraphs[0].text = items[0][0]; tf.paragraphs[0].level = items[0][1]
        for t, lvl in items[1:]:
            p = tf.add_paragraph(); p.text = t; p.level = lvl
        if ph.get('nowrap'):
            tf._txBody.bodyPr.set('wrap', 'none')      # PowerPoint inherits this from the layout; python-pptx writes an empty bodyPr
        self.typed.append((len(self.prs.slides), L['name'], key, items, ph))

    def picture(self, phs, key, path, stretch_w=None):
        shp, ph = phs[key]
        left, top, width, height = shp.left, shp.top, shp.width, shp.height   # inherited from the layout slot
        if stretch_w:                                    # set the full box BEFORE inserting (python-pptx crops to the slot aspect;
            width = Inches(stretch_w)                    # setting width alone would write an xfrm with zero height)
            shp.left, shp.top, shp.width, shp.height = left, top, width, height
        pic = shp.insert_picture(path)
        pic.left, pic.top, pic.width, pic.height = left, top, width, height   # and AFTER: the new pic element has no xfrm
        return pic

    def delete(self, phs, key):
        el = phs[key][0]._element; el.getparent().remove(el)

    def save(self):
        out = os.path.join(OUTDIR, f'fill-{self.label}.pptx')
        self.prs.save(out)
        return out

    def fit_report(self):
        """Measured line count vs box for every typed placeholder; returns the list of overflows."""
        over = []
        for slide_no, lname, key, items, ph in self.typed:
            f = ph['font']; b = ph['box']
            fallback = 'Arial' if f['face'] == 'Source Sans 3' else f['face']
            total = 0
            for t, lvl in items:
                w = b['w'] - (INDENT * (lvl + 1) if ph.get('bullets') else 0)
                if ph.get('nowrap'):
                    total += 1
                else:
                    total += len(measure.wrap(t, fallback, bool(f.get('bold')), f['size'], w, f.get('tracking', 0) or 0))
            need = total * f['line'] / 72
            widest = max(measure.width_in(t, fallback, bool(f.get('bold')), f['size']) for t, _ in items)
            fits = need <= b['h'] + 1e-6 and (not ph.get('nowrap') or widest <= b['w'] + 1e-6)
            if not fits:
                over.append(f'  slide {slide_no} {lname} / {key}: {total} line(s) x {f["line"]} pt = {need:.2f} in > box {b["h"]:.2f} in' + (f' (nowrap widest {widest:.2f} > {b["w"]:.2f})' if ph.get('nowrap') and widest > b['w'] else ''))
        return over

    def assert_nasa_ack(self):
        for i, slide in enumerate(self.prs.slides, 1):
            if slide.slide_layout.name.startswith('ATR - NASA'):
                text = '\n'.join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
                head, tail = NASA_ACK.split('[xxxx]')[0], NASA_ACK.split('[NASA program name]')[1]   # the fixed wording around the two values
                mid = 'and was part of the'
                assert head in text and mid in text and tail in text, f'{self.label} slide {i}: NASA slide without the acknowledgement sentence as slide text'


def build(factor, label, plot, logo):
    d = Deck(label)
    report = []
    for L in d.spec['layouts']:
        slide, phs, _ = d.add(L['name'])
        first_bullets = True
        for key, (shp, ph) in phs.items():
            if ph['ph_type'] == 'pic':
                d.picture(phs, key, logo if key == 'sponsor' else plot)
                continue
            n = ph['max_chars']; lines = ph['max_lines']
            if key == 'ack':
                text = [NASA_ACK]                              # NASA sentence stays verbatim
            elif key == 'footer_right':
                text = ['Internal · 2026-09-28' if factor <= 1 else 'Internal draft · 2026-09-28 · v3 · do not distribute']
            else:
                text = filler(n, lines, factor)
                if ph.get('bullets') and first_bullets and lines >= 3:
                    text = [(text[0], 0), (text[1], 1), (text[2], 2)] + text[3:]   # one level-2 and one level-3 item per layout
                    first_bullets = False
            d.text(phs, key, text, L)
            chars = sum(len(t if isinstance(t, str) else t[0]) for t in text)
            report.append(f"  {L['name']} / {key}: {chars} chars in {len(text)} para(s) (budget {n} chars, {lines} lines)")
    d.assert_nasa_ack()
    out = d.save()
    print(f'{label}: wrote {out}'); print('\n'.join(report))
    return d, out


def build_edge(plot, logo):
    d = Deck('edge')
    # --- NASA: real citation, two-line title, one wide figure (figure 2 deleted) -------------------------------------
    slide, ph, L = d.add('ATR - NASA Research Quad')
    d.text(ph, 'title', 'Predictive overlays keep rovers on course at 2 s of delay', L)     # 57 chars, 2 lines
    d.text(ph, 'citation', 'Kim et al. (2026), IEEE Trans. Robot., doi:10.1109/TRO.2026.1234567', L)
    d.text(ph, 'meta', ['Dr. Jong-Hoon Kim', 'CMS  ·  Sept. 2026'], L)
    d.text(ph, 'q1', ['Remote rovers receive every command late; past 1 s of delay drivers slow down and drift off the planned path.', 'Can a predictive overlay restore driving accuracy?'], L)
    d.text(ph, 'q2', ['12 volunteers drove a simulated rover at 0.5, 1 and 2 s of delay, with and without the overlay.', 'NASA resources used: [testbed], [dataset]'], L)
    d.picture(ph, 'q3fig1', plot, stretch_w=4.375)                                                     # 3:1 into 3.77:1: small crop
    d.delete(ph, 'q3fig2')                                                                             # delete figure 2: nothing may remain
    d.text(ph, 'q3cap', 'Fig. 1: path error (%) by delay (s), n = 12', L)
    d.text(ph, 'q4', ['The overlay held path error near the no-delay level at 2 s, where direct video driving failed.', 'One operator could drive safely at delays that today require full autonomy.'], L)
    d.text(ph, 'ack', NASA_ACK.replace('[xxxx]', '80NSSC21K0123').replace('[NASA program name]', 'Carbon Monitoring System (CMS)'), L)
    # --- Project Status: sponsor logo, cropped 3:1 plot, typed milestone rows, footer marking at the budget -------------
    slide, ph, L = d.add('ATR - Project Status Quad')
    d.text(ph, 'title', 'Gesture-controlled robot for remote lab tours', L)
    d.text(ph, 'meta', ['PI: J.-H. Kim  ·  POC: J. Kim', 'NSF CISE  ·  Award 2412345', 'Oct. 2025 to Sept. 2027'], L)   # 2.70 in wide beside the sponsor slot
    d.picture(ph, 'sponsor', logo)
    d.text(ph, 'q1', ['Remote visitors tour the lab by gesture.', 'Mobile base with a gimbal camera.'], L)
    d.picture(ph, 'q1img', plot)                                                                       # 3:1 into 1.34:1: heavy crop, visible
    d.text(ph, 'q2', ['Visitor hand pose maps to velocity commands.', 'A predictive display hides the 200 ms delay.', 'Front end at 30 fps on a fixed route.'], L)
    d.text(ph, 'q3', ['Kickoff review  ·  Jan 26  ·  Complete', 'Prototype v1 demo  ·  Mar 26  ·  On track', 'User study, n = 12  ·  May 26  ·  At risk', 'Paper submission  ·  Jul 26  ·  Late', 'Field deployment  ·  Oct 26  ·  Not started'], L)
    d.text(ph, 'q4', ['Impact: remote visits and K-12 outreach without travel.', 'Deliverables: ROS package, dataset, paper.', 'Next: outdoor trials on the campus loop.'], L)
    d.text(ph, 'footer_right', 'Proprietary · 2026-09-28', L)                                      # 24 chars: at the budget
    # --- Weekly Summary: normal fill with a link label --------------------------------------------------------------
    slide, ph, L = d.add('ATR - Weekly Summary')
    d.text(ph, 'title', 'Telepresence robot: weekly summary', L)
    d.text(ph, 'meta', ['Week of Sept. 28, 2026', '[Student name]  ·  [Advisor]', 'ATR Lab  ·  Kent State University'], L)
    d.text(ph, 'acc', ['Gesture front end runs at 30 fps on the laptop.', 'Base follows the fixed route with 12 cm mean error.', 'Draft of the user-study protocol reviewed.'], L)
    d.text(ph, 'path', ['Add the predictive display by Oct. 9.', 'Pilot with two lab members on Oct. 12.', 'Need: decision on the camera gimbal purchase.'], L)
    d.text(ph, 'risks', ['Gimbal lead time six weeks; order now.', 'Need weekend building access.'], L)
    d.text(ph, 'notes', ['Read: predictive display papers from the 1990s Mars rover work.'], L)
    d.text(ph, 'media', ['Gesture robot demo (video)', 'Watch: lab YouTube channel'], L)
    d.text(ph, 'footer_right', 'Internal · 2026-09-28', L)
    # --- NASA: both figure slots deleted (no orphan frames), level-2 and level-3 bullets, prompts otherwise short -------
    slide, ph, L = d.add('ATR - NASA Research Quad')
    d.text(ph, 'title', 'Orphan test: both figure slots deleted', L)
    d.text(ph, 'citation', 'Doe et al. (2026), J. Field Robotics, doi:10.1002/rob.00000', L)
    d.text(ph, 'meta', ['PI Name', 'Program  ·  Sept. 2026'], L)
    d.text(ph, 'q1', [('Level one', 0), ('Level two', 1), ('Level three stays 14 pt Arial', 2)], L)
    d.text(ph, 'q2', ['Analysis', 'NASA resources used: [x]'], L)
    d.delete(ph, 'q3fig1'); d.delete(ph, 'q3fig2')
    d.text(ph, 'q3cap', 'Fig. 1: slots deleted; nothing may remain above', L)
    d.text(ph, 'q4', ['Significance'], L)
    d.text(ph, 'ack', NASA_ACK, L)
    d.assert_nasa_ack()
    out = d.save()
    print(f'edge: wrote {out}')
    return d, out


def main():
    os.makedirs(OUTDIR, exist_ok=True)
    plot = os.path.join(OUTDIR, 'test-plot-3x1.png'); logo = os.path.join(OUTDIR, 'test-logo-square.png')
    make_plot(plot); make_logo(logo)
    decks = [build(1.0, 'budget', plot, logo), build(1.5, 'over', plot, logo), build_edge(plot, logo)]
    failed = False
    for d, path in decks:
        over = d.fit_report()
        if d.label == 'over':
            print(f'{d.label}: {len(over)} measured overflow(s) (expected: the deck is filled to 150 %)')
            if not over:
                print('  FAIL: the 150 % deck did not overflow anywhere, so the budgets are too loose'); failed = True
        else:
            print(f'{d.label}: {len(over)} measured overflow(s) (expected 0)')
            for o in over: print('  OVER' + o)
            failed |= bool(over)
        subprocess.run([RENDER, path, OUTDIR, '110'], check=True)
    print('rendered', [os.path.basename(p) for _, p in decks], 'into', OUTDIR)
    print('RESULT:', 'FAIL' if failed else 'PASS')
    sys.exit(1 if failed else 0)


if __name__ == '__main__':
    main()
